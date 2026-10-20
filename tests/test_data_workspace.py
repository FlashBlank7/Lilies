"""Workspace behavior: real parsing, scoped resources and real workflow execution."""
from concurrent.futures import ThreadPoolExecutor
import time
import pytest

from tests.test_projects import configured as configured
from tests.test_projects import graph, node, edge, ref
from tests.test_modeling import modeling as modeling, real_compute as real_compute, wait_task


def setup(configured, content='id,x,y,label\n0001,1,3,0\n0002,2,6,1\n0003,3,9,0\n'):
    client, _, project, settings = configured
    base = '/api/v1/projects/' + project['id']
    path = settings.workspace_root / project['id'] / 'requirement-package/data.csv'
    path.parent.mkdir(exist_ok=True)
    path.write_text(content)
    obj = client.get(base + '/data-workspace/object', params={'object_id': 'file:requirement-package/data.csv'})
    assert obj.status_code == 200, obj.text
    return client, base, obj.json(), path


def preview(client, base, obj, **kwargs):
    response = client.post(base + '/data-workspace/preview', json={'object_id': obj['id'], 'version': obj['version'], **kwargs})
    assert response.status_code == 200, response.text
    return response.json()


def run(client, base, obj, tool, **kwargs):
    response = client.post(base + '/data-workspace/operations', json={'request_key': tool + str(time.time_ns()),
        'source': {'id': obj['id'], 'version': obj['version']}, 'tool': tool, **kwargs})
    assert response.status_code == 202, response.text
    operation = response.json()
    task = wait_task(client, base, {'id': operation['task_id']})
    result = client.get(base + '/data-workspace/operations/' + operation['id'])
    assert result.status_code == 200, result.text
    return result.json(), task


def test_preview_preserves_empty_rows_tokens_and_exclusive_states(configured):
    client, base, obj, _ = setup(configured, 'id,x,label\n0001,1,0\n,,\n0003,/,1\n0004,3.5-4.2,0\n0005,bad,2\n')
    mapping = {'id_column': 'id', 'field_rules': {'x': {'missing_tokens': {'/': 'not_measured'}}}}
    value = preview(client, base, obj, mapping=mapping)
    assert value['total_rows'] == 5
    assert [r['source_line'] for r in value['rows']] == [2, 3, 4, 5, 6]
    assert value['rows'][0]['raw']['id'] == '0001'
    assert value['rows'][3]['raw']['x'] == '3.5-4.2'
    assert value['states']['counts'] == {'valid': 1, 'missing': 2, 'pending': 1, 'invalid': 1}
    subset = preview(client, base, obj, mapping=mapping, selection={'filters': [{'column': 'x', 'operator': 'state', 'value': 'pending'}]})
    assert subset['selected_rows'] == 1 and subset['rows'][0]['source_line'] == 5


def test_paged_full_table_filter_does_not_limit_to_preview(configured):
    content = ','.join(['x','y',*[f'field{i}' for i in range(28)]]) + '\n' + '\n'.join(','.join(map(str,[i,i*2,*range(28)])) for i in range(100001))
    client, base, obj, _ = setup(configured, content)
    value = preview(client, base, obj, offset=20, limit=10, x='x', y='y', selection={'filters': [{'column': 'x', 'operator': 'ge', 'value': 90000}]})
    assert value['total_rows'] == 100001
    assert value['selected_rows'] == 10001
    assert len(value['rows']) == 10 and value['rows'][0]['row_id'] == 90020
    assert len(value['points']) == 1000 and value['points_sampled']
    assert len(value['columns']) == 30
    assert len(__import__('json').dumps(value)) < 200000


def test_excel_preview_uses_selected_sheet_and_header(configured):
    from openpyxl import Workbook
    client, base, _, path = setup(configured)
    workbook = Workbook()
    workbook.active.title = '说明'
    workbook.create_sheet('样本')
    sheet = workbook['样本']
    sheet.append(['标题'])
    sheet.append(['id', 'x'])
    sheet.append(['副标题'])
    sheet.append(['001', 2])
    sheet.append([None, None])
    sheet.append(['003', 5])
    workbook.save(path.with_suffix('.xlsx'))
    obj = client.get(base+'/data-workspace/object', params={'object_id':'file:requirement-package/data.xlsx'}).json()
    result = preview(client, base, obj, mapping={'sheet':'样本','header_row':1,'skip_after_header':1,'id_column':'id'})
    assert result['sheets'] == ['说明', '样本']
    assert result['total_rows'] == 3
    assert result['rows'][0]['raw']['id'] == '001'
    assert result['rows'][0]['source_line'] == 4


def test_reject_changed_source_cross_project_and_symlink(configured, tmp_path):
    client, base, obj, path = setup(configured)
    path.write_text('x,y\n4,5\n')
    assert client.post(base+'/data-workspace/preview',json={'object_id':obj['id'],'version':obj['version']}).status_code == 409
    request={'request_key':'changed','source':{'id':obj['id'],'version':obj['version']},'tool':'profile'}
    assert client.post(base+'/data-workspace/operations',json=request).status_code == 409
    outside = tmp_path/'outside.csv'
    outside.write_text('secret')
    path.with_name('link.csv').symlink_to(outside)
    assert client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/link.csv'}).status_code == 422
    assert client.get(base+'/data-workspace/object',params={'object_id':'file:../outside.csv'}).status_code == 422
    dataset=client.post(base+'/datasets',json={'source_path':'requirement-package/data.csv'}).json()
    other=client.post('/api/v1/projects',json={'name':'other'}).json()
    assert client.get('/api/v1/projects/'+other['id']+'/data-workspace/object',params={'object_id':'dataset:'+dataset['id']}).status_code == 404
    assert client.get(base+'/tasks').json() == []


def test_workflow_drop_uses_declared_input_and_deduplicates_execution(configured):
    client, app, project, _ = configured
    client, base, obj, _ = setup(configured)
    graph(client, project['id'], [node('start','start',inputs=[{'name':'source_path','type':'string','resource_kind':'file'}]),
        node('end','end',outputs={'source':ref('$inputs','source_path')})], [edge('start','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    request={'request_key':'same-drop','source':{'id':obj['id'],'version':obj['version']},'tool':'workflow:'+project['id'],'workflow_revision':revision,'input_name':'source_path'}
    with ThreadPoolExecutor(2) as pool:
        responses=list(pool.map(lambda _:client.post(base+'/data-workspace/operations',json=request), range(2)))
    assert all(r.status_code==202 for r in responses), [r.text for r in responses]
    assert responses[0].json()['task_id']==responses[1].json()['task_id']
    task=wait_task(client,base,{'id':responses[0].json()['task_id']})
    assert task['status']=='succeeded', task
    space=client.get(base+'/data-workspace/catalog').json()
    assert len(space['workflows'])==1
    request['parameters']={'different':True}
    assert client.post(base+'/data-workspace/operations',json=request).status_code==409


def test_document_image_and_unsupported_objects_have_explicit_views(configured):
    from docx import Document
    client, base, _, path = setup(configured)
    doc=Document()
    doc.add_paragraph('倒焊项目说明')
    doc.save(path.with_name('notes.docx'))
    obj=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/notes.docx'}).json()
    assert '倒焊项目说明' in preview(client,base,obj)['text']
    path.with_name('opaque.bin').write_bytes(b'\x00\x01')
    obj=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/opaque.bin'}).json()
    assert preview(client,base,obj)=={'object':obj}
    response=client.post(base+'/data-workspace/operations',json={'request_key':'wrong','source':{'id':obj['id'],'version':obj['version']},'tool':'spearman'})
    assert response.status_code==422 and '表格数据' in response.text


def test_real_selection_analysis_training_prediction_and_saved_chain(real_compute):
    configured,_=real_compute
    client,base,obj,path=setup(configured,'chip,x,y,label\n'+'\n'.join(f'{i//2:03d},{i%9},{i/3},{i%2}' for i in range(60)))
    mapping={'target':'label','group_column':'chip','id_column':'chip','field_rules':{'x':{'dtype':'number'},'y':{'dtype':'number'},'label':{'dtype':'category','categories':['0','1'],'category_labels':{'0':'合格','1':'需复核'}}}}
    selected,task=run(client,base,obj,'select',mapping=mapping,selection={'filters':[{'column':'x','operator':'ge','value':1}]})
    assert task['status']=='succeeded',task
    data=next(o for o in selected['results'] if o.get('dataset_id'))
    table=preview(client,base,data)
    assert table['selected_rows']==53 and table['mapping']['field_rules']['label']['category_labels']['1']=='需复核'
    processed=next(o for o in selected['results'] if o.get('path','').endswith('/processed.csv'))
    via_file=preview(client,base,processed)
    assert via_file['rows'][0]['raw']['chip']=='000'
    assert via_file['rows'][0]['source_line']==3
    assert via_file['mapping']['field_rules']==table['mapping']['field_rules']
    analysis,task=run(client,base,data,'spearman',selection={'columns':['x','y']})
    assert task['status']=='succeeded',task
    assert any(o['kind']=='analysis' for o in analysis['results'])
    saved=client.post(base+'/data-workspace/operations/'+analysis['id']+'/save',json={'name':'选取与相关分析','request_key':'save-1'})
    assert saved.status_code==201,saved.text
    repeated=client.post(base+'/data-workspace/operations/'+analysis['id']+'/save',json={'name':'选取与相关分析','request_key':'save-1'})
    assert repeated.json()['workflow_id']==saved.json()['workflow_id']
    flow=client.get('/api/v1/applications/'+saved.json()['workflow_id']+'/draft').json()['snapshot']['workflow']
    assert len([n for n in flow['nodes'] if n['type']=='data_analysis'])==3
    path.with_name('new.csv').write_text('chip,x,y,label\n'+'\n'.join(f'{i:03d},{i%4},{100-i},{i%2}' for i in range(24)))
    new=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/new.csv'}).json()
    saved_draft=client.get('/api/v1/applications/'+saved.json()['workflow_id']+'/draft').json()
    reused,task=run(client,base,new,'workflow:'+saved.json()['workflow_id'],workflow_revision=saved_draft['revision'],input_name='source_path')
    assert task['status']=='succeeded',task
    correlation=next(o for o in reused['results'] if o.get('value',{}).get('operation')=='spearman')
    assert correlation['value']['rows']==18
    trained,task=run(client,base,data,'train',selection={'columns':['x','y']},parameters={'evaluation':{'problem':'classification','metric':'macro_f1','split':'group','holdout_fraction':0},'candidate':{'models':['linear'],'batch_size':1},'budget':{'seconds':60,'trials':1,'trial_seconds':60}})
    assert task['status']=='succeeded',task
    model=next(o for o in trained['results'] if o['kind']=='model')
    predicted,task=run(client,base,data,'predict',model={'id':model['id'],'version':model['version']})
    assert task['status']=='succeeded',task
    result=next(o for o in predicted['results'] if o['name']=='逐行预测结果')
    predictions=preview(client,base,result)
    assert predictions['total_rows']==53
    assert predictions['rows'][0]['source_line']==3
    assert predictions['rows'][0]['raw']['chip']=='000'
    assert predictions['rows'][0]['values']['source_row']==1
    assert set(predictions['columns']) >= {'prediction','prediction_name','source_line'}
    saved_predict=client.post(base+'/data-workspace/operations/'+predicted['id']+'/save',json={'name':'固定模型预测','request_key':'save-2'}).json()
    flow=client.get('/api/v1/applications/'+saved_predict['workflow_id']+'/draft').json()['snapshot']['workflow']
    assert not any(n['type']=='model_train' for n in flow['nodes'])
    candidate_count=len(client.get(base+'/data-workspace/catalog').json()['models'])
    saved_draft=client.get('/api/v1/applications/'+saved_predict['workflow_id']+'/draft').json()
    again,task=run(client,base,new,'workflow:'+saved_predict['workflow_id'],workflow_revision=saved_draft['revision'],input_name='source_path')
    assert task['status']=='succeeded',task
    assert len(client.get(base+'/data-workspace/catalog').json()['models'])==candidate_count
    result=next(o for o in again['results'] if o['name']=='逐行预测结果')
    assert preview(client,base,result)['total_rows']==18
    dropped,task=run(client,base,data,'outliers',selection={'columns':['x']},parameters={'method':'bounds','lower':2,'handling':'drop'})
    assert task['status']=='succeeded',task
    sub=next(o for o in dropped['results'] if o.get('dataset_id'))
    assert preview(client,base,sub)['rows'][0]['source_row']==2


def test_dataset_version_rejects_changed_frozen_content(configured):
    client, app, project, _ = configured
    client, base, _, _ = setup(configured)
    dataset = client.post(base+'/datasets', json={'source_path':'requirement-package/data.csv'}).json()
    obj = client.get(base+'/data-workspace/object', params={'object_id':'dataset:'+dataset['id']}).json()
    file = app.state.services.modeling.path(project['id'], dataset['id']) / dataset['files']['source']['name']
    file.write_text('id,x,y,label\n001,999,1,0\n')
    response=client.post(base+'/data-workspace/operations',json={'request_key':'tampered','source':{'id':obj['id'],'version':obj['version']},'tool':'profile'})
    assert response.status_code==409
    assert client.post(base+'/data-workspace/preview',json={'object_id':obj['id'],'version':obj['version']}).status_code==409
    assert client.get(base+'/tasks').json()==[]


def test_failure_does_not_reuse_another_runs_success_report(configured):
    client, app, project, settings = configured
    client, base, obj, _ = setup(configured)
    report = settings.workspace_root / project['id'] / 'results/report.html'
    report.parent.mkdir(exist_ok=True)
    report.write_text('<h1>Only the successful run</h1>')
    graph(client,project['id'],[node('start','start',inputs=[{'name':'source_path','resource_kind':'file'}]),
        node('end','end',outputs={'report':'results/report.html','input':ref('$inputs','source_path')})],[edge('start','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    success,task=run(client,base,obj,'workflow:'+project['id'],workflow_revision=revision,input_name='source_path')
    assert task['status']=='succeeded'
    assert any(value.get('path') == 'results/report.html' for value in success['results'])
    graph(client,project['id'],[node('start','start',inputs=[{'name':'source_path','resource_kind':'file'}]),
        node('step','code',code='raise ValueError("验收中的故意失败")'),node('end','end',outputs={'result':ref('step','output')})],[edge('start','step'),edge('step','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    failed,task=run(client,base,obj,'workflow:'+project['id'],workflow_revision=revision,input_name='source_path')
    assert task['status']=='failed'
    assert failed['results']==[]
    assert failed['task_id']!=success['task_id']
    assert client.get(base+'/data-workspace/operations/'+success['id']).json()['status']=='succeeded'


def test_failed_workflow_call_keeps_its_completed_step_artifact(configured):
    client, _, project, settings = configured
    client, base, obj, _ = setup(configured)
    artifact = settings.workspace_root / project['id'] / 'results/partial.csv'
    artifact.parent.mkdir(exist_ok=True)
    artifact.write_text('x,y\n1,2\n')
    graph(client, project['id'], [
        node('start', 'start', inputs=[{'name':'source_path','resource_kind':'file'}]),
        node('completed', 'template_transform', template='results/partial.csv', variables={}),
        node('failed', 'code', code='raise ValueError("后续步骤失败，保留已完成产物")'),
        node('end', 'end', outputs={'result':ref('failed','output')}),
    ], [edge('start','completed'), edge('completed','failed'), edge('failed','end')])
    revision = client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    operation, task = run(client, base, obj, 'workflow:'+project['id'], workflow_revision=revision, input_name='source_path')
    assert task['status'] == 'failed'
    assert [value.get('path') for value in operation['results']] == ['results/partial.csv']
    assert preview(client, base, operation['results'][0])['total_rows'] == 1


def test_numeric_equality_and_literal_category_identity_use_different_semantics(configured):
    client,base,obj,_=setup(configured,'code,x\n001,1.0\n1,2.5\n002,\n')
    mapping={'field_rules':{'code':{'dtype':'text'}}}
    numeric=preview(client,base,obj,mapping=mapping,selection={'filters':[{'column':'x','operator':'eq','value':'1'}]})
    assert numeric['selected_rows']==1 and numeric['rows'][0]['raw']['code']=='001'
    text=preview(client,base,obj,mapping=mapping,selection={'filters':[{'column':'code','operator':'eq','value':'1'}]})
    assert text['selected_rows']==1 and text['rows'][0]['row_id']==1


def test_file_input_requires_declared_fields_before_creating_task(configured):
    client, _, project, _=configured
    client,base,obj,_=setup(configured)
    graph(client,project['id'],[node('start','start',inputs=[{'name':'source_path','resource_kind':'file','required_columns':['missing']}]),node('end','end',outputs={})],[edge('start','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    response=client.post(base+'/data-workspace/operations',json={'source':{'id':obj['id'],'version':obj['version']},'request_key':'missing-columns','tool':'workflow:'+project['id'],'workflow_revision':revision,'input_name':'source_path'})
    assert response.status_code==422 and 'missing' in response.text
    assert client.get(base+'/tasks').json()==[]


def test_scatter_sampling_uses_valid_pairs_instead_of_total_rows(configured):
    client,base,obj,_=setup(configured,'x,y\n1,2\n'+',\n'*1200+'3,4\n')
    result=preview(client,base,obj,x='x',y='y')
    assert result['total_rows']==1202
    assert len(result['points'])==2 and not result['points_sampled']


def test_saved_excel_chain_can_change_layout_for_csv_without_losing_first_row(real_compute):
    from openpyxl import Workbook
    configured,_=real_compute
    client,base,_,path=setup(configured)
    workbook=Workbook(); sheet=workbook.active
    sheet.append(['x','y']); sheet.append(['副表头','副表头'])
    for i in range(12):sheet.append([i,i*2])
    workbook.save(path.with_suffix('.xlsx'))
    obj=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/data.xlsx'}).json()
    operation,task=run(client,base,obj,'spearman',mapping={'skip_after_header':1},selection={'columns':['x','y']})
    assert task['status']=='succeeded',task
    saved=client.post(base+'/data-workspace/operations/'+operation['id']+'/save',json={'name':'换表相关分析','request_key':'layout'}).json()
    path.write_text('x,y\n'+''.join(f'{i},{30-i}\n' for i in range(12)))
    obj=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/data.csv'}).json()
    draft=client.get('/api/v1/applications/'+saved['workflow_id']+'/draft').json()
    result,task=run(client,base,obj,'workflow:'+saved['workflow_id'],workflow_revision=draft['revision'],input_name='source_path',parameters={'source_sheet':0,'source_header_row':0,'source_skip_after_header':0})
    assert task['status']=='succeeded',task
    assert next(o['value']['rows'] for o in result['results'] if o.get('value',{}).get('operation')=='spearman')==12


def test_preview_does_not_hide_columns_that_start_with_sqlalchemy_prefix(configured):
    client,base,obj,_=setup(configured,'_sample,_sales,x\n0,001,1\n1,002,2\n')
    result=preview(client,base,obj)
    assert result['rows'][0]['values']['_sample']==0
    assert result['rows'][0]['raw']['_sales']=='001'


def test_pdf_preview_renders_one_bounded_page_and_rejects_unknown_page(configured):
    import base64
    from pypdf import PdfWriter
    client,base,_,path=setup(configured)
    writer=PdfWriter();writer.add_blank_page(width=612,height=792);writer.add_blank_page(width=612,height=792)
    writer.write(path.with_suffix('.pdf'))
    obj=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/data.pdf'}).json()
    result=preview(client,base,obj,offset=1)
    assert result['page']==1 and result['pages']==2
    assert base64.b64decode(result['image'].split(',')[1]).startswith(b'\x89PNG')
    assert len(result['image'])<200000
    response=client.post(base+'/data-workspace/preview',json={'object_id':obj['id'],'version':obj['version'],'offset':2})
    assert response.status_code==422


def test_agent_uses_workspace_objects_and_waits_for_the_actual_task(configured):
    from agent_platform.project_agent_tools import WorkspaceProjectTools
    client,app,project,_=configured
    client,base,obj,_=setup(configured)
    graph(client,project['id'],[node('start','start',inputs=[{'name':'source_path','resource_kind':'file'}]),node('end','end',outputs={'ready':True})],[edge('start','end')])
    tools=WorkspaceProjectTools(app.state.services,project['id'],app.state.services.local_agents)
    listed=client.portal.call(tools.call,'project_data_workspace',{'action':'list'})
    assert any(o['id']==obj['id'] for o in listed['objects'])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    response=client.portal.call(tools.call,'project_data_workspace',{'action':'run','wait':True,'request':{'request_key':'agent-1','source':{'id':obj['id'],'version':obj['version']},'tool':'workflow:'+project['id'],'workflow_revision':revision,'input_name':'source_path'}})
    assert response['status']=='succeeded'
    assert client.get(base+'/tasks/'+response['task_id']).json()['status']=='succeeded'
    assert client.get(base+'/data-workspace/operations/'+response['id']).json()['task_id']==response['task_id']


@pytest.mark.parametrize('source_kind',['file','dataset'])
def test_default_secondary_file_is_frozen_and_shown_as_an_input(configured,source_kind):
    client,_,project,_=configured
    client,base,obj,path=setup(configured)
    if source_kind=='dataset':
        dataset=client.post(base+'/datasets',json={'source_path':obj['path']}).json()
        obj=client.get(base+'/data-workspace/object',params={'object_id':'dataset:'+dataset['id']}).json()
    path.with_name('second.csv').write_text('x\noriginal\n')
    graph(client,project['id'],[node('start','start',inputs=[{'name':'source_path','resource_kind':'file'},
        {'name':'aux_path','resource_kind':'file','default':'requirement-package/second.csv'},
        {'name':'header','type':'number','purpose':'header_row','default':1}]),
        node('end','end',outputs={'second':ref('$inputs','aux_path')})],[edge('start','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    operation,task=run(client,base,obj,'workflow:'+project['id'],workflow_revision=revision,input_name='source_path')
    assert task['status']=='succeeded',task
    assert len(operation['inputs'])==2
    frozen=next(o for o in operation['results'] if o.get('name')=='second.csv')
    path.with_name('second.csv').write_text('x\nchanged\n')
    assert preview(client,base,frozen)['rows'][0]['raw']['x']=='original'
    saved=client.post(base+'/data-workspace/operations/'+operation['id']+'/save',json={'name':'双输入复用','request_key':'two-inputs'}).json()
    draft=client.get('/api/v1/applications/'+saved['workflow_id']+'/draft').json()
    inputs=draft['snapshot']['workflow']['nodes'][0]['config']['inputs']
    auxiliary=next(f for f in inputs if f['name']!='source_path' and f.get('resource_kind')=='file')
    assert auxiliary['resource_kind']=='file'
    assert next(f for f in inputs if f.get('purpose')=='header_row')['default']==1
    replacement=client.get(base+'/data-workspace/object',params={'object_id':'file:requirement-package/second.csv'}).json()
    reused,task=run(client,base,obj,'workflow:'+saved['workflow_id'],workflow_revision=draft['revision'],input_name='source_path',
        bindings={auxiliary['name']:{'id':replacement['id'],'version':replacement['version']}})
    assert task['status']=='succeeded',task
    changed=next(o for o in reused['results'] if o.get('name')=='second.csv')
    assert preview(client,base,changed)['rows'][0]['raw']['x']=='changed'
    assert preview(client,base,frozen)['rows'][0]['raw']['x']=='original'


def test_saved_result_workflow_preserves_result_input_type(configured):
    client,_,project,_=configured
    client,base,obj,_=setup(configured)
    graph(client,project['id'],[node('start','start',inputs=[{'name':'source_path','resource_kind':'file'}]),
        node('end','end',outputs={'answer':42})],[edge('start','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    original,task=run(client,base,obj,'workflow:'+project['id'],workflow_revision=revision,input_name='source_path')
    result=next(o for o in original['results'] if o['kind']=='analysis')
    graph(client,project['id'],[node('start','start',inputs=[{'name':'analysis','type':'object','resource_kind':'result'}]),
        node('end','end',outputs={'received':ref('$inputs','analysis')})],[edge('start','end')])
    revision=client.get('/api/v1/applications/'+project['id']+'/draft').json()['revision']
    operation,task=run(client,base,result,'workflow:'+project['id'],workflow_revision=revision,input_name='analysis')
    assert task['status']=='succeeded',task
    saved=client.post(base+'/data-workspace/operations/'+operation['id']+'/save',json={'name':'复用分析结果','request_key':'result-input'}).json()
    draft=client.get('/api/v1/applications/'+saved['workflow_id']+'/draft').json()
    field=draft['snapshot']['workflow']['nodes'][0]['config']['inputs'][0]
    assert field['resource_kind']=='result' and field['type']=='object'
    reused,task=run(client,base,result,'workflow:'+saved['workflow_id'],workflow_revision=draft['revision'],input_name=field['name'])
    assert task['status']=='succeeded',task
    assert reused['results'][0]['value']['result']==operation['results'][0]['value']

"""Original layout, row locations and the full business vocabulary survive subsets."""
import io
import json
import zipfile

import pytest

from agent_platform.modeling_models import DataAnalysisPlan, DataMapping
from agent_platform import modeling_worker as worker
from tests.test_modeling import modeling, real_compute, wait_task  # noqa: F401
from tests.test_projects import configured, graph, node, edge, ref, start  # noqa: F401


def source(path):
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active; sheet.title = 'Sheet2'
    sheet.append(['id', 'height', 'pressure', 'quality'])
    sheet.append(['编号', 'mm', 'N', '编码'])
    for i in range(64):
        sheet.append([f'{i:03d}', i + 1, i % 11 + 2, str(i % 3)])
    for cell, value in [('B6', '3.5-4.2'), ('B7', '/'), ('D8', '-1')]:
        sheet[cell] = value
    for cell in sheet[9]:
        cell.value = None  # A physical blank record, not permission to delete it.
    book.save(path)
    return DataMapping(sheet='Sheet2', skip_after_header=1, id_column='id', target='quality', invalid_rows='exclude',
        field_rules={'height': {'pending_tokens': ['/']}, 'pressure': {}, 'quality': {
            'dtype': 'category', 'categories': ['0', '1', '2'],
            'category_labels': {'0': '合格', '1': '压连不足', '2': '过压损伤', '-1': '未判定'},
            'missing_tokens': {'-1': 'unjudged'}}}).model_dump()


@pytest.mark.parametrize('indices', [None, [7, 0, 6, 2], []])
def test_prepare_roundtrip_and_reordered_subset_preserve_every_selected_position(tmp_path, indices):
    pytest.importorskip('pandas')
    path = tmp_path / 'source.xlsx'; mapping = source(path)
    config = {'source': str(path), 'mapping': mapping, 'output': str(tmp_path / 'output'),
              'analysis': DataAnalysisPlan(operation='prepare', row_indices=indices).model_dump()}
    original = path.read_bytes()
    result = worker.analyze(config)
    state = json.loads((tmp_path / 'output/data-state.json').read_text())
    selected = list(range(64)) if indices is None else indices
    assert result['rows'] == 64 and result['output_rows'] == len(selected)
    assert state['rules'] == mapping['field_rules']
    assert all(state['mapping'][k] == 0 for k in ('sheet', 'header_row', 'skip_after_header'))
    frame = worker.read_data({**config, 'source': str(tmp_path / 'output/processed.csv'),
                             'state': str(tmp_path / 'output/data-state.json'), 'mapping': state['mapping']})
    cells = [r for r in frame.attrs['data_state']['cells'] if r['column'] == 'height']
    assert [r['source_row'] for r in cells] == sorted(selected)  # Stored cells retain source order.
    assert {r['row']: r['source_line'] for r in cells} == {i: old + 3 for i, old in enumerate(selected)}
    assert len(frame) == len(selected) and path.read_bytes() == original
    if indices is None:
        assert cells[0]['raw'] == '1' and cells[0]['source_line'] == 3
        assert cells[6]['status'] == 'missing'
        with pytest.raises(ValueError, match='布局'):
            worker.read_data({**config, 'source': str(tmp_path / 'output/processed.csv'),
                             'state': str(tmp_path / 'output/data-state.json'),
                             'mapping': {**state['mapping'], 'skip_after_header': 1}})
    config['analysis']['row_indices'] = [64]
    with pytest.raises(ValueError, match='越界'):
        worker.analyze(config)


def test_csv_rejects_excel_sheet_and_retains_empty_records(tmp_path):
    pytest.importorskip('pandas')
    path = tmp_path / 'data.csv'; path.write_text('x\n1\n\n3\n')
    assert len(worker.read_table(path, {})) == 3
    with pytest.raises(ValueError, match='CSV/TSV'):
        worker.read_table(path, {'sheet': 'Sheet2', 'skip_after_header': 1})
    for indices in ([1, 1], [-1]):
        with pytest.raises(ValueError):
            DataAnalysisPlan(operation='prepare', row_indices=indices)


def test_prepare_training_and_same_workflow_replacement_keep_rows_and_absent_class_definitions(real_compute):
    from openpyxl import load_workbook
    config, service = real_compute
    client, app, project, settings = config
    root = settings.workspace_root / project['id']; package = root / 'requirement-package'; package.mkdir(exist_ok=True)
    path = package / 'data.xlsx'; mapping = source(path); original = path.read_bytes()
    base = '/api/v1/projects/' + project['id']
    dataset = client.post(base + '/datasets', json={'source_path': 'requirement-package/data.xlsx', 'mapping': mapping}).json()
    # Explicit subset excludes missing/pending features and -1/blank labels but retains their definitions.
    selected = [i for i in range(64) if i not in {3, 4, 5, 6}]
    response = client.post(base + '/agent-tools', json={'name': 'project_modeling', 'arguments': {
        'action': 'analyze', 'dataset_id': dataset['id'], 'analysis': {'operation': 'prepare', 'row_indices': selected}}})
    assert response.status_code == 200, response.text
    prepared = response.json()
    assert prepared['rows'] == 64 and prepared['output_rows'] == 60
    derived = client.get(base + '/datasets/' + prepared['dataset_id']).json()
    repeated = client.post(base + '/datasets', json={'source_path': derived['files']['source']['original']})
    assert repeated.status_code == 201, repeated.text
    assert repeated.json()['mapping'] == derived['mapping']
    assert repeated.json()['mapping']['field_rules'] == mapping['field_rules']
    for bad in ({'skip_after_header': 1}, {'sheet': 'Sheet2'}, {'field_rules': {}}):
        response = client.post(base + '/datasets', json={'source_path': derived['files']['source']['original'], 'mapping': bad})
        assert response.status_code == 422, response.text
    graph(client, project['id'], [node('start', 'start'),
        node('train', 'model_train', dataset_id=repeated.json()['id'], evaluation={'problem': 'classification', 'metric': 'macro_f1', 'folds': 2},
             features={'columns': ['height', 'pressure']}, candidate={'engine': 'sklearn', 'models': ['forest'], 'parameters': {'n_estimators': 8}, 'batch_size': 1},
             budget={'seconds': 180, 'trials': 1, 'trial_seconds': 45}),
        node('end', 'end', outputs={'training': ref('train', 'output')})], [edge('start', 'train'), edge('train', 'end')])
    task = wait_task(client, base, start(client, base, 'prepare-training', workflow_id=project['id']))
    assert task['status'] == 'succeeded', task
    trained = task['outputs']['training']
    split = json.loads((service.path(project['id'], trained['study_id']) / 'output/split.json').read_text())
    assert len(split['samples']) == 60
    response = client.put(base + '/models/prepared', json={'name': '完整定义', 'study_id': trained['study_id'], 'candidate_id': trained['id'], 'slot': 0})
    assert response.status_code == 200, response.text
    bundle = client.get(base + f'/modeling/studies/{trained["study_id"]}/candidates/{trained["id"]}/download')
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
        assert json.loads(archive.read('predict.example.json'))['mapping']['field_rules'] == mapping['field_rules']
    response = client.post(base + '/agent-tools', json={'name': 'project_skills', 'arguments': {
        'action': 'read', 'skill_id': 'modeling', 'reference': 'prepare-workflow.json'}})
    workflow = json.loads(response.json()['content'])
    for field in workflow['nodes'][0]['config']['inputs']:
        if field['name'] != 'source_path':
            field.update(required=False, default={'mapping': mapping, 'analysis_columns': ['height', 'pressure'], 'model_ref': 'prepared'}[field['name']])
    graph(client, project['id'], workflow['nodes'], workflow['edges'])
    revision = client.get('/api/v1/applications/' + project['id'] + '/draft').json()['revision']
    book = load_workbook(path); book['Sheet2']['B3'] = '待复测'; book['Sheet2']['C4'] = 99; book.save(package / 'changed.xlsx')
    outputs = []
    for n, name in enumerate(['data.xlsx', 'changed.xlsx']):
        task = wait_task(client, base, start(client, base, f'prepare-predict-{n}', workflow_id=project['id'], inputs={'source_path': 'requirement-package/' + name}))
        assert task['status'] == 'succeeded', task
        out = task['outputs']; outputs.append(out)
        assert out['表格与源行']['rows'] == out['表格与源行']['output_rows'] == 64
        assert out['相关分析']['rows'] == 64 and out['相关分析']['pairs'][0]['samples'] == 61 - n
        assert out['逐条预测']['rows'] == 64 and out['逐条预测']['unavailable_rows'] == 3 + n
        state = json.loads((root / out['逐条预测']['data_state_path']).read_text())
        assert state['rules'] == mapping['field_rules']
        quality = next(c for c in state['cells'] if c['column'] == 'quality' and c['source_row'] == 5)
        assert quality['raw'] == '-1' and quality['status'] == 'missing' and quality['reason'] == 'unjudged'
        assert quality['source_line'] == 8
        height = [c for c in state['cells'] if c['column'] == 'height']
        assert {c['source_line'] for c in height} == set(range(3, 67))
        assert height[0]['status'] == ('valid' if n == 0 else 'invalid')
    assert outputs[0]['相关分析']['pairs'][0]['rho'] != outputs[1]['相关分析']['pairs'][0]['rho']
    assert client.get('/api/v1/applications/' + project['id'] + '/draft').json()['revision'] == revision
    assert path.read_bytes() == original

import io
import json

import httpx
import pytest

from agent_platform.project_knowledge import KnowledgeSearch, KnowledgeSource, parse_document
from tests.test_projects import configured, graph, node, edge, ref, start, settled  # noqa: F401


@pytest.fixture
def knowledge(configured, monkeypatch):
    client, app, project, settings = configured
    base = '/api/v1/projects/' + project['id']
    service = app.state.services.projects.knowledge
    calls = []

    async def embeddings(connection, texts):
        calls.extend(texts)
        return [[1., 0.] if any(word in text for word in ['cat', 'feline', '猫']) else [0., 1.] for text in texts]

    monkeypatch.setattr(service, 'embeddings', embeddings)
    return client, app, project, settings, base, service, calls


def setup_index(knowledge):
    client, app, project, settings, base, service, calls = knowledge
    assert client.put(base + '/embedding-model', json={'provider': 'api', 'model': 'embed-v1',
        'base_url': 'http://127.0.0.1:11434/v1', 'runtime_enabled': True}).status_code == 200
    response = client.put(base + '/knowledge/manual', json={'name': '手册', 'chunk_size': 100, 'chunk_overlap': 20})
    assert response.status_code == 200, response.text
    response = client.post(base + '/knowledge/manual/documents', json={'title': 'Pets', 'text': 'A cat enjoys sleeping.', 'expected_revision': 1})
    assert response.status_code == 200, response.text
    response = client.post(base + '/knowledge/manual/build', json={'expected_revision': 2})
    assert response.status_code == 200, response.text
    return response.json()


def test_save_without_resources_then_build_search_and_freeze_version(knowledge):
    client, app, project, settings, base, service, calls = knowledge
    graph(client, project['id'], [node('start', 'start'), node('search', 'knowledge_search', knowledge_ref='manual', query='feline'),
        node('end', 'end', outputs={'result': ref('search', 'output')})], [edge('start', 'search'), edge('search', 'end')])
    missing = settled(client, base, start(client, base, 'before-index'))
    assert missing['status'] == 'failed' and 'search' in missing['error'] and '索引' in missing['error']
    saved = setup_index(knowledge)
    run = settled(client, base, start(client, base, 'ready'))
    assert run['status'] == 'succeeded', run
    output = run['outputs']['result']
    assert output['version'] == saved['active_version'] and output['results'][0]['text'] == 'A cat enjoys sleeping.'
    assert output['results'][0]['citation'] == '[1]' and output['results'][0]['location']['line'] == 1
    assert 'vector' not in output['results'][0]
    count = len(calls)
    assert client.post(base + '/knowledge/manual/build', json={'expected_revision': 2}).json()['active_version'] == saved['active_version']
    assert len(calls) == count  # Ready builds do not pay for the same embeddings again.
    response = client.post(base + '/knowledge/manual/documents', json={'title': 'Cars', 'text': 'Cars need wheels.', 'expected_revision': 2})
    assert response.json()['status'] == 'pending'
    assert client.post(base + '/knowledge/manual/search', json={'query': 'wheels'}).status_code == 422
    async def old_version():
        return await service.search(project['id'], 'manual', KnowledgeSearch(query='feline'), version_id=saved['active_version'])
    assert client.portal.call(old_version)['results'][0]['title'] == 'Pets'
    rebuilt = client.post(base + '/knowledge/manual/build', json={'expected_revision': 3}).json()
    assert rebuilt['active_version'] != saved['active_version']
    assert client.get(base + '/tasks/' + run['id']).json()['outputs'] == run['outputs']
    other = client.post('/api/v1/projects', json={'name': '另一个项目'}).json()['id']
    assert client.get(f'/api/v1/projects/{other}/knowledge').json() == []
    assert client.get(f'/api/v1/projects/{other}/knowledge/manual').status_code == 404
    with pytest.raises(ValueError, match='当前项目'):
        client.portal.call(service.version, other, 'manual', saved['active_version'])


def test_index_commit_rejects_concurrent_edit_and_keeps_previous_version(knowledge, monkeypatch):
    client, app, project, settings, base, service, calls = knowledge
    saved = setup_index(knowledge)
    client.post(base + '/knowledge/manual/documents', json={'title': 'Next', 'text': 'cat next version', 'expected_revision': 2}).raise_for_status()
    async def edited(connection, texts):
        await service.add(project['id'], 'manual', KnowledgeSource(title='Concurrent', text='new material', expected_revision=3))
        return [[1., 0.] for _ in texts]
    monkeypatch.setattr(service, 'embeddings', edited)
    result = client.post(base + '/knowledge/manual/build', json={'expected_revision': 3})
    assert result.status_code == 409 and '期间资料已修改' in result.text
    current = client.get(base + '/knowledge/manual').json()
    assert current['active_version'] == saved['active_version'] and current['revision'] == 4
    assert not service.builds


def test_failed_build_and_changed_embedding_never_fall_back(knowledge, monkeypatch):
    client, app, project, settings, base, service, calls = knowledge
    saved = setup_index(knowledge)
    client.put(base + '/embedding-model', json={'provider': 'api', 'model': 'another', 'base_url': 'http://127.0.0.1:11434/v1', 'runtime_enabled': True}).raise_for_status()
    assert client.get(base + '/knowledge/manual').json()['status'] == 'pending'
    async def old_version():
        return await service.search(project['id'], 'manual', KnowledgeSearch(query='cat'), version_id=saved['active_version'])
    with pytest.raises(ValueError, match='连接已变化'):
        client.portal.call(old_version)
    async def fail(connection, texts):
        raise ValueError('embedding failed')
    monkeypatch.setattr(service, 'embeddings', fail)
    assert client.post(base + '/knowledge/manual/build', json={'expected_revision': 2}).status_code == 422
    assert client.get(base + '/knowledge/manual').json()['active_version'] == saved['active_version']
    assert not service.builds


def test_sources_preserve_page_text_and_reject_traversal_and_duplicate_content(knowledge):
    client, app, project, settings, base, service, calls = knowledge
    setup_index(knowledge)
    duplicate = client.post(base + '/knowledge/manual/documents', json={'title': 'Pets', 'text': 'A cat enjoys sleeping.', 'expected_revision': 2})
    assert duplicate.json()['revision'] == 2 and len(duplicate.json()['documents']) == 1
    root = app.state.services.projects.workspace(project['id'])
    (root / 'results').mkdir(exist_ok=True)
    (root / 'results' / 'document.txt').write_text('line one\n猫喜欢睡觉\nline three')
    added = client.post(base + '/knowledge/manual/documents', json={'source_path': 'results/document.txt', 'expected_revision': 2})
    assert added.status_code == 200, added.text
    for path in ['../secret', '/etc/passwd', 'results/../../secret']:
        assert client.post(base + '/knowledge/manual/documents', json={'source_path': path, 'expected_revision': 3}).status_code == 422
    (root / 'results' / 'symlink.txt').symlink_to(root / 'results' / 'document.txt')
    assert client.post(base + '/knowledge/manual/documents', json={'source_path': 'results/symlink.txt', 'expected_revision': 3}).status_code == 422
    doc_id = added.json()['documents'][-1]['id']
    removed = client.delete(base + f'/knowledge/manual/documents/{doc_id}?expected_revision=3')
    assert removed.status_code == 200 and len(removed.json()['documents']) == 1


def test_docx_and_pdf_parsing_keep_original_locations():
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    word = Document(); word.add_paragraph('猫的护理'); table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '饮水'; table.cell(0, 1).text = '每天更换'
    data = io.BytesIO(); word.save(data)
    sections = parse_document(data.getvalue(), 'care.docx')
    assert sections[0] == {'text': '猫的护理', 'paragraph': 1}
    assert '每天更换' in sections[1]['text']
    pdf = PdfWriter(); page = pdf.add_blank_page(width=300, height=300)
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 250 Td (Original page text) Tj ET')
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({
        NameObject('/F1'): DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})})})
    page[NameObject('/Contents')] = pdf._add_object(stream)
    data = io.BytesIO(); pdf.write(data)
    assert parse_document(data.getvalue(), 'care.pdf')[0] == {'text': 'Original page text', 'page': 1}


def test_remote_egress_and_response_validation(configured, monkeypatch):
    from agent_platform.model_connections import ModelConnection
    client, app, project, settings = configured
    service = app.state.services.projects.knowledge
    remote = ModelConnection(provider='api', model='embed', base_url='https://example.test/v1', api_key='private-key', runtime_enabled=True)
    with pytest.raises(ValueError, match='出口已关闭'):
        client.portal.call(service.embeddings, remote, ['private document'])
    local = remote.model_copy(update={'base_url': 'http://127.0.0.1:9001/v1'})
    original = httpx.AsyncClient
    bodies = []
    payload = {'data': [{'index': 1, 'embedding': [0., 3.]}, {'index': 0, 'embedding': [2., 0.]}]}
    def respond(request):
        bodies.append(json.loads(request.content))
        return httpx.Response(200, json=payload)
    monkeypatch.setattr(httpx, 'AsyncClient', lambda **kwargs: original(**kwargs, transport=httpx.MockTransport(respond)))
    assert client.portal.call(service.embeddings, local, ['first', 'second']) == [[1., 0.], [0., 1.]]
    assert bodies[0]['input'] == ['first', 'second']
    payload['data'][1]['index'] = 1
    with pytest.raises(ValueError, match='数量或序号'):
        client.portal.call(service.embeddings, local, ['first', 'second'])


def test_knowledge_permissions_and_embedding_secret(configured):
    from tests.test_users import signup, project as create_project
    client, app, _, settings = configured
    client.headers.pop('Authorization')
    _, owner = signup(client, 'knowledge-owner')
    member, collaborator = signup(client, 'knowledge-member')
    pid = create_project(client, owner)
    base = f'/api/v1/projects/{pid}'
    assert client.get(base + '/knowledge').status_code == 401
    assert client.get(base + '/knowledge', headers=collaborator).status_code == 404
    client.post(base + '/access-members', headers=owner, json={'name': 'knowledge-member'}).raise_for_status()
    connection = {'provider': 'api', 'model': 'embedding', 'base_url': 'https://example.test/v1', 'api_key': 'private-embedding-key', 'runtime_enabled': True}
    assert client.put(base + '/embedding-model', headers=collaborator, json=connection).status_code == 403
    saved = client.put(base + '/embedding-model', headers=owner, json=connection)
    assert saved.status_code == 200 and 'private-embedding-key' not in saved.text
    public = client.get(base + '/embedding-model', headers=collaborator)
    assert public.status_code == 200 and 'api_key' not in public.json() and public.json()['has_api_key']
    assert client.put(base + '/knowledge/team', headers=collaborator, json={'name': '团队资料'}).status_code == 200
    assert client.get(base + '/knowledge', headers=owner).json()[0]['knowledge_ref'] == 'team'
    client.delete(base + '/access-members/' + member['user']['id'], headers=owner).raise_for_status()
    assert client.get(base + '/knowledge/team', headers=collaborator).status_code == 404
    assert client.post(base + '/knowledge/team/search', headers=collaborator, json={'query': 'private'}).status_code == 404


@pytest.mark.parametrize('filename', ['broken.pdf', 'broken.docx'])
def test_malformed_document_returns_fixable_error_without_partial_save(knowledge, filename):
    client, app, project, settings, base, service, calls = knowledge
    setup_index(knowledge)
    root = app.state.services.projects.workspace(project['id']) / 'results'
    root.mkdir(exist_ok=True)
    (root / filename).write_bytes(b'not a valid document')
    result = client.post(base + '/knowledge/manual/documents', json={'source_path': 'results/' + filename, 'expected_revision': 2})
    assert result.status_code == 422 and '重新导出' in result.text
    current = client.get(base + '/knowledge/manual').json()
    assert current['revision'] == 2 and len(current['documents']) == 1 and current['status'] == 'ready'

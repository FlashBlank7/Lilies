"""Portable document preparation, excerpt checks, and report-only replay."""
import csv
import hashlib
import io
import json
import re
from collections import Counter
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile
import xml.etree.ElementTree as ET

STATUSES=['说法一致','存在分歧','信息互补','口径不同','与要求一致','与要求不一致','资料不足']
ROLES=['说明','要求/标准','实际记录','其他']


def source(value):
    p=Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ['requirement-package','results']:
        raise ValueError('请选择当前项目资料或结果文件')
    if any(v.is_symlink() for v in [p,*p.parents]) or not p.resolve().is_relative_to(Path.cwd().resolve()):raise ValueError('不能读取项目外文件或符号链接')
    if not p.is_file():raise ValueError('文件不存在：'+str(p))
    if p.stat().st_size>10_000_000:raise ValueError('文件超过10 MB，请拆分所需资料')
    return p


def digest(raw):return hashlib.sha256(raw).hexdigest()


def bounded(value,label,limit=2000):
    if not isinstance(value,str) or len(value)>limit:raise ValueError(f'{label}需要不超过{limit}字的文字')
    return value.strip()


def read_document(path):
    raw=path.read_bytes();suffix=path.suffix.lower();parts=[];notes=[]
    if suffix in ['.txt','.md']:
        parts=[(f'第{i}行',v) for i,v in enumerate(raw.decode('utf-8-sig').splitlines(),1)]
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:raise ValueError('请提供未加密PDF')
        if len(reader.pages)>60:raise ValueError('单份PDF超过60页，请选择所需章节另存后处理')
        for i,page in enumerate(reader.pages,1):
            value=page.extract_text() or '';parts.append((f'第{i}页',value))
            if not value.strip():notes.append(f'第{i}页没有可提取正文，未执行OCR')
        notes.append('PDF仅提取文字；图形、图内分支及版面关系未识别')
    elif suffix=='.docx':
        with ZipFile(io.BytesIO(raw)) as archive:
            if archive.getinfo('word/document.xml').file_size>10_000_000:raise ValueError('DOCX正文过大，请拆分')
            root=ET.fromstring(archive.read('word/document.xml'))
        ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        for i,p in enumerate(root.findall('.//{'+ns+'}p'),1):
            text=''.join(e.text or '' if e.tag=='{'+ns+'}t' else '\t' if e.tag=='{'+ns+'}tab' else '\n' if e.tag in ['{'+ns+'}br','{'+ns+'}cr'] else '' for e in p.iter())
            parts.append((f'第{i}段（含表格内段落）',text))
        if root.findall('.//{'+ns+'}drawing') or root.findall('.//{'+ns+'}pict'):notes.append('DOCX中的图片未识别')
        notes.append('DOCX按正文段落定位，不推断页码；页眉、批注和附件未读取')
    elif suffix in ['.csv','.tsv']:
        values=list(csv.reader(io.StringIO(raw.decode('utf-8-sig')),delimiter='\t' if suffix=='.tsv' else ','))
        if len(values)>10000:raise ValueError('表格超过1万行，请选择所需部分')
        for i,row in enumerate(values,1):parts.append((f'第{i}条记录',' | '.join(f'列{j}={v}' for j,v in enumerate(row,1))))
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            count=0
            for sheet in book:
                for i,row in enumerate(sheet.values,1):
                    count+=1
                    if count>10000:raise ValueError('工作簿超过1万行，请选择所需部分')
                    text=' | '.join(f'{get_column_letter(j)}{i}={v}' for j,v in enumerate(row,1) if v is not None)
                    parts.append((f'{sheet.title} · 第{i}行',text))
        finally:book.close()
        notes.append('工作簿仅使用单元格值和已有公式缓存，不重新计算公式或识别图片')
    else:raise ValueError('支持TXT、Markdown、文字PDF、DOCX、CSV、TSV或XLSX')
    parts=[(where,text) for where,text in parts if text.strip()]
    if not parts:raise ValueError('没有提取到正文；请提供文字版，扫描件需先转为文字')
    if sum(len(t) for _,t in parts)>30000:raise ValueError('单份正文超过3万字符，请拆分；没有截断后假装读完')
    return raw,parts,notes


def folder():
    p=Path('results')/('source-comparison-'+uuid4().hex);p.mkdir(parents=True);return p


def prepare(inputs):
    files=inputs.get('sources')
    if not isinstance(files,list) or not 1<=len(files)<=8 or any(not isinstance(v,dict) for v in files):raise ValueError('请添加1至8份需要比较的材料')
    question=bounded(inputs.get('question',''),'关注事项')
    sources=[];chunks=[];seen=set();total=0
    for i,item in enumerate(files,1):
        path=source(item.get('path',''));role=item.get('role') or '说明';context=bounded(item.get('context',''),'来源说明',500)
        if role not in ROLES:raise ValueError('请选择有效的材料用途')
        if str(path) in seen:raise ValueError('同一文件被重复选择，请删除重复行')
        seen.add(str(path));raw,parts,notes=read_document(path);sid='S'+str(i)
        for where,text in parts:
            for offset in range(0,len(text),1200):
                part=text[offset:offset+1200];total+=len(part)
                if total>30000:raise ValueError('合计正文超过3万字符，请选择相关材料或拆分；尚未调用模型')
                if len(chunks)>=400:raise ValueError('正文片段超过400处，请选择相关章节或表格部分')
                chunks.append({'id':sid+'.'+str(len(chunks)+1),'source_id':sid,'location':where+f' · 字符{offset+1}至{offset+len(part)}','text':part})
        sources.append({'id':sid,'path':str(path),'name':path.name,'sha256':digest(raw),'role':role,'context':context,
            'characters':sum(len(t) for _,t in parts),'notes':notes})
    for s in sources:
        same=[v['id'] for v in sources if v['sha256']==s['sha256'] and v['id']!=s['id']]
        if same:s['notes'].append('与'+','.join(same)+'字节相同，不是新增内容；不同文件本身也不证明来源独立')
    data={'version':1,'question':question,'sources':sources,'chunks':chunks,'characters':total}
    root=folder();path=root/'source-map.json';path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    prompt=json.dumps({'question':question,'sources':sources,'chunks':chunks},ensure_ascii=False)
    return {'snapshot_path':str(path),'sha256':digest(path.read_bytes()),'prompt':prompt,'characters':total,'source_count':len(sources)}


def snapshot(prepared):
    raw=source(prepared['snapshot_path']).read_bytes()
    if digest(raw)!=prepared['sha256']:raise ValueError('本次原文快照已被修改，请重新分析资料')
    data=json.loads(raw)
    if data.get('version')!=1:raise ValueError('不支持的原文快照版本')
    return data


def normalize(text):return re.sub(r'\s+','',text)


def check(answer,data):
    if not isinstance(answer,dict) or not isinstance(answer.get('items'),list) or len(answer['items'])>12:raise ValueError('模型输出缺少有效事项列表，未形成可用分析')
    sources={s['id']:s for s in data['sources']};chunks={c['id']:c for c in data['chunks']};items=[]
    for index,item in enumerate(answer['items'],1):
        if not isinstance(item,dict) or item.get('status') not in STATUSES:raise ValueError('比较事项的判断类型无效')
        topic=bounded(item.get('topic'),'事项',300);reason=bounded(item.get('reason'),'分析理由');next_step=bounded(item.get('next_step'),'下一步建议')
        raw_statements=item.get('statements')
        if not isinstance(raw_statements,list) or len(raw_statements)>8:raise ValueError('比较事项需要有效的来源陈述列表')
        statements=[];issues=[];seen=set()
        for statement in raw_statements:
            if not isinstance(statement,dict):raise ValueError('来源陈述格式不正确')
            sid=statement.get('source_id');text=bounded(statement.get('statement'),'来源陈述');scope=bounded(statement.get('scope'),'适用范围',1000)
            evidence=[];problems=[];quotes=statement.get('quotes')
            if sid not in sources:problems.append('引用了本次未提供的来源')
            if sid in seen:problems.append('同一事项重复列出该来源，请合并核对')
            seen.add(sid)
            if not isinstance(quotes,list) or len(quotes)>2:raise ValueError('每条陈述最多两处引用')
            for q in quotes:
                if not isinstance(q,dict):raise ValueError('引用格式不正确')
                cid=q.get('chunk_id');quote=bounded(q.get('quote'),'引用原文',600);chunk=chunks.get(cid)
                if not chunk or chunk['source_id']!=sid:problems.append('片段不属于所声明的来源')
                elif len(normalize(quote))<4 or normalize(quote) not in normalize(chunk['text']):problems.append('所引文字不能在指定原文片段定位')
                else:evidence.append({'chunk_id':cid,'quote':quote,'location':chunk['location'],'context':chunk['text']})
            if not evidence:problems.append('没有可定位的引用原文')
            statements.append({'source_id':sid,'statement':text,'scope':scope,'evidence':evidence,'issues':problems})
            issues.extend(f'{sid}：{p}' for p in problems)
        valid=[s for s in statements if not s['issues']];distinct={sources[s['source_id']]['sha256'] for s in valid};status=item['status']
        if status in ['说法一致','存在分歧','信息互补','口径不同','与要求一致','与要求不一致'] and len(distinct)<2:
            issues.append('有效引用不足两份不同内容的材料，不能据此作跨来源判断')
        if status in ['与要求一致','与要求不一致'] and (not any(sources[s['source_id']]['role']=='要求/标准' for s in valid)
                or not any(sources[s['source_id']]['role']!='要求/标准' for s in valid)):
            issues.append('要求对照缺少员工指定的要求材料或实际对应材料')
        if not statements:issues.append('模型没有提供来源陈述，仅保留待补充事项')
        items.append({'id':'T'+str(index),'topic':topic,'model_status':status,'status':'待核对' if issues else status,
            'reason':reason,'statements':statements,'missing_sources':[s for s in sources if s not in seen],
            'issues':issues,'next_step':next_step})
    return items


def analyze(inputs):
    prepared=inputs['prepared'];data=snapshot(prepared);answer=inputs['answer'];check(answer,data)
    path=folder()/'analysis.json'
    path.write_text(json.dumps({'format':'source-comparison-v1','prepared':{'snapshot_path':prepared['snapshot_path'],'sha256':prepared['sha256']},
        'answer':answer,'model_usage':inputs.get('model_usage',{})},ensure_ascii=False,indent=2),encoding='utf-8')
    return {'analysis_path':str(path)}


def escape(value):
    return re.sub(r'([\\`*_{}\[\]<>()!|#])',r'\\\1',str(value)).replace('\r',' ').replace('\n',' / ')


def write_csv(path,fields,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader()
        writer.writerows({k:"'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for k,v in r.items()} for r in rows)


def report(inputs):
    path=source(inputs['source_result_path']);raw=path.read_bytes();saved=json.loads(raw)
    if saved.get('format')!='source-comparison-v1':raise ValueError('请选择本流程生成的analysis.json')
    data=snapshot(saved['prepared']);all_items=check(saved['answer'],data)
    status=inputs.get('status_filter') or '全部';keyword=bounded(inputs.get('keyword',''),'筛选关键词',200)
    if status not in ['全部','待核对',*STATUSES]:raise ValueError('请选择有效的报告状态')
    items=[i for i in all_items if (status=='全部' or i['status']==status) and (not keyword or keyword.casefold() in json.dumps(i,ensure_ascii=False).casefold())]
    sources={s['id']:s for s in data['sources']};root=folder()
    text=f'# 多份材料的事实与分歧整理\n\n关注事项：{escape(data["question"])}\n\n本次读取 {len(sources)} 份材料、{data["characters"]} 个正文字符；展示 {len(items)}/{len(all_items)} 个模型整理事项。\n\n'
    text+='以下“来源陈述”来自材料，“比较说明”是模型解释。代码只核对引用文字及位置，不证明陈述真实或解释正确。没有提取到某项内容，不等于材料中不存在或功能未实现。\n\n'
    if status!='全部' or keyword:text+=f'本报告筛选：{escape(status)}；关键词：{escape(keyword or "无")}。完整分析仍保留，可清空筛选重新生成。\n\n'
    text+='## 材料与解析范围\n\n'
    for s in sources.values():
        text+=f'- {s["id"]} · {escape(s["name"])} · 用途：{escape(s["role"])}；员工补充：{escape(s["context"] or "未填写")}。\n'
        for note in s['notes']:text+=f'  - {escape(note)}。\n'
    text+='\n## 按事项比较\n\n| 事项 | 当前判断 | '+ ' | '.join(escape(s['id']+' '+s['name']) for s in sources.values())+' |\n'
    text+='| --- | --- | '+' | '.join('---' for _ in sources)+' |\n'
    csv_rows=[]
    for item in items:
        row={'事项':item['topic'],'当前判断':item['status'],'模型原判断':item['model_status'],'比较说明':item['reason'],'下一步':item['next_step'],'引用问题':'；'.join(item['issues'])}
        cells=[]
        for sid,s in sources.items():
            matches=[v for v in item['statements'] if v['source_id']==sid]
            cell=' / '.join(v['statement']+('（引用待核对）' if v['issues'] else '') for v in matches) or '本次未提取该项陈述'
            row[sid+' '+s['name']]=cell;cells.append(escape(cell[:180]))
        csv_rows.append(row);text+='| '+ ' | '.join([escape(item['topic']),escape(item['status']),*cells])+' |\n'
    if not items:text+='\n没有符合当前筛选的事项。原始分析与材料仍可查看；没有据此判定资料没有问题。\n'
    text+='\n## 逐项理由与原文\n\n'
    for item in items:
        text+=f'### {item["id"]} · {escape(item["topic"])}\n\n当前判断：{escape(item["status"])}。模型原判断：{escape(item["model_status"])}。\n\n比较说明（模型）：{escape(item["reason"])}\n\n'
        for s in item['statements']:
            text+=f'- {escape(s["source_id"])} 来源陈述：{escape(s["statement"])}；范围：{escape(s["scope"] or "未明确")}。\n'
            for q in s['evidence']:text+=f'  - [{escape(q["chunk_id"])}] {escape(q["location"])}：{escape(q["quote"])}\n'
        if item['missing_sources']:text+='\n未提取到该项陈述的材料：'+', '.join(item['missing_sources'])+'；不等同于文中不存在，请按需核对原文。\n'
        for issue in item['issues']:text+='\n引用或比较范围待核对：'+escape(issue)+'。\n'
        text+='\n下一步建议：'+escape(item['next_step'] or '结合原文核对，不自动执行后续操作')+'\n\n'
    source_text='# 本次提取的正文与定位\n\n原文片段与文件内容哈希见source-map.json；文中提到的其他资料没有自动读取。\n\n'
    for s in sources.values():
        source_text+=f'## {s["id"]} · {escape(s["name"])}\n\n文件：{escape(s["path"])}\n\n内容SHA256：{s["sha256"]}\n\n'
        for c in data['chunks']:
            if c['source_id']==s['id']:source_text+=f'### {c["id"]} · {escape(c["location"])}\n\n{escape(c["text"])}\n\n'
    result={'rows':len(items),'total_items':len(all_items),'status_counts':dict(Counter(i['status'] for i in all_items)),
        'items':items,'sources':data['sources'],'source_characters':data['characters'],'analysis_path':str(path),'analysis_sha256':digest(raw),
        'model_usage':saved.get('model_usage',{}),'status_filter':status,'keyword':keyword,'source_path':str(root/'comparison.csv')}
    fields=['事项','当前判断','模型原判断','比较说明','下一步','引用问题',*[s['id']+' '+s['name'] for s in sources.values()]]
    write_csv(root/'comparison.csv',fields,csv_rows)
    (root/'report.md').write_text(text,encoding='utf-8');(root/'source-texts.md').write_text(source_text,encoding='utf-8')
    (root/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    artifacts=[{'file_path':str(root/n),'label':label} for n,label in [('report.md','比较报告 Markdown'),('comparison.csv','事项与来源对照 CSV'),('result.json','本次筛选结果 JSON'),('source-texts.md','本次原文与定位 Markdown')]]
    artifacts += [{'file_path':str(path),'label':'可再次筛选的分析 JSON'},{'file_path':saved['prepared']['snapshot_path'],'label':'原文、来源与版本 JSON'}]
    return {**result,'markdown':text,'artifacts':artifacts}


def main(inputs):return {'prepare':prepare,'analyze':analyze,'report':report}[inputs['operation']](inputs)

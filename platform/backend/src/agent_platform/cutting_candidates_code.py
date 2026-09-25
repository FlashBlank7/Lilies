"""Bounded single-stock patterns with explicit fixed or ranged lengths."""
import csv
import hashlib
import json
from decimal import Decimal, InvalidOperation, ROUND_CEILING, ROUND_FLOOR
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value):
    p=Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ('requirement-package','results'):
        raise ValueError('请选择当前项目的资料或结果文件')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p,*p.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size>10_000_000:raise ValueError('文件不存在或超过10 MB，请检查资料')
    return p


def table(value,sheet,required):
    p=source(value)
    if p.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(p,read_only=True,data_only=True)
        try:
            if not sheet and len(book.sheetnames)>1:raise ValueError('Excel有多张表，请指定工作表：'+'、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:raise ValueError('工作表不存在：'+sheet)
            rows=list(islice((book[sheet] if sheet else book.active).values,1002))
        finally:book.close()
    elif p.suffix.lower() in ('.csv','.tsv'):
        with p.open(encoding='utf-8-sig',newline='') as f:rows=list(islice(csv.reader(f,delimiter='\t' if p.suffix.lower()=='.tsv' else ','),1002))
    else:raise ValueError('请选择CSV、TSV或XLSX')
    if not rows or len(rows)>1001:raise ValueError('需要表头和数据，单表最多1000条，请选择本次物料与需求')
    fields=['' if v is None else str(v).strip() for v in rows[0]]
    if any(not f for f in fields) or len(fields)!=len(set(fields)):raise ValueError('表头不能为空或重名')
    if set(required)-set(fields):raise ValueError(p.name+' 缺少字段：'+'、'.join(sorted(set(required)-set(fields))))
    output=[]
    for rownum,row in enumerate(rows[1:],2):
        if all(v is None or not str(v).strip() for v in row):continue
        if len(row)!=len(fields):raise ValueError(f'{p.name} 第{rownum}行列数与表头不同')
        output.append((rownum,dict(zip(fields,['' if v is None else str(v).strip() for v in row]))))
    if not output:raise ValueError(p.name+' 没有数据记录')
    return p,output


def number(value,label,positive=False):
    try:
        n=Decimal(str(value).strip())
        if n.is_finite() and 0<=n<=Decimal('1000000000') and n.as_tuple().exponent>=-6 and (not positive or n>0):return n
    except (InvalidOperation,ValueError):pass
    raise ValueError(label+' 需要'+('正数' if positive else '非负数')+'，最多6位小数且不超过10亿')


def integer(value,label,lo,hi):
    n=number(value,label)
    if n!=int(n) or not lo<=n<=hi:raise ValueError(f'{label} 需要{lo}至{hi}的整数')
    return int(n)


def prepare(inputs):
    unit=str(inputs.get('unit') or '').strip()
    if not unit:raise ValueError('请说明两张表共同的长度单位，不会自动换算')
    mode=inputs.get('kerf_mode')
    if mode not in ('每段各计一道切缝','仅相邻段间计切缝'):raise ValueError('请选择实际切缝计数方式')
    config={'unit':unit,'kerf':str(number(inputs.get('kerf'),'单次切缝')),
            'end_allowance':str(number(inputs.get('end_allowance',0),'总预留长度')),'kerf_mode':mode,
            'max_pieces':integer(inputs.get('max_pieces',6),'最多段数',1,12),
            'max_types':integer(inputs.get('max_types',3),'最多需求种类',1,6),
            'search_limit':integer(inputs.get('search_limit',20000),'组合分支上限',100,200000)}
    length_mode=inputs.get('length_mode') or '固定长度'
    if length_mode not in ('固定长度','长度范围'):raise ValueError('请选择需求长度方式')
    config['length_mode']=length_mode
    if length_mode=='长度范围':
        allocation=inputs.get('allocation_mode') or '按区间比例分配'
        if allocation not in ('按区间比例分配','先中点再按需求表顺序分配'):raise ValueError('请选择范围内长度分配方式')
        config.update(allocation_mode=allocation,length_precision=integer(inputs.get('length_precision',6),'长度小数位',0,6))
    quantum=Decimal(1).scaleb(-config.get('length_precision',6))
    data={'config':config,'sources':{}}
    for kind,ident,limit in [('stock','stock_id',30),('demand','demand_id',12)]:
        ranged=kind=='demand' and length_mode=='长度范围'
        length_fields=['min_length','max_length'] if ranged else ['length']
        path,rows=table(inputs.get(kind+'_path'),str(inputs.get(kind+'_sheet') or ''),[ident,'material']+length_fields+(['quantity'] if kind=='demand' else []))
        if len(rows)>limit:raise ValueError(f'本次最多{limit}条'+('物料' if kind=='stock' else '需求')+'，请拆分本次输入')
        seen=set();normalized=[]
        for rownum,row in rows:
            label=f'{path.name} 第{rownum}行'
            if not row[ident] or row[ident] in seen:raise ValueError(label+' 标识为空或重复：'+row[ident])
            seen.add(row[ident])
            if not row['material']:raise ValueError(label+' 缺少明确物料类型，不能猜测是否兼容')
            if 'unit' in row and row['unit']!=unit:raise ValueError(label+' 长度单位与表单不同或缺失，请先统一单位')
            item={ident:row[ident],'material':row['material'],'source_row':rownum}
            if ranged:
                lower=number(row['min_length'],label+' 长度下限',True);upper=number(row['max_length'],label+' 长度上限',True)
                if lower>upper:raise ValueError(label+' 长度下限不能大于上限')
                effective_lower=lower.quantize(quantum,rounding=ROUND_CEILING)
                effective_upper=upper.quantize(quantum,rounding=ROUND_FLOOR)
                if effective_lower>effective_upper:raise ValueError(label+' 长度范围内没有满足所选小数位的长度，请增加小数位或核对范围')
                item.update(length=str(effective_lower),upper_length=str(effective_upper),min_length=str(lower),max_length=str(upper))
            else:item['length']=str(number(row['length'],label+' 长度',True))
            if kind=='demand':item['quantity']=integer(row['quantity'],label+' 剩余需求',0,1000000)
            normalized.append(item)
        data[kind]=normalized
        data['sources'][kind]={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':inputs.get(kind+'_sheet') or ''}
    folder=Path('results')/('cutting-input-'+uuid4().hex);folder.mkdir(parents=True)
    snapshot=folder/'input.json';snapshot.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    return {'snapshot_path':str(snapshot),'sha256':hashlib.sha256(snapshot.read_bytes()).hexdigest()}


def text_number(n):return format(n,'f')


def allocate(selected,available,config):
    """One declared policy per count pattern, not all continuous assignments."""
    lower=[Decimal(d['length']) for d,n in selected]
    if config.get('length_mode','固定长度')=='固定长度':return lower
    upper=[Decimal(d['upper_length']) for d,n in selected]
    counts=[n for d,n in selected]
    minimum=sum(n*v for n,v in zip(counts,lower));maximum=sum(n*v for n,v in zip(counts,upper))
    if available<minimum:raise ValueError('组合可用长度不足，请重新检查输入和损耗')
    middle=[(a+b)/2 for a,b in zip(lower,upper)]
    middle_total=sum(n*v for n,v in zip(counts,middle))
    if available>=maximum:values=upper
    elif config['allocation_mode']=='先中点再按需求表顺序分配' and available>=middle_total:
        values=middle;extra=available-middle_total
        for i,n in enumerate(counts):
            addition=min(upper[i]-values[i],extra/n)
            values[i]+=addition;extra-=addition*n
    else:
        ratio=(available-minimum)/(maximum-minimum)
        values=[a+ratio*(b-a) for a,b in zip(lower,upper)]
    quantum=Decimal(1).scaleb(-config['length_precision'])
    # Round only down within the already representable bounds. Any rounding
    # remainder stays explicit; no claim to maximize usage across assignments.
    return [v.quantize(quantum,rounding=ROUND_FLOOR) for v in values]


def generate(inputs):
    prepared=inputs['prepared'];path=source(prepared['snapshot_path'])
    if hashlib.sha256(path.read_bytes()).hexdigest()!=prepared['sha256']:raise ValueError('输入快照已改变，请重新读取资料')
    data=json.loads(path.read_text(encoding='utf-8'));cfg=data['config']
    kerf=Decimal(cfg['kerf']);allowance=Decimal(cfg['end_allowance'])
    candidates=[];patterns=[];summary=[];states=0
    def cuts(n):return n if cfg['kerf_mode']=='每段各计一道切缝' else max(0,n-1)
    for stock in data['stock']:
        available=Decimal(stock['length'])-allowance
        demand=sorted([d for d in data['demand'] if d['material']==stock['material'] and d['quantity']>0],
                      key=lambda x:x['source_row'] if cfg.get('length_mode')=='长度范围' else x['demand_id'])
        before=len(candidates)
        def walk(index,counts,count,length,types):
            nonlocal states
            states+=1
            if states>cfg['search_limit']:raise ValueError('达到组合分支上限，尚未完整生成；请减少本次物料、段数或需求种类，或在表单提高上限后重跑。没有发布截断候选表。')
            if index==len(demand):
                if not count:return
                if len(candidates)>=10000:raise ValueError('候选超过1万，请减少本次物料、段数或需求种类后重跑；未发布截断结果')
                selected=[(d,n) for d,n in zip(demand,counts) if n]
                fingerprint=json.dumps([(d['demand_id'],n) for d,n in selected],ensure_ascii=False)
                cid=stock['stock_id']+'-'+hashlib.sha256(fingerprint.encode()).hexdigest()[:12]
                loss=cuts(count)*kerf
                assigned=allocate(selected,available-loss,cfg)
                length=sum(n*v for (d,n),v in zip(selected,assigned));used=length+loss
                candidates.append(dict(candidate_id=cid,stock_id=stock['stock_id'],material=stock['material'],unit=cfg['unit'],
                    stock_length=stock['length'],available_length=text_number(available),piece_count=count,demand_types=types,
                    product_length=text_number(length),cuts=cuts(count),kerf_loss=text_number(loss),used_length=text_number(used),
                    remaining_length=text_number(available-used),end_allowance=cfg['end_allowance'],
                    combination='；'.join(d['demand_id']+' × '+str(n) for d,n in selected),
                    allocation_method=cfg.get('allocation_mode','固定长度')))
                for (d,n),v in zip(selected,assigned):patterns.append(dict(candidate_id=cid,stock_id=stock['stock_id'],demand_id=d['demand_id'],quantity=n,
                    length=text_number(v),min_length=d.get('min_length',d['length']),max_length=d.get('max_length',d['length']),unit=cfg['unit']))
                return
            d=demand[index];segment=Decimal(d['length'])
            maximum=min(d['quantity'],cfg['max_pieces']-count,int(available//segment))
            for n in range(maximum+1):
                if types+bool(n)>cfg['max_types']:break
                next_length=length+n*segment
                if next_length+cuts(count+n)*kerf>available:break
                walk(index+1,counts+[n],count+n,next_length,types+bool(n))
        if available>0 and demand:walk(0,[],0,Decimal(0),0)
        generated=len(candidates)-before
        reason='' if generated else '预留后没有可用长度' if available<=0 else '没有同类型且数量大于0的需求' if not demand else '长度与切缝条件下没有可产出的需求段'
        summary.append(dict(stock_id=stock['stock_id'],material=stock['material'],candidates=generated,reason=reason))
    folder=Path('results')/('cutting-candidates-'+uuid4().hex);folder.mkdir(parents=True);artifacts=[]
    def save_csv(name,fields,rows,label):
        with (folder/name).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
        artifacts.append(dict(file_path=str(folder/name),label=label))
    save_csv('candidates.csv',['candidate_id','stock_id','material','unit','stock_length','available_length','piece_count','demand_types','product_length','cuts','kerf_loss','used_length','remaining_length','end_allowance','combination','allocation_method'],candidates,'单料候选 CSV')
    save_csv('patterns.csv',['candidate_id','stock_id','demand_id','quantity','length','min_length','max_length','unit'],patterns,'需求数量与分配长度 CSV')
    save_csv('stocks.csv',['stock_id','material','candidates','reason'],summary,'每根物料结果与缺项 CSV')
    result=dict(rows=len(candidates),stocks=summary,search_states=states,complete_within_declared_limits=True,config=cfg,sources=data['sources'],
                source_path=str(folder/'candidates.csv'),patterns_path=str(folder/'patterns.csv'))
    result['comparison_inputs']={'source_path':result['source_path'],'id_column':'candidate_id','group_column':'stock_id','sheet':'',
        'constraints':[dict(field='used_length',operator='不大于',other_field='available_length',value='')],
        'mode':'按优先级逐项比较','objectives':[]} if candidates else None
    def md(v):return str(v).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('|','\\|').replace('\n',' ').replace('\r',' ')
    lines=['# 按物料与需求生成下料候选','',f"共{len(data['stock'])}根物料、{len(candidates)}个单料组合，检查{states}个分支。只覆盖本次明确的段数、种类和需求长度条件。",
        '',f"长度单位：{md(cfg['unit'])}；单次切缝：{cfg['kerf']}；每根共预留：{cfg['end_allowance']}；计数：{cfg['kerf_mode']}。",
        '', '| 物料 | 类型 | 候选数 | 无方案原因 |','|---|---|---:|---|']
    lines.extend('| '+' | '.join(md(s[k]) for k in ['stock_id','material','candidates','reason'])+' |' for s in summary)
    if cfg.get('length_mode')=='长度范围':
        lines+=['',f"本次按长度范围制备：{cfg['allocation_mode']}，最多{cfg['length_precision']}位小数。各段先满足范围下限，再按所选方式分配可用长度；同一需求在同一候选中的每段等长。",
            '需求明细保留原上下限及实际分配长度；上下限先收紧到所选小数位，分配结果向下取整，产生的余量保留。区间为连续可选长度，若有离散允许尺寸需另行建模。',
            '每个数量组合只给出一种按所选方式分配的长度，不枚举区间内所有尺寸，也不证明所给长度全局最优。先中点方式在可用量不足中点时按范围比例分配；达到中点后按需求表行顺序分配，改变行顺序可能改变结果。']
    lines+=['','## 查看组合（最多展示前30个，完整记录见下载）','','| 物料 | 需求组合 | 段数 | 产品长度 | 切缝损耗 | 剩余长度 |','|---|---|---:|---:|---:|---:|']
    lines.extend('| '+' | '.join(md(r[k]) for k in ['stock_id','combination','piece_count','product_length','kerf_loss','remaining_length'])+' |' for r in candidates[:30])
    lines+=['','## 如何使用这些候选','',
        '数量明细表示组合，不规定实际切割顺序。余量不等于废料，可复用性、设备、缺陷分区及短料配对还需具体资料。',
        '可以将候选表交给已有方案比较流程，按stock_id分别比较；先说明业务偏好，再选择指标。仅改变比较偏好可以直接复用本次候选。' if candidates else '没有候选时先根据上表核对需求类型、长度、损耗及预留；不自动放宽条件或调用空表比较。',
        '各根独立使用同一份剩余需求，不能把首选直接合并为整体排程；共享库存、总需求、设备和时序需要后续分配检查。本次没有训练、消耗库存或回写现场。']
    result['markdown']='\n'.join(lines)
    (folder/'report.md').write_text(result['markdown'],encoding='utf-8')
    artifacts.append(dict(file_path=str(folder/'report.md'),label='候选说明 Markdown'))
    (folder/'summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    artifacts.extend([dict(file_path=str(folder/'summary.json'),label='本次配置与汇总 JSON'),dict(file_path=str(path),label='物料与需求快照 JSON')])
    result['artifacts']=artifacts
    return result


def main(inputs):
    if inputs.get('operation')=='prepare':return prepare(inputs)
    if inputs.get('operation')=='generate':return generate(inputs)
    raise ValueError('未知处理步骤')

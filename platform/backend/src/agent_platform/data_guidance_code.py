"""Portable, editable code for the data guidance workflow; no provider calls here."""
import csv
import hashlib
import json
import math
import statistics
from collections import Counter
from pathlib import Path
from uuid import uuid4


def source(value):
    path = Path(str(value))
    if path.is_absolute() or '..' in path.parts or not path.parts or path.parts[0] not in ('requirement-package', 'results'):
        raise ValueError('请选择当前项目内的 CSV、TSV 或 Excel 文件')
    if not path.resolve().is_relative_to(Path.cwd().resolve()) or any(p.is_symlink() for p in [path, *path.parents]):
        raise ValueError('不能读取项目外文件或符号链接')
    if not path.is_file():
        raise ValueError('资料文件不存在，请重新选择')
    if path.stat().st_size > 10_000_000:
        raise ValueError('首版分析支持 10 MB 以内的表格，请按业务范围拆分后处理')
    return path


def read_table(path, sheet=''):
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if not sheet and len(book.sheetnames) > 1:
                raise ValueError('此文件有多张工作表，请填写要分析的工作表名称：' + '、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:
                raise ValueError('工作表不存在：' + sheet)
            data = book[sheet] if sheet else book.active
            rows = []
            for row in data.values:
                rows.append(['' if v is None else str(v) for v in row])
                if len(rows) > 100001:
                    raise ValueError('首版最多分析 10 万行，请拆分数据')
        finally:
            book.close()
    elif path.suffix.lower() in ('.csv', '.tsv'):
        with path.open(encoding='utf-8-sig', newline='') as stream:
            rows = []
            for row in csv.reader(stream, delimiter='\t' if path.suffix.lower() == '.tsv' else ','):
                rows.append(row)
                if len(rows) > 100001:
                    raise ValueError('首版最多分析 10 万行，请拆分数据')
    else:
        raise ValueError('请选择 CSV、TSV 或 XLSX；旧版 XLS 请先另存为 XLSX')
    fields = [v.strip() for v in rows[0]] if rows else []
    if not fields or any(not f for f in fields) or len(fields) != len(set(fields)):
        raise ValueError('第一行需要非空、不重名的字段名，请检查表头')
    data = [row for row in rows[1:] if any(v.strip() for v in row)]
    if not data:
        raise ValueError('资料只有表头或空行，没有可分析的数据')
    if any(len(row) != len(fields) for row in data):
        raise ValueError('存在与表头列数不一致的记录，请检查分隔符和单元格')
    return fields, data, len(rows) - 1 - len(data)


def write_csv(path, fields, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream); writer.writerow(fields); writer.writerows(rows)


def profile(inputs):
    path = source(inputs['source_path'])
    fields, rows, empty = read_table(path, inputs.get('sheet', ''))
    folder = Path('results/data-guidance') / str(uuid4()); folder.mkdir(parents=True)
    repeated = Counter(tuple(row) for row in rows)
    columns = []
    for index, name in enumerate(fields):
        values = [r[index].strip() for r in rows if r[index].strip()]
        counts = Counter(values); numbers = []
        for value in values:
            try:
                number = float(value)
                if math.isfinite(number): numbers.append(number)
            except ValueError:
                pass
        numeric = bool(values) and len(numbers) == len(values)
        item = {'name': name, 'type': '数值' if numeric else '文本或混合', 'missing': len(rows)-len(values),
                'unique': len(counts), 'constant': len(counts) == 1,
                'repeated_values': sum(c-1 for c in counts.values()),
                'smallest_value_count': min(counts.values()) if counts else 0,
                'top_values': [{'value': v[:120], 'count': n} for v,n in counts.most_common(5)]}
        if numeric:
            item.update(min=min(numbers), max=max(numbers), mean=statistics.mean(numbers), median=statistics.median(numbers))
        columns.append(item)
    facts = {'source_path': str(path), 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
             'sheet': inputs.get('sheet', ''), 'rows': len(rows), 'column_count': len(fields), 'empty_rows': empty,
             'duplicate_rows': sum(c-1 for c in repeated.values()), 'columns': columns,
             'note': '重复是逐单元格相同记录的额外次数，未自动删除。字段类型是格式检查，不代表标签、因果或业务语义。Excel 使用保存的单元格值，不执行公式。'}
    (folder/'facts.json').write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding='utf-8')
    write_csv(folder/'columns.csv', ['字段','类型','缺失数','不同值数量','最小取值频数','常量'],
              [[c['name'],c['type'],c['missing'],c['unique'],c['smallest_value_count'],c['constant']] for c in columns])
    write_csv(folder/'duplicates.csv', fields+['相同记录数量'], [list(row)+[count] for row,count in repeated.items() if count>1])
    # Full facts remain downloadable; the model sees a bounded overview and five example rows.
    overview = {**facts, 'columns': columns[:40], 'columns_omitted': max(0,len(columns)-40),
                'sample': [{f:r[i][:100] for i,f in enumerate(fields[:20])} for r in rows[:5]]}
    prompt = json.dumps({'user_question': inputs.get('question','帮我看看这份数据'),
        'business_context': inputs.get('business_context',''), 'computed_facts': overview}, ensure_ascii=False)
    if len(prompt) > 24000:
        overview.pop('sample'); overview['columns'] = columns[:20]; overview['columns_omitted'] = max(0,len(columns)-20)
        prompt = json.dumps({'user_question': inputs.get('question','')[:3000],
            'business_context': inputs.get('business_context','')[:3000], 'computed_facts': overview}, ensure_ascii=False)
    return {'facts': facts, 'folder': str(folder), 'prompt': prompt,
            'artifacts': [{'file_path': str(folder/name), 'label': label} for name,label in [
                ('facts.json','完整数据事实'), ('columns.csv','字段检查表 CSV'), ('duplicates.csv','重复记录 CSV')]]}


def assess(inputs):
    advice = inputs['advice']
    questions = [q.strip() for q in advice.get('questions',[]) if isinstance(q,str) and q.strip()][:3]
    needs = advice.get('needs_input') is True and bool(questions)
    context = advice['analysis'] + ('\n\n需要了解：\n'+'\n'.join('- '+q for q in questions) if needs else '')
    return {'needs_input': needs, 'context': context,
            'advice': {**advice, 'questions': questions, 'needs_input': needs}}


def main(inputs):
    operation = inputs['operation']
    if operation == 'profile': return profile(inputs)
    if operation == 'assess': return assess(inputs)
    if operation == 'followup':
        return {'prompt': json.dumps({'original_context': inputs['profile']['prompt'], 'previous_analysis': inputs['advice'],
            'user_answer': inputs['answer'], 'instruction': '依据用户回答更新建议；不清楚也是有效答案。本次不再提问，给出已有结论与具体缺项。'}, ensure_ascii=False)}
    if operation != 'export': raise ValueError('未知分析步骤')
    prepared, advice = inputs['profile'], inputs['advice']
    facts = prepared['facts']; path = source(facts['source_path'])
    if hashlib.sha256(path.read_bytes()).hexdigest() != facts['source_sha256']:
        raise ValueError('等待期间原文件内容改变，请以新资料重新运行，避免混用分析依据')
    folder = Path(prepared['folder'])
    # Always keep final reports independent, including resumed/recomputed variants.
    output = Path('results/data-guidance') / str(uuid4()); output.mkdir(parents=True)
    suggestions = [s for s in advice.get('next_steps',[]) if isinstance(s,str)]
    report = '# 数据摸底与分析建议\n\n## 代码计算的事实\n\n'
    report += f"来源：{facts['source_path']}\n\n{facts['rows']} 行、{facts['column_count']} 列；空行 {facts['empty_rows']} 条；完全重复的额外记录 {facts['duplicate_rows']} 条。原始资料未修改。\n\n"
    report += '## 分析与待确认事项\n\n'+advice['analysis']+'\n\n## 后续建议\n\n'+'\n'.join('- '+s for s in suggestions)
    report += '\n\n建议不是已执行的任务。数值相关或预测表现不能证明因果关系。完整字段检查及重复记录见下载。\n'
    (output/'analysis.md').write_text(report, encoding='utf-8')
    (output/'analysis.json').write_text(json.dumps({'facts':facts,'advice':advice,'answer':inputs.get('answer'),
        'suggestions':suggestions}, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'markdown': report, 'suggestions': suggestions,
        'artifacts': prepared['artifacts'] + [{'file_path':str(output/'analysis.md'),'label':'分析报告 Markdown'},
                                            {'file_path':str(output/'analysis.json'),'label':'完整分析 JSON'}]}

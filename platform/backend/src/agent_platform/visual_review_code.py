"""Code-node implementation: freeze images, collect explicit human input, export."""
import csv
import hashlib
import io
import json
import re
import shutil
from datetime import datetime,timezone
from itertools import islice
from pathlib import Path
from uuid import uuid4


def digest(raw):return hashlib.sha256(raw).hexdigest()


def project_file(value):
    p=Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ['requirement-package','results']:
        raise ValueError('请选择当前项目资料或结果的相对路径')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p,*p.parents]):raise ValueError('不能读取项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size>10_000_000:raise ValueError('文件不存在或超过10 MB：'+str(p))
    return p


def resolve_image(value):
    value=str(value).strip()
    if not value:raise ValueError('图片路径缺失')
    if '/' not in value and '\\' not in value:
        matches=[p for p in Path('requirement-package').rglob('*') if p.name==value and p.is_file()]
        if len(matches)!=1:raise ValueError('图片文件名缺失或不唯一，请填写项目完整路径：'+value)
        value=str(matches[0])
    return project_file(value)


def read_table(path,sheet):
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            if not sheet and len(book.sheetnames)!=1:raise ValueError('Excel有多张表，请指定工作表')
            if sheet and sheet not in book.sheetnames:raise ValueError('Excel工作表不存在：'+sheet)
            values=list(islice((book[sheet] if sheet else book.active).values,10002))
        finally:book.close()
    elif path.suffix.lower() in ['.csv','.tsv']:
        with path.open(encoding='utf-8-sig',newline='') as f:values=list(islice(csv.reader(f,delimiter='\t' if path.suffix.lower()=='.tsv' else ','),10002))
    else:raise ValueError('清单请选择CSV、TSV或XLSX')
    if len(values)<2 or len(values)>10001:raise ValueError('清单需要表头和数据，最多1万条；本次复核另外选择不超过20条')
    fields=[str(x).strip() if x is not None else '' for x in values[0]]
    if any(not x for x in fields) or len(set(fields))!=len(fields):raise ValueError('表头不能为空或重复')
    rows=[]
    for n,values in enumerate(values[1:],2):
        if all(x is None or not str(x).strip() for x in values):continue
        if len(values)!=len(fields):raise ValueError(f'第{n}行列数不一致')
        rows.append(dict(zip(fields,[str(x).strip() if x is not None else '' for x in values])))
    if not rows:raise ValueError('清单没有数据')
    return fields,rows


def folder():
    p=Path('results')/('visual-review-'+uuid4().hex);p.mkdir(parents=True);return p


def integer(value,label,low,high):
    if isinstance(value,bool):raise ValueError(label+'需要整数')
    try:n=int(value)
    except (ValueError,TypeError,OverflowError):raise ValueError(label+'需要整数') from None
    if n!=float(value) or not low<=n<=high:raise ValueError(f'{label}需要{low}至{high}的整数')
    return n


def snapshot_image(value,root):
    from PIL import Image
    p=resolve_image(value);raw=p.read_bytes()
    if len(raw)>5_000_000:raise ValueError('单张图片超过5 MB：'+str(p))
    try:
        with Image.open(io.BytesIO(raw)) as im:
            if im.format not in ['PNG','JPEG','WEBP'] or im.width*im.height>20_000_000 or getattr(im,'n_frames',1)>1:raise ValueError('需要单帧PNG/JPEG/WebP且不超过2000万像素')
            fmt=im.format;size=list(im.size);im.verify()
    except Exception as e:raise ValueError('无法读取图片 '+str(p)+'：'+str(e)) from e
    sha=digest(raw);destination=root/(sha+{'PNG':'.png','JPEG':'.jpg','WEBP':'.webp'}[fmt]);destination.write_bytes(raw)
    return {'source_path':str(p),'path':str(destination),'sha256':sha,'size':size}


def escape(value):return re.sub(r'([\\`*_{}\[\]<>()!|#])',r'\\\1',str(value)).replace('\n',' / ').replace('\r',' ')


def prepare(inputs):
    path=project_file(inputs.get('source_path',''));fields,rows=read_table(path,inputs.get('sheet') or '')
    keys={k:str(inputs.get(k) or default) for k,default in [('id_column','sample_id'),('image_column','image_path'),('product_column','product_type'),('prediction_column','machine_prediction')]}
    missing=[keys[k] for k in ['id_column','image_column'] if keys[k] not in fields]
    if missing:raise ValueError('清单缺少字段：'+'、'.join(missing))
    ids=[r[keys['id_column']] for r in rows]
    if any(not x for x in ids) or len(set(ids))!=len(ids):raise ValueError('样本标识为空或重复，请区分产品、检测次数或版本，不按行号猜测')
    start=integer(inputs.get('start_row',1),'起始记录',1,len(rows));limit=integer(inputs.get('sample_limit',10),'复核数量',1,20)
    criteria=str(inputs.get('criteria') or '').strip()
    if len(criteria)>8000:raise ValueError('判定依据最多8000字')
    selected=rows[start-1:start-1+limit];root=folder();samples=[]
    try:
        for offset,row in enumerate(selected,start):
            try:image=snapshot_image(row[keys['image_column']],root);reference=snapshot_image(row['reference_path'],root) if row.get('reference_path') else None
            except ValueError as e:raise ValueError('样本 '+row[keys['id_column']]+'：'+str(e)) from e
            warnings=[];product=row.get(keys['product_column'],'')
            if not product:warnings.append('没有产品类型，不能直接迁移判断方法')
            if not criteria:warnings.append('尚未提供判定依据；不能仅凭外观猜测是否合格')
            if reference:
                if not row.get('reference_version'):warnings.append('参考图片没有版本说明')
                if not row.get('reference_product'):warnings.append('参考图片的适用产品不明')
                elif product!=row['reference_product']:warnings.append('参考图片与本样本产品类型不同')
                if row.get('reference_available_at') and row.get('captured_at'):
                    try:
                        later=datetime.fromisoformat(row['reference_available_at'].replace('Z','+00:00'));earlier=datetime.fromisoformat(row['captured_at'].replace('Z','+00:00'))
                        if later.tzinfo is None or earlier.tzinfo is None:raise ValueError()
                        if later>earlier:warnings.append('参考图片形成于检测之后，不能当作检测当时已知依据')
                    except ValueError:warnings.append('参考或检测时间格式/时区不明确，无法判断当时可用性')
                else:warnings.append('缺少参考形成或检测时间，无法确认参考在检测当时可用')
            samples.append(dict(index=offset,sample_id=row[keys['id_column']],product=product,machine=row.get(keys['prediction_column'],''),image=image,reference=reference,metadata=row,warnings=warnings))
        counts={}
        for sample in samples:counts[sample['image']['sha256']]=counts.get(sample['image']['sha256'],0)+1
        for sample in samples:
            if counts[sample['image']['sha256']]>1:sample['warnings'].append('本次其他样本使用相同图片，不应当作独立训练/测试样本')
        data=dict(format='visual-review-input-v1',source={'path':str(path),'sha256':digest(path.read_bytes())},total=len(rows),start=start,selected=len(samples),not_selected=len(rows)-len(samples),criteria=criteria,samples=samples)
        snapshot=root/'input.json';raw=json.dumps(data,ensure_ascii=False,indent=2).encode();snapshot.write_bytes(raw)
        return {'path':str(snapshot),'sha256':digest(raw),'samples':[{'index':s['index']} for s in samples]}
    except BaseException:shutil.rmtree(root);raise


def frozen(prepared):
    raw=project_file(prepared['path']).read_bytes()
    if digest(raw)!=prepared['sha256']:raise ValueError('本次清单快照改变，请重新开始复核')
    data=json.loads(raw)
    if data.get('format')!='visual-review-input-v1':raise ValueError('不是本流程输入记录')
    return data


def sample_for(inputs):
    data=frozen(inputs['prepared']);index=inputs['sample']['index']
    sample=next((s for s in data['samples'] if s['index']==index),None)
    if sample is None:raise ValueError('样本不属于本次复核')
    for img in [sample['image'],sample['reference']]:
        if img and digest(project_file(img['path']).read_bytes())!=img['sha256']:raise ValueError('本次图片快照已改变，请重新开始复核')
    return data,sample


def context(inputs):
    data,s=sample_for(inputs)
    markdown=f"### 样本 {escape(s['sample_id'])}（本清单第{s['index']}条）\n\n产品：{escape(s['product'] or '未说明')}；机器原判断：{escape(s['machine'] or '未提供')}。\n\n判定依据：{escape(data['criteria'] or '未提供，可选择无法判断')}\n\n"
    markdown+='\n'.join('- '+escape(w) for w in s['warnings'])
    labels={'batch':'批次','reference_product':'参考适用产品','reference_version':'参考版本','reference_available_at':'参考形成时间','captured_at':'检测时间'}
    metadata=[label+'：'+escape(s['metadata'][key]) for key,label in labels.items() if s['metadata'].get(key)]
    if metadata:markdown+='\n\n相关记录：\n\n'+'\n'.join('- '+line for line in metadata)
    return {'markdown':markdown,'images':[{'path':s['image']['path'],'label':'待复核图片'}]+([{'path':s['reference']['path'],'label':'参考图片（适用范围请核对）'}] if s['reference'] else [])}


def collect(inputs):
    _,s=sample_for(inputs);answer=inputs.get('answer',{});label=answer.get('label')
    if label not in ['合格','不合格','无法判断']:raise ValueError('复判结论请选择合格、不合格或无法判断')
    reason=str(answer.get('reason') or '').strip();method=str(answer.get('method') or '').strip()
    if max(len(reason),len(method))>8000:raise ValueError('每项说明最多8000字')
    requested=answer.get('reference_candidate') is True
    candidate=requested and label!='无法判断' and bool(reason) and bool(s['product'])
    record={'sample_id':s['sample_id'],'index':s['index'],'product':s['product'],'machine_prediction':s['machine'],'human_label':label,'reason':reason,'method':method,
            'agreement':'无法判断' if label=='无法判断' else '没有可比较机器判断' if s['machine'] not in ['合格','不合格'] else '一致' if label==s['machine'] else '不一致',
            'reference_requested':requested,'reference_candidate':candidate,'candidate_note':'仅为参考候选，尚未应用到参考库或模型' if candidate else '明确结论、产品类型、理由或人工选择缺失，不列候选',
            'reviewed_at':datetime.now(timezone.utc).isoformat(),'sample':s}
    path=folder()/'review-item.json';raw=json.dumps(record,ensure_ascii=False,indent=2).encode();path.write_bytes(raw)
    return {'path':str(path),'sha256':digest(raw)}


def finish(inputs):
    data=frozen(inputs['prepared']);reviews=[]
    for result in inputs['reviews']:
        raw=project_file(result['path']).read_bytes()
        if digest(raw)!=result['sha256']:raise ValueError('已保存复核记录改变，请检查原运行')
        reviews.append(json.loads(raw))
    if [r['sample_id'] for r in reviews]!=[s['sample_id'] for s in data['samples']]:raise ValueError('复核结果与本次样本不一致')
    record={'format':'visual-review-result-v1','prepared':inputs['prepared'],'reviews':reviews,'item_files':inputs['reviews']}
    path=folder()/'review.json';path.write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
    return report({'source_result_path':str(path)})


def csv_file(path,rows,fields):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader()
        w.writerows({k:("'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v) for k,v in row.items()} for row in rows)


def report(inputs):
    source=project_file(inputs['source_result_path']);saved=json.loads(source.read_text())
    if saved.get('format')!='visual-review-result-v1':raise ValueError('请选择本流程生成的review.json')
    data=frozen(saved['prepared']);reviews=saved['reviews'];root=folder()
    rows=[{**{k:r[k] for k in ['sample_id','product','machine_prediction','human_label','agreement','reason','method','reference_candidate','candidate_note','reviewed_at']},
           'image_path':r['sample']['image']['path'],'image_sha256':r['sample']['image']['sha256'],'warnings':'；'.join(r['sample']['warnings'])} for r in reviews]
    candidates=[r for r in rows if r['reference_candidate']];unknown=[r for r in rows if r['human_label']=='无法判断'];different=[r for r in rows if r['agreement']=='不一致']
    fields=list(rows[0])
    csv_file(root/'reviews.csv',rows,fields);csv_file(root/'reference-candidates.csv',candidates,fields);csv_file(root/'unknown.csv',unknown,fields)
    csv_file(root/'labeled.csv',[r for r in rows if r['human_label']!='无法判断'],fields)
    markdown=f"# 看图复核与经验记录\n\n清单共{data['total']}条，本次复核{len(rows)}条，未选入{data['not_selected']}条。未知{len(unknown)}条，与机器判断不一致{len(different)}条，参考候选{len(candidates)}条。\n\n这是本次人工记录，不是视觉模型准确率或生产放行结论。尚未训练、修改参考库或写回现场。\n\n"
    markdown+='| 样本 | 产品 | 机器原判断 | 人工复判 | 对照 | 理由或缺项 |\n| --- | --- | --- | --- | --- | --- |\n'
    for r in rows:markdown+='| '+' | '.join(escape(r[k]) for k in ['sample_id','product','machine_prediction','human_label','agreement','reason'])+' |\n'
    markdown+='\n## 可复用经验与适用条件\n\n'
    for r in reviews:
        markdown+='### '+escape(r['sample_id'])+'\n\n'+escape(r['method'] or '未填写方法；不从结论自动生成规则')+'\n\n'
        markdown+='观察与限制：'+escape('；'.join(r['sample']['warnings']) or '本次没有检测到元数据缺项，仍需核对业务适用条件')+'\n\n'
        markdown+='[查看本次图片]('+r['sample']['image']['path']+')\n\n'
    markdown+='## 如何继续\n\n未知项补资料后另开复核；有价值的方法可由员工整理成项目Skill。参考候选仍需匹配产品、版本与时点，不自动成为模型训练集；相同图片或同一产品的多次检测不能随机拆进训练和测试。分析数值/类别误差可再调用已有反馈对照流程。\n'
    (root/'report.md').write_text(markdown,encoding='utf-8')
    artifacts=[{'file_path':str(root/name),'label':label} for name,label in [('report.md','复核与经验 Markdown'),('reviews.csv','逐条复核 CSV'),('reference-candidates.csv','参考样本候选 CSV'),('unknown.csv','仍需补充 CSV'),('labeled.csv','明确人工标签 CSV')]]
    artifacts+=[{'file_path':str(source),'label':'可再次导出的复核 JSON'},{'file_path':saved['prepared']['path'],'label':'原清单与图片版本 JSON'}]
    return {'reviewed':len(rows),'unknown':len(unknown),'different':len(different),'candidates':len(candidates),'source_result_path':str(source),'source_path':str(root/'reviews.csv'),
            'labeled_path':str(root/'labeled.csv'),'feedback_inputs':{'source_path':str(root/'labeled.csv'),'mode':'类别判断','key_columns':'sample_id','prediction_column':'machine_prediction','actual_column':'human_label','group_columns':'product'},'artifacts':artifacts,'markdown':markdown}


def main(inputs):return {'prepare':prepare,'context':context,'collect':collect,'finish':finish,'report':report}[inputs['operation']](inputs)

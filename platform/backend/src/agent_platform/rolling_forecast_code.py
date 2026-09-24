"""Portable temporal preparation and reporting; fitting uses native model_train/model_predict."""
import csv
import hashlib
import io
import json
import math
import re
import zipfile
from bisect import bisect_right
from collections import Counter,defaultdict
from datetime import datetime,timedelta,timezone
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value,limit=20_000_000):
    p=Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ['requirement-package','results']:
        raise ValueError('请选择当前项目资料或结果文件')
    if any(v.is_symlink() for v in [p,*p.parents]) or not p.resolve().is_relative_to(Path.cwd().resolve()):raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size>limit:raise ValueError('文件不存在或超过允许大小，请拆分')
    return p


def read_table(value,sheet=''):
    path=source(value,10_000_000)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            if not sheet and len(book.sheetnames)>1:raise ValueError('Excel有多张表，请指定工作表')
            if sheet and sheet not in book.sheetnames:raise ValueError('指定工作表不存在')
            records=list(islice((book[sheet] if sheet else book.active).values,20002))
        finally:book.close()
    elif path.suffix.lower() in ['.csv','.tsv']:
        records=list(islice(csv.reader(io.StringIO(path.read_text(encoding='utf-8-sig')),delimiter='\t' if path.suffix.lower()=='.tsv' else ','),20002))
    else:raise ValueError('请选择CSV、TSV或XLSX')
    if not records or len(records)>20001:raise ValueError('单表需要表头，最多2万条记录')
    fields=[str(v).strip() if v is not None else '' for v in records[0]]
    if any(not v for v in fields) or len(set(fields))!=len(fields):raise ValueError('表头不能为空或重名')
    rows=[]
    for i,r in enumerate(records[1:],2):
        if all(v is None or str(v).strip()=='' for v in r):continue
        if len(r)!=len(fields):raise ValueError(f'第{i}条记录的列数与表头不同')
        rows.append((i,{k:'' if v is None else str(v).strip() for k,v in zip(fields,r)}))
    if not rows:raise ValueError('没有可用记录')
    return path,fields,rows


def number(value,label):
    try:
        n=float(value)
        if math.isfinite(n) and abs(n)<=1e20 and not isinstance(value,bool):return n
    except (TypeError,ValueError):pass
    raise ValueError(label+'需要有限数值，绝对值不超过1e20')


def integer(value,label,low,high):
    n=number(value,label)
    if not n.is_integer() or not low<=n<=high:raise ValueError(f'{label}需要{low}至{high}之间的整数')
    return int(n)


def zone(value):
    if not isinstance(value,str) or not re.fullmatch(r'[+-]\d{2}:\d{2}',value):raise ValueError('默认时区请填写UTC偏移，例如+08:00')
    h,m=int(value[1:3]),int(value[4:6])
    if h>14 or m>59 or (h==14 and m):raise ValueError('UTC偏移超出范围')
    return timezone(timedelta(minutes=(h*60+m)*(1 if value[0]=='+' else -1)))


def stamp(value,tz,label):
    try:
        t=datetime.fromisoformat(str(value).strip().replace('Z','+00:00'))
        if t.tzinfo is None:t=t.replace(tzinfo=tz)
        return t.astimezone(timezone.utc)
    except (ValueError,TypeError):raise ValueError(label+'不是有效ISO时间') from None


def folder(prefix):
    p=Path('results')/(prefix+'-'+uuid4().hex);p.mkdir(parents=True);return p


def save_csv(path,fields,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(rows)


def prepare(inputs):
    path,fields,rows=read_table(inputs['source_path'],inputs.get('sheet') or '')
    keys={k:str(inputs.get(k) or '').strip() for k in ['time_column','value_column','series_column','available_column']}
    if any(keys[k] not in fields for k in ['time_column','value_column']) or any(keys[k] and keys[k] not in fields for k in ['series_column','available_column']):
        raise ValueError('请检查时间、数值、序列标识和可用时间列是否存在')
    config={**keys,'timezone':inputs.get('timezone','+00:00')};tz=zone(config['timezone'])
    for key,default,low,high in [('step_seconds',86400,1,31536000),('lags',7,1,100),('horizon',3,1,96),('origins',3,1,20),
        ('origin_stride',3,1,1000),('training_window',0,0,10000),('seasonal_period',7,1,1000)]:
        config[key]=integer(inputs.get(key,default),key,low,high)
    config['forecast_last']=inputs.get('forecast_last',False)
    if not isinstance(config['forecast_last'],bool):raise ValueError('最后时点预测请选择是或否')
    config['model_choice']=inputs.get('model_choice','仅线性回归')
    if config['model_choice'] not in ['仅线性回归','线性与随机森林比较']:raise ValueError('请选择支持的模型比较方式')
    config['direction_tolerance']=number(inputs.get('direction_tolerance',0),'方向容差')
    if config['direction_tolerance']<0:raise ValueError('方向容差不能为负')
    for key in ['lower_boundary','upper_boundary']:config[key]=None if inputs.get(key) in [None,''] else number(inputs[key],key)
    if config['lower_boundary'] is not None and config['upper_boundary'] is not None and config['lower_boundary']>=config['upper_boundary']:raise ValueError('低边界必须小于高边界')
    series=defaultdict(list);seen={};issues=[]
    for record,row in rows:
        name=row[keys['series_column']] if keys['series_column'] else '全部序列'
        if not name:raise ValueError(f'第{record}条记录缺少序列标识')
        observed=stamp(row[keys['time_column']],tz,f'第{record}条观测时刻')
        available=stamp(row[keys['available_column']],tz,f'第{record}条可用时刻') if keys['available_column'] else observed
        if available<observed:raise ValueError(f'第{record}条记录的实际数值不能在观测发生前可用')
        ident=(name,observed,available)
        if ident in seen:raise ValueError(f'第{record}与第{seen[ident]}条记录的序列、观测和可用时刻重复，请先明确版本')
        seen[ident]=record
        try:value=number(row[keys['value_column']],'数值')
        except ValueError:value=None;issues.append(dict(source_record=record,series=name,reason='数值缺失或无效，保留时间位置，不填0'))
        series[name].append(dict(observed=observed.isoformat(),available=available.isoformat(),value=value,source_record=record))
    if len(series)>10:raise ValueError('一次最多10个序列，请拆分项目资料')
    jobs=[]
    for name,records in series.items():
        records.sort(key=lambda r:(r['observed'],r['available']));first=stamp(records[0]['observed'],tz,'时间');last=stamp(records[-1]['observed'],tz,'时间')
        step=config['step_seconds'];length=(last-first).total_seconds()/step
        if length>10000:raise ValueError(f'{name}的时间跨度超过1万时段，请拆分或调整实际频率')
        for r in records:
            delta=(stamp(r['observed'],tz,'时间')-first).total_seconds()/step
            if not delta.is_integer():raise ValueError(f"第{r['source_record']}条观测不在指定时间网格，检查step_seconds和时区")
            r['index']=int(delta)
        last_index=int(length)
        start_index=last_index-config['horizon']-(config['origins']-1)*config['origin_stride']
        if inputs.get('first_origin'):
            index=(stamp(inputs['first_origin'],tz,'首个起点')-first).total_seconds()/step
            if not index.is_integer():raise ValueError('首个起点不在观测时间网格上')
            start_index=int(index)
        for i in range(config['origins']):
            index=start_index+i*config['origin_stride']
            if index<config['lags']+8 or index+config['horizon']>last_index:raise ValueError(f'{name}没有足够历史和完整预测范围，请减少起点/步长或补充资料')
            jobs.append(dict(series=name,index=index,origin=(first+timedelta(seconds=index*step)).isoformat(),kind='backtest'))
        if config['forecast_last']:jobs.append(dict(series=name,index=last_index,origin=last.isoformat(),kind='forecast'))
    if len(jobs)>20:raise ValueError('一次最多20次逐起点训练，请减少序列或回测起点数量')
    config['first_origin']=inputs.get('first_origin') or ''
    data=dict(config=config,series=dict(series),issues=issues,source={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()},jobs=jobs)
    root=folder('forecast-input');snapshot=root/'input.json';snapshot.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    return {'snapshot_path':str(snapshot),'sha256':hashlib.sha256(snapshot.read_bytes()).hexdigest(),'jobs':jobs}


def load(prepared):
    raw=source(prepared['snapshot_path']).read_bytes()
    if hashlib.sha256(raw).hexdigest()!=prepared['sha256']:raise ValueError('预测输入快照已改变，请重新读取资料')
    return json.loads(raw)


def window(inputs):
    data=load(inputs['prepared']);job=inputs['job'];cfg=data['config']
    if job not in data['jobs']:raise ValueError('起点不属于本次固定配置')
    records=data['series'][job['series']];versions=defaultdict(list)
    for r in records:versions[r['index']].append(r)
    available={i:[r['available'] for r in values] for i,values in versions.items()}
    first=datetime.fromisoformat(records[0]['observed']);step=cfg['step_seconds'];cutoff=job['origin'];oi=job['index'];features=['horizon']+['lag_'+str(i) for i in range(cfg['lags'])]
    def at(i,when):
        k=bisect_right(available.get(i,[]),when)-1
        return versions[i][k] if k>=0 else None
    def time(i):return (first+timedelta(seconds=i*step)).isoformat()
    def history(i):return [at(i-j,time(i)) for j in range(cfg['lags'])]
    def values(hist):return {'lag_'+str(j):r['value'] if r and r['value'] is not None else '' for j,r in enumerate(hist)}
    train=[];usage=[];excluded=Counter();start=max(0,oi-cfg['training_window']) if cfg['training_window'] else 0
    for i in range(start,oi):
        hist=history(i)
        if not any(r and r['value'] is not None for r in hist):excluded['历史滞后值全部缺失']+=cfg['horizon'];continue
        for h in range(1,cfg['horizon']+1):
            label=at(i+h,cutoff)
            if i+h>oi or not label or label['value'] is None:excluded['标签在本起点尚不可用或无效']+=1;continue
            sid=f'{i}:{h}'
            train.append(dict(sample_id=sid,origin_time=time(i),label_available_time=label['available'],horizon=h,**values(hist),target=label['value']))
            usage.append(dict(sample_id=sid,origin_time=time(i),target_time=time(i+h),label_source_record=label['source_record'],
                label_available_time=label['available'],feature_source_records=';'.join(str(r['source_record']) if r else '' for r in hist)))
    if len(train)<24 or len({r['origin_time'] for r in train})<12:raise ValueError(f"{job['series']}在{cutoff}可用训练样本不足：{len(train)}条；至少24条且12个不同样本起点")
    if len(train)>20000:raise ValueError('一个起点超过2万训练样本，请限制历史窗口或预测步数')
    hist=history(oi)
    if not any(r and r['value'] is not None for r in hist):raise ValueError(f"{job['series']}在{cutoff}全部滞后输入不可用，请补充资料或修改起点")
    known=[at(i,cutoff) for i in sorted(versions) if i<=oi];known=[r for r in known if r and r['value'] is not None]
    latest=known[-1] if known else None;predict=[];truth=[]
    for h in range(1,cfg['horizon']+1):
        sid=f'{oi}:{h}';target_index=oi+h;actual=versions.get(target_index,[])
        actual=actual[-1] if actual else None
        seasonal_index=target_index-math.ceil(h/cfg['seasonal_period'])*cfg['seasonal_period']
        seasonal=at(seasonal_index,cutoff)
        predict.append(dict(sample_id=sid,horizon=h,**values(hist)))
        truth.append(dict(sample_id=sid,series=job['series'],origin=cutoff,kind=job['kind'],horizon=h,target_time=time(target_index),
            actual=actual['value'] if actual else None,actual_source_record=actual['source_record'] if actual else None,
            actual_available_time=actual['available'] if actual else '',last_value=latest['value'] if latest else None,
            last_value_time=latest['observed'] if latest else '',seasonal_naive=seasonal['value'] if seasonal else None,
            seasonal_source_record=seasonal['source_record'] if seasonal else None))
    root=folder('forecast-window');save_csv(root/'train.csv',['sample_id','origin_time','label_available_time',*features,'target'],train)
    save_csv(root/'prediction-input.csv',['sample_id',*features],predict)
    save_csv(root/'sample-sources.csv',['sample_id','origin_time','target_time','label_source_record','label_available_time','feature_source_records'],usage)
    metadata=dict(job=job,train_rows=len(train),excluded=dict(excluded),features=features,truth=truth,source=data['source'],config=cfg,
                  max_label_available=max(r['label_available_time'] for r in train),forecast_feature_sources=[r['source_record'] if r else None for r in hist])
    (root/'window.json').write_text(json.dumps(metadata,ensure_ascii=False),encoding='utf-8')
    return {'train_path':str(root/'train.csv'),'predict_path':str(root/'prediction-input.csv'),'window_path':str(root/'window.json'),
        'features':features,'candidate':{'engine':'sklearn','models':['linear'] if cfg['model_choice']=='仅线性回归' else ['linear','forest'],'batch_size':1 if cfg['model_choice']=='仅线性回归' else 2},
        'artifacts':[{'file_path':str(root/name),'label':label} for name,label in [('train.csv','本起点训练样本 CSV'),('prediction-input.csv','本起点预测输入 CSV'),('sample-sources.csv','样本特征与标签来源 CSV'),('window.json','起点与输入版本 JSON')]]}


def collect(inputs):
    window=inputs['window'];meta=json.loads(source(window['window_path']).read_text());prediction=inputs['prediction'];version=prediction.get('model_version')
    if not version:raise ValueError('缺少实际模型版本，不能以基线冒充模型预测')
    _,_,records=read_table(prediction['project_path']);byid={}
    for _,r in records:
        if r['sample_id'] in byid:raise ValueError('预测样本标识重复')
        byid[r['sample_id']]=number(r['prediction'],'预测值')
    if set(byid)!={r['sample_id'] for r in meta['truth']}:raise ValueError('预测结果与本起点输入不一致')
    rows=[{**r,'model_prediction':byid[r['sample_id']],**version} for r in meta['truth']]
    root=folder('forecast-origin');path=root/'predictions.csv';save_csv(path,list(rows[0]),rows)
    return {'source_path':str(path),'window_path':window['window_path'],'job':meta['job'],'train_rows':meta['train_rows'],'excluded':meta['excluded'],
        'model_version':version,'trials':[{k:t.get(k) for k in ['slot','model','status','metrics','error']} for t in inputs['training'].get('trials',[])],
        'artifacts':window['artifacts']+[{'file_path':str(path),'label':'本起点预测、实际与基线 CSV'}]}


def sign(v,tol):return 1 if v>tol else -1 if v<-tol else 0


def metrics(rows,key,cfg):
    if not rows:return {'rows':0,'mae':None,'rmse':None,'bias':None,'direction_accuracy':None}
    errors=[r[key]-r['actual'] for r in rows];direction=[sign(r[key]-r['last_value'],cfg['direction_tolerance'])==sign(r['actual']-r['last_value'],cfg['direction_tolerance']) for r in rows if r['last_value'] is not None]
    result={'rows':len(rows),'mae':sum(abs(v) for v in errors)/len(errors),'rmse':math.sqrt(sum(v*v for v in errors)/len(errors)),
        'bias':sum(errors)/len(errors),'direction_accuracy':sum(direction)/len(direction) if direction else None}
    for name,field,op in [('lower','lower_boundary',lambda a,b:a<=b),('upper','upper_boundary',lambda a,b:a>=b)]:
        threshold=cfg[field]
        if threshold is not None:
            actual=[op(r['actual'],threshold) for r in rows];pred=[op(r[key],threshold) for r in rows];tp=sum(a and b for a,b in zip(actual,pred))
            result.update({name+'_events':sum(actual),name+'_predicted_events':sum(pred),name+'_recall':tp/sum(actual) if sum(actual) else None,name+'_precision':tp/sum(pred) if sum(pred) else None})
    return result


def rank_correlation(a,b):
    def ranks(values):
        ordered=sorted(range(len(values)),key=values.__getitem__);result=[0.]*len(values);i=0
        while i<len(ordered):
            j=i+1
            while j<len(ordered) and values[ordered[j]]==values[ordered[i]]:j+=1
            for k in range(i,j):result[ordered[k]]=(i+j-1)/2
            i=j
        return result
    if len(a)<2:return None
    x,y=ranks(a),ranks(b);center=(len(x)-1)/2
    xx=sum((v-center)**2 for v in x);yy=sum((v-center)**2 for v in y)
    return sum((u-center)*(v-center) for u,v in zip(x,y))/math.sqrt(xx*yy) if xx and yy else None


def report(inputs):
    data=load(inputs['prepared']);cfg=data['config'];results=inputs['results'];all_rows=[]
    if len(results)!=len(data['jobs']):raise ValueError('部分预测起点尚未完成，请查看原任务状态')
    bystep=defaultdict(list);byorigin=defaultdict(list);root=folder('rolling-forecast');methods=['model_prediction','last_value','seasonal_naive']
    for result in results:
        _,_,records=read_table(result['source_path'])
        for _,r in records:
            r['horizon']=int(r['horizon'])
            for k in ['actual',*methods]:r[k]=None if r[k]=='' else number(r[k],k)
            all_rows.append(r)
            if r['kind']=='backtest':bystep[(r['series'],r['horizon'])].append(r);byorigin[(r['series'],r['origin'])].append(r)
    summaries=[];origin_summary=[]
    for (name,horizon),rows in bystep.items():
        common=[r for r in rows if r['actual'] is not None and all(r[m] is not None for m in methods)]
        for key in methods:
            own=[r for r in rows if r['actual'] is not None and r[key] is not None]
            summaries.append({'series':name,'horizon':horizon,'method':key,'requested':len(rows),'own_rows':len(own),
                'own_mae':metrics(own,key,cfg)['mae'],**metrics(common,key,cfg)})
    for (name,origin),rows in byorigin.items():
        common=[r for r in rows if r['actual'] is not None and all(r[m] is not None for m in methods)]
        for key in methods:origin_summary.append({'series':name,'origin':origin,'method':key,**metrics(common,key,cfg),
            'rank_correlation':rank_correlation([r[key] for r in common],[r['actual'] for r in common])})
    save_csv(root/'predictions.csv',list(all_rows[0]),all_rows)
    def union_csv(name,rows):save_csv(root/name,list(dict.fromkeys(k for r in rows for k in r)),rows)
    union_csv('by-horizon.csv',summaries);union_csv('by-origin.csv',origin_summary)
    result=dict(source_path=str(root/'predictions.csv'),rows=len(all_rows),config=cfg,source=data['source'],issues=data['issues'],
        by_horizon=summaries,by_origin=origin_summary,origins=[{k:r[k] for k in ['job','train_rows','excluded','model_version','trials','window_path']} for r in results],
        scope='历史回测按预测决策计数，不等于独立样本或收益；未来预测尚无实际值，不计入评价。没有验证真实发布时间的输入不构成生产时点验收。')
    labels={'model_prediction':'历史内验证选出的模型','last_value':'最近可知值','seasonal_naive':'季节基线'}
    def cell(v):return str(v).replace('|','\\|').replace('\n',' ').replace('\r',' ')[:180]
    def fmt(v):return '无法计算' if v is None else f'{v:.5g}'
    text=f"# 多步预测与滚动回测\n\n完成 {len(results)} 次逐起点训练与预测，输出 {len(all_rows)} 条预测。历史回测 {sum(j['kind']=='backtest' for j in data['jobs'])} 个起点，最后时点预测 {sum(j['kind']=='forecast' for j in data['jobs'])} 个。\n\n"
    text+='每次只在起点已可知标签内训练，内部验证选择模型；未来实际值只用于评价。模型与预处理保存在项目建模记录，逐起点版本可下载。\n\n'
    if not cfg['available_column']:text+='**未提供实际可用时间：本次假定观测时刻即可使用，现实发布时间尚未验证。**\n\n'
    text+='## 按序列与步长比较\n\n| 序列 | 步长 | 方法 | 共同记录/计划数 | MAE | RMSE | 方向正确率 |\n| --- | --- | --- | --- | --- | --- | --- |\n'
    for s in summaries[:60]:text+='| '+' | '.join(map(cell,[s['series'],s['horizon'],labels[s['method']],f"{s['rows']}/{s['requested']}",fmt(s['mae']),fmt(s['rmse']),fmt(s['direction_accuracy'])]))+' |\n'
    text+='\n预览最多60行。误差使用三种方法都可评价的相同记录；各方法自身可评价数量及误差另见下载，缺项不会算作0。MAE/RMSE越小越好，方向以起点最近可知值为参照，使用填写的平稳容差。不同量纲序列不合并排名。\n\n'
    text+='## 各起点覆盖与缺项\n\n'
    for r in results[:20]:text+=f"- {cell(r['job']['series'])} · {r['job']['origin']} · {'历史回测' if r['job']['kind']=='backtest' else '最后时点预测'}：训练 {r['train_rows']} 条，排除原因 {cell(r['excluded'])}。\n"
    text+=f"\n每步为 {cfg['step_seconds']} 秒，使用 {cfg['lags']} 个滞后时段；训练窗口 {'逐步扩大' if not cfg['training_window'] else str(cfg['training_window'])+' 时段'}，季节周期 {cfg['seasonal_period']}。缺时段不压缩，不用未来实际值递归预测。实际值取所给文件最新版本，输入和标签来源可追溯。\n\n"
    text+='## 边界事件与相对排序\n\n上下边界只按本次填写的阈值计算，精确率/召回率见逐步和逐起点指标CSV。无真实事件时召回率无法计算，无预测事件时精确率无法计算。每个起点另有未来各步的秩相关，常量序列或不足两点不计算；这些指标不等同预测点位误差或业务收益。\n\n'
    text+='## 下一步\n\n先检查模型是否在共同记录上超过简单基线，以及不同步长和起点是否稳定。修改滞后窗口、历史长度或模型后新建运行；回测结果用于调参后不能再称独立测试，应保留新的后续时间段复验。未来预测没有实际值时保留预测，收到实测后可用已有反馈流程比较。\n\n'+result['scope']
    (root/'summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');(root/'report.md').write_text(text,encoding='utf-8')
    with zipfile.ZipFile(root/'origin-inputs.zip','w',zipfile.ZIP_DEFLATED) as archive:
        archive.write(source(inputs['prepared']['snapshot_path']), 'input.json')
        for i,r in enumerate(results,1):
            for a in r['artifacts']:
                path=source(a['file_path']);archive.write(path,f'origin-{i}/{path.name}')
    artifacts=[{'file_path':str(root/name),'label':label} for name,label in [('predictions.csv','预测、实际与基线 CSV'),('by-horizon.csv','逐步误差与边界事件 CSV'),('by-origin.csv','逐起点误差与相对排序 CSV'),('summary.json','配置、模型版本与覆盖 JSON'),('report.md','滚动回测说明 Markdown'),('origin-inputs.zip','各起点样本与来源 ZIP')]]
    return {**result,'markdown':text,'artifacts':artifacts,'stage_artifacts':[a for r in results for a in r['artifacts']]}


def main(inputs):
    return {'prepare':prepare,'window':window,'collect':collect,'report':report}[inputs['operation']](inputs)

"""Portable deterministic row preparation; no fitted statistics or model calls."""
import csv
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, localcontext
from itertools import islice
from pathlib import Path
from uuid import uuid4

METHOD_VERSION = 'row-preparation-1'
OPS = ['两列相加','第一列减第二列','两列差的绝对值','两列相乘','第一列除以第二列','乘固定数','加固定数']


def source(value, limit=10_000_000):
    p = Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ('requirement-package','results'):
        raise ValueError('请选择当前项目资料或结果文件')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p,*p.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size > limit:
        raise ValueError(f'资料不存在或超过{limit//1_000_000} MB，请检查或拆分文件')
    return p


def table(value, sheet=''):
    p=source(value)
    if p.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(p,read_only=True,data_only=True)
        try:
            if not sheet and len(book.sheetnames)>1:raise ValueError('Excel有多张表，请指定工作表：'+'、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:raise ValueError('工作表不存在：'+sheet)
            rows=list(islice((book[sheet] if sheet else book.active).values,20002))
        finally:book.close()
    elif p.suffix.lower() in ('.csv','.tsv'):
        with p.open(encoding='utf-8-sig',newline='') as f:rows=list(islice(csv.reader(f,delimiter='\t' if p.suffix.lower()=='.tsv' else ','),20002))
    else:raise ValueError('请选择CSV、TSV或XLSX；旧XLS请先另存')
    if len(rows)<2 or len(rows)>20001:raise ValueError('需要表头和数据，首版最多2万条记录')
    fields=[str(v).strip() if v is not None else '' for v in rows[0]]
    if not fields or len(fields)>200 or any(not f or f.startswith('preparation_') for f in fields) or len(set(fields))!=len(fields):
        raise ValueError('表头须非空、不重名、最多200列；preparation_前缀留作排除记录说明')
    result=[]
    for n,raw in enumerate(rows[1:],2):
        if not any(v is not None and str(v).strip() for v in raw):raw=['']*len(fields)
        if len(raw)!=len(fields):raise ValueError(f'第{n}条记录列数与表头不同')
        result.append((n,dict(zip(fields,['' if v is None else str(v) for v in raw]))))
    return fields,result


def digest(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,allow_nan=False).encode()).hexdigest()


def processor_hash():
    # Platform code nodes receive the hash of the actual executed source. Direct
    # local imports are also useful for checking the portable implementation.
    return globals().get('__workflow_code_sha256__') or hashlib.sha256(Path(__file__).read_bytes()).hexdigest()


def number(value):
    try:
        n=Decimal(str(value).strip())
        if n.is_finite() and len(n.as_tuple().digits)<=40 and (n==0 or -40<=n.adjusted()<=40):return n
    except (InvalidOperation,ValueError):pass
    raise ValueError('不是有效数值（需有限值，最多40位有效数字，非零数量级1e-40至1e40）')


def names(value):
    values=value if isinstance(value,list) else str(value or '').splitlines()
    result=[str(v).strip() for v in values if str(v).strip()]
    if len(set(result))!=len(result):raise ValueError('列清单中存在重复名称')
    return result


def method_config(inputs):
    m={k:names(inputs.get(k)) for k in ('required_columns','numeric_columns','id_columns','drop_columns')}
    m.update(version=METHOD_VERSION,label_column=str(inputs.get('label_column') or '').strip(),
        label_map=inputs.get('label_map') or [],features=inputs.get('features') or [],
        duplicate_policy=inputs.get('duplicate_policy','保留并说明'),invalid_policy=inputs.get('invalid_policy','保留空值并说明'))
    return m


def validate_method(m, fields, prediction):
    if m.get('version')!=METHOD_VERSION:raise ValueError('处理方法版本与当前代码不同，请使用原工作流版本或明确建立新方法')
    if m.get('duplicate_policy') not in ['保留并说明','仅保留第一次出现']:raise ValueError('请选择有效的重复处理方式')
    if m.get('invalid_policy') not in ['保留空值并说明','排除有转换问题的记录']:raise ValueError('请选择有效的转换问题处理方式')
    for k in ('required_columns','numeric_columns','id_columns','drop_columns'):
        if not isinstance(m.get(k),list) or any(not isinstance(v,str) or not v for v in m[k]) or len(set(m[k]))!=len(m[k]):raise ValueError('处理方法的列清单无效：'+k)
    label=m.get('label_column')
    if not isinstance(label,str):raise ValueError('处理方法缺少标签列配置')
    if label and label in m['numeric_columns']:raise ValueError('标签列不应再做数值转换；请明确填写标签对应关系，避免改变标签含义')
    if label and label in m['drop_columns']:raise ValueError('标签不应同时列为排除列；新数据模式会自动移除标签')
    mapping=m.get('label_map');features=m.get('features')
    if not isinstance(mapping,list) or len(mapping)>200 or not isinstance(features,list) or len(features)>50:raise ValueError('标签对应关系最多200行，特征最多50行')
    if mapping and not label:raise ValueError('填写标签对应关系前请指定标签列')
    seen=set()
    for item in mapping:
        if not isinstance(item,dict) or any(item.get(k) is None or not str(item[k]).strip() for k in ['from','to']):raise ValueError('标签对应关系的原值和新值都需要填写')
        key=str(item['from']).strip()
        if key in seen:raise ValueError('标签原值重复，请保留唯一对应关系：'+key)
        seen.add(key)
    available=set(fields)-set(m['drop_columns'])-({label} if label else set())
    for i,f in enumerate(features,1):
        if not isinstance(f,dict):raise ValueError(f'特征第{i}行不是配置')
        out=str(f.get('output') or '').strip();left=str(f.get('left') or '').strip();right=str(f.get('right') or '').strip();op=f.get('operation')
        if not out or out in set(fields)|available or out.startswith('preparation_'):raise ValueError(f'特征第{i}行请使用新的、不重名的输出列')
        if op not in OPS:raise ValueError(f'特征第{i}行请选择计算方式')
        if left not in available:raise ValueError(f'特征第{i}行第一列不存在、被排除或是标签：{left}')
        constant=str(f.get('constant') if f.get('constant') is not None else '').strip()
        if op in ['乘固定数','加固定数']:
            if right:raise ValueError(f'特征第{i}行固定数计算不填写第二列')
            try:number(constant)
            except ValueError as exc:raise ValueError(f'特征第{i}行固定数{exc}') from None
        elif right not in available or constant:raise ValueError(f'特征第{i}行需要有效第二列，不填固定数；标签和排除列不能作为特征')
        f.update(output=out,left=left,right=right,constant=constant)
        available.add(out)
    required=set(m['numeric_columns']+m['id_columns']) | (set(m['required_columns'])-({label} if prediction else set()))
    if label and not prediction:required.add(label)
    missing=required-set(fields)-{f['output'] for f in features}
    if missing:raise ValueError('缺少处理方法要求的列：'+'、'.join(sorted(missing)))
    if set(m['numeric_columns']+m['id_columns'])-set(fields):raise ValueError('数值列和标识列应来自原表，不是生成的特征')
    if set(m['required_columns']) & set(m['drop_columns']):raise ValueError('同一列不能同时要求必填又从输出排除')


def prepare(inputs):
    path=source(inputs['source_path']);fields,rows=table(path,inputs.get('sheet') or '')
    purpose=inputs.get('purpose','有标签样本或一般分析')
    if purpose not in ['有标签样本或一般分析','新数据预测前处理']:raise ValueError('请选择本次资料用途')
    previous=str(inputs.get('method_path') or '').strip()
    expected_processor=''
    if previous:
        try:
            saved=json.loads(source(previous).read_text(encoding='utf-8'));m=saved['method']
            if not isinstance(m,dict):raise TypeError()
        except (json.JSONDecodeError,KeyError,TypeError):raise ValueError('请选择本流程保存的method.json处理方法文件') from None
        if digest(m)!=saved.get('method_sha256'):raise ValueError('处理方法文件已改变，请选择原文件或在表单中建立新方法')
        expected_processor=saved.get('processor_sha256')
        if not expected_processor:raise ValueError('旧方法未记录实际处理代码，请使用原工作流重新生成方法')
    else:m=method_config(inputs)
    validate_method(m,fields,purpose=='新数据预测前处理')
    payload=dict(fields=fields,rows=rows,method=m,purpose=purpose,method_from=previous,expected_processor=expected_processor,
        source={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':inputs.get('sheet') or ''})
    folder=Path('results')/('sample-input-'+uuid4().hex);folder.mkdir(parents=True)
    snapshot=folder/'input.json';snapshot.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    return {'snapshot_path':str(snapshot),'sha256':hashlib.sha256(snapshot.read_bytes()).hexdigest(),'rows':len(rows)}


def calculate(row, f):
    try:a=number(row[f['left']]);b=number(f['constant'] if f['operation'] in ['乘固定数','加固定数'] else row[f['right']])
    except ValueError as exc:raise ValueError('计算所需字段缺值或'+str(exc)) from None
    with localcontext() as ctx:
        ctx.prec=40
        op=f['operation']
        if op=='第一列除以第二列' and b==0:raise ValueError('分母为0，未用极小数代替')
        value=a+b if op in ['两列相加','加固定数'] else a-b if op=='第一列减第二列' else abs(a-b) if op=='两列差的绝对值' else a/b if op=='第一列除以第二列' else a*b
        number(value)
        return str(value)


def transform(inputs):
    prepared=inputs['prepared'];p=source(prepared['snapshot_path'],100_000_000)
    if hashlib.sha256(p.read_bytes()).hexdigest()!=prepared['sha256']:raise ValueError('输入快照已改变，请重新读取资料')
    data=json.loads(p.read_text(encoding='utf-8'));m=data['method'];fields=data['fields'];prediction=data['purpose']=='新数据预测前处理'
    actual_processor=processor_hash()
    if data.get('expected_processor') and actual_processor!=data['expected_processor']:
        raise ValueError('处理代码与保存方法时不同，请使用原工作流版本；若确实要改方法，清空沿用文件并重新配置')
    validate_method(m,fields,prediction)
    label=m['label_column'];mapping={str(x['from']).strip():str(x['to']).strip() for x in m['label_map']}
    output_fields=[f for f in fields if f not in m['drop_columns'] and not (prediction and f==label)]+[f['output'] for f in m['features']]
    if not output_fields:raise ValueError('排除后没有可输出字段，请保留数据列')
    original_groups=defaultdict(list);ids=defaultdict(list)
    for n,row in data['rows']:
        if not any(v.strip() for v in row.values()):continue
        original_groups[tuple(row[f] for f in fields)].append(n)
        key=tuple(row[f].strip() for f in m['id_columns'])
        if key and all(key):ids[key].append((n,tuple(row[f] for f in fields)))
    conflicts={}
    for entries in ids.values():
        if len({v for _,v in entries})>1:
            records=[r[0] for r in entries]
            conflicts.update((n,records) for n,_ in entries)
    issues=[];status=[];accepted=[];excluded=[];reasons_count=Counter();label_before=Counter();label_after=Counter();first={};warning_rows=0
    for n,original in data['rows']:
        row=dict(original);reasons=[];warnings=[];bad=False
        def issue(field,reason,before='',after='',action='说明'):
            issues.append(dict(source_record=n,field=field,reason=reason,original_value=before,prepared_value=after,action=action))
        def exclude(reason):
            if reason not in reasons:reasons.append(reason);reasons_count[reason]+=1
        if not any(v.strip() for v in original.values()):exclude('全空记录')
        else:
            raw_key=tuple(original[f] for f in fields)
            if raw_key in first:
                text=f'与第{first[raw_key]}条记录的全部原单元格相同'
                warnings.append(text);issue('',text,action='排除' if m['duplicate_policy']=='仅保留第一次出现' else '保留')
                if m['duplicate_policy']=='仅保留第一次出现':exclude('完全相同的额外重复记录')
            else:first[raw_key]=n
            if n in conflicts:
                text='同一标识的内容不同，相关记录：'+','.join(map(str,conflicts[n][:20]))
                warnings.append(text);issue(','.join(m['id_columns']),text,action='保留并待核对')
            if m['id_columns'] and any(not row[f].strip() for f in m['id_columns']):
                warnings.append('样本标识缺值，无法核对其重复');issue(','.join(m['id_columns']),'样本标识缺值，无法核对其重复')
            for field in m['numeric_columns']:
                before=row[field]
                if not before.strip():continue
                try:row[field]=str(number(before))
                except ValueError as exc:
                    row[field]='';bad=True;warnings.append(field+'：'+str(exc));issue(field,str(exc),before,'','置空并说明')
                else:
                    if row[field]!=before:issue(field,'数值格式统一',before,row[field],'转换')
            if label and not prediction:
                raw=row[label].strip();label_before[raw or '（缺失）']+=1
                if not raw:exclude('标签缺失')
                elif mapping and raw not in mapping:exclude('标签未列入对应关系');issue(label,'标签未列入对应关系',row[label],'','排除')
                else:
                    row[label]=mapping.get(raw,raw)
                    if row[label]!=original[label]:issue(label,'按明确对应关系处理标签',original[label],row[label],'转换')
            for f in m['features']:
                try:row[f['output']]=calculate(row,f)
                except ValueError as exc:
                    row[f['output']]='';bad=True;warnings.append(f['output']+'：'+str(exc));issue(f['output'],str(exc),'','','置空并说明')
            for field in m['required_columns']:
                if prediction and field==label:continue
                if not row[field].strip():exclude('必填列缺失：'+field)
            if bad and m['invalid_policy']=='排除有转换问题的记录':exclude('数值转换或特征计算有问题')
        if reasons:
            excluded.append({**original,'preparation_source_record':n,'preparation_reason':'；'.join(reasons)})
            output_record=''
        else:
            accepted.append({f:row[f] for f in output_fields});output_record=len(accepted)+1
            if label and not prediction:label_after[row[label]]+=1
        if warnings:warning_rows+=1
        status.append(dict(source_record=n,output_record=output_record,status='排除' if reasons else '保留（有说明）' if warnings else '保留',reasons='；'.join(reasons),notes='；'.join(warnings)))
    folder=Path('results')/('prepared-samples-'+uuid4().hex);folder.mkdir(parents=True);artifacts=[]
    def save_csv(name,headers,rows,title):
        with (folder/name).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=headers);writer.writeheader();writer.writerows(rows)
        artifacts.append(dict(file_path=str(folder/name),label=title))
    save_csv('samples.csv',output_fields,accepted,'处理后的数据 CSV')
    save_csv('excluded.csv',fields+['preparation_source_record','preparation_reason'],excluded,'排除记录与原始值 CSV')
    save_csv('row-status.csv',['source_record','output_record','status','reasons','notes'],status,'每条记录的去向与原因 CSV')
    save_csv('issues.csv',['source_record','field','reason','original_value','prepared_value','action'],issues,'字段问题与转换明细 CSV')
    saved=dict(method=m,method_sha256=digest(m),processor_sha256=actual_processor,input_fields=fields,output_fields=output_fields)
    (folder/'method.json').write_text(json.dumps(saved,ensure_ascii=False,indent=2),encoding='utf-8');artifacts.append(dict(file_path=str(folder/'method.json'),label='可复用处理方法 JSON'))
    result=dict(source_path=str(folder/'samples.csv'),method_path=str(folder/'method.json'),method_sha256=digest(m),processor_sha256=actual_processor,method=m,method_from=data['method_from'],
        purpose=data['purpose'],source=data['source'],rows=len(data['rows']),retained=len(accepted),excluded=len(excluded),warning_rows=warning_rows,
        exclusion_reasons=dict(reasons_count),conflicting_records=len(conflicts),exact_duplicate_extras=sum(len(v)-1 for v in original_groups.values()),
        labels_before=dict(label_before),labels_after=dict(label_after),output_fields=output_fields,
        preview=[{f:r[f] for f in output_fields[:10]} for r in accepted[:8]],status_preview=status[:12],artifacts=artifacts)
    (folder/'summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');artifacts.append(dict(file_path=str(folder/'summary.json'),label='数量、方法与来源汇总 JSON'))
    return result


def report(result):
    def cell(value):return str(value).replace('\\','\\\\').replace('|','\\|').replace('\n',' ').replace('\r',' ')[:200]
    m=result['method']
    text=f"# 样本清洗与逐行特征准备\n\n输入 {result['rows']} 条记录，保留 {result['retained']} 条，排除 {result['excluded']} 条；{result['warning_rows']} 条有需要了解的说明。原始资料保持不变。\n\n"
    if not result['retained']:text+='没有剩余样本，当前结果仅供检查；请根据排除原因修正资料或方法，再进行后续任务。\n\n'
    text+='## 本次如何处理\n\n'+('沿用已保存的处理方法，其余方法表单参数未使用。' if result['method_from'] else '使用本次表单中明确配置的方法。')+'用途：'+result['purpose']+'。\n\n'
    text+='- 完全相同的重复记录：'+m['duplicate_policy']+'；可选数值/特征问题：'+m['invalid_policy']+'。\n'
    text+='- 必填列：'+('、'.join(map(cell,m['required_columns'])) or '未设置')+'；数值列：'+('、'.join(map(cell,m['numeric_columns'])) or '未设置')+'。\n'
    text+='- 排除输出列：'+('、'.join(map(cell,m['drop_columns'])) or '未设置')+'。被排除字段和标签不用于新特征。\n'
    if result['purpose']=='新数据预测前处理':text+='- 本次不要求标签存在，输出移除标签；其他检查和特征沿用相同方法。\n'
    elif m['label_column']:
        text+='- 标签按明确对应关系处理，未列出的原值不会强转整数或猜测。保留样本的标签数量：'+('、'.join(cell(k)+'：'+str(v) for k,v in result['labels_after'].items()) or '无')+'。数量很少的类别需要单独判断可靠性。\n'
    text+='\n## 排除和需要核对的事项\n\n'
    for reason,count in result['exclusion_reasons'].items():text+='- '+cell(reason)+'：'+str(count)+' 条。\n'
    if not result['exclusion_reasons']:text+='本次没有排除记录。\n'
    text+=f"\n完全相同的额外记录 {result['exact_duplicate_extras']} 条；同一标识内容不同涉及 {result['conflicting_records']} 条。冲突不擅自选第一条，需核对样本单位及标签来源。一个记录可能有多个排除原因，原因数量不能直接相加。\n\n"
    text+='## 特征怎么算\n\n'
    if not m['features']:text+='未添加新特征。\n'
    for i,f in enumerate(m['features'],1):text+=f"- {i}. {cell(f['output'])}：{cell(f['left'])}，{f['operation']}，{cell(f['constant'] if f['operation'] in ['乘固定数','加固定数'] else f['right'])}。\n"
    text+='\n公式按顺序逐行计算，保留40位十进制计算精度；分母0或无效值返回空值和原因，不添加极小数制造有效结果。没有在整表估计均值、标准差或筛选特征；这些需要拟合的操作继续交给训练折。\n\n'
    text+='## 记录去向预览\n\n| 输入记录号 | 输出记录号 | 去向 | 原因与说明 |\n| --- | --- | --- | --- |\n'
    for r in result['status_preview']:text+='| '+' | '.join(map(cell,[r['source_record'],r['output_record'],r['status'],'；'.join(v for v in [r['reasons'],r['notes']] if v)]))+' |\n'
    if result['preview']:
        headers=list(result['preview'][0]);text+='\n## 处理后数据预览\n\n| '+' | '.join(map(cell,headers))+' |\n| '+' | '.join('---' for _ in headers)+' |\n'
        for r in result['preview']:text+='| '+' | '.join(cell(r[k]) for k in headers)+' |\n'
    text+='\n预览最多12条去向、8条数据和10个字段；完整结果请下载。输出数据不包含去向、问题说明等辅助字段，避免被自动当作模型特征。\n\n## 下一步\n\n可以将处理后的数据交给分析或训练流程；设备/批次/时间隔离、测试集及预测时点仍由任务决定。使用已有模型前，先选择训练时保存的method.json处理新资料，并选择新数据模式；模型服务不会自动执行本流程。改方法要清空沿用文件并明确修改配置，不能将旧模型与新字段语义混用。\n'
    path=Path(result['source_path']).parent/'report.md';path.write_text(text,encoding='utf-8')
    return {**result,'markdown':text,'artifacts':result['artifacts']+[dict(file_path=str(path),label='清洗与特征说明 Markdown')]}


def main(inputs):
    if inputs['operation']=='prepare':return prepare(inputs)
    if inputs['operation']=='transform':return transform(inputs)
    if inputs['operation']=='report':return report(inputs['result'])
    raise ValueError('未知的样本处理步骤')

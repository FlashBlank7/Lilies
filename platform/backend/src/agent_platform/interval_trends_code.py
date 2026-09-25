"""Portable period-mean labels. Future measurements never become input features."""
import hashlib
import io
import json
import csv
import re
from datetime import date, datetime, timezone
from fractions import Fraction
from pathlib import Path
from uuid import uuid4

class Period:
    def __init__(self, ordinal, freq):
        self.ordinal, self.freq = ordinal, freq

    def start(self):
        if self.freq == 'D':
            return date.fromordinal(self.ordinal)
        year, offset = divmod(self.ordinal, 12 if self.freq == 'M' else 4)
        return date(year, offset + 1 if self.freq == 'M' else offset * 3 + 1, 1)

    def __str__(self):
        d = self.start()
        return d.isoformat() if self.freq == 'D' else (f'{d.year:04d}-{d.month:02d}' if self.freq == 'M'
                                                      else f'{d.year:04d}Q{(d.month - 1) // 3 + 1}')


def main(inputs):
    path = Path(str(inputs.get('source_path') or ''))
    if (path.is_absolute() or '..' in path.parts or not path.parts
            or path.parts[0] not in {'requirement-package', 'results'} or not path.is_file()
            or any(p.is_symlink() for p in [path, *path.parents])
            or not path.resolve().is_relative_to(Path.cwd().resolve())):
        raise ValueError('请选择本项目内的历史数值表')
    if path.stat().st_size > 50_000_000:
        raise ValueError('历史表超过50 MB，请先按对象或时间范围拆分')
    raw = path.read_bytes()
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = str(inputs.get('sheet') or '')
            if not sheet and len(book.sheetnames) != 1:
                raise ValueError('Excel有多个工作表，请指定工作表')
            if sheet and sheet not in book.sheetnames:
                raise ValueError('所选Excel工作表不存在')
            records = [[str(v) if v is not None else '' for v in row]
                       for row in book[sheet or book.sheetnames[0]].iter_rows(values_only=True)]
        finally:
            book.close()
    elif path.suffix.lower() in {'.csv', '.tsv'}:
        records = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), delimiter='\t' if path.suffix.lower() == '.tsv' else ','))
    else:
        raise ValueError('请选择CSV、TSV或XLSX表格')
    if not 2 <= len(records) <= 20001:
        raise ValueError('需要1至20000条历史记录')
    fields = [v.strip() for v in records[0]]
    if not all(fields) or len(fields) != len(set(fields)):
        raise ValueError('表头不能为空或重复')
    if any(len(row) != len(fields) for row in records[1:]):
        raise ValueError('记录的列数与表头不一致，请检查分隔符或空行')
    frame = [dict(zip(fields, row)) for row in records[1:]]
    period_col = str(inputs.get('period_column') or 'period')
    value_col = str(inputs.get('value_column') or 'value')
    series_col = str(inputs.get('series_column') or '')
    available_col = str(inputs.get('available_column') or '')
    for column in [period_col, value_col, series_col, available_col]:
        if column and column not in fields:
            raise ValueError('缺少所选列：' + column)
    frequency = {'日': 'D', '月': 'M', '季': 'Q-DEC'}.get(inputs.get('frequency', '月'))
    if not frequency:
        raise ValueError('请选择日、月或季的原表周期')

    def integer(name, default, maximum):
        try:
            n = float(inputs.get(name, default))
            if n.is_integer() and 1 <= n <= maximum:
                return int(n)
        except (ValueError, TypeError):
            pass
        label = {'history_periods': '历史参考周期数', 'window_periods': '每个未来区间的周期数',
                 'windows': '未来区间数量', 'origin_stride': '样本起点间隔'}[name]
        raise ValueError(label + '需要1至' + str(maximum) + '之间的整数')

    history = integer('history_periods', 3, 60)
    width = integer('window_periods', 1, 60)
    windows = integer('windows', 3, 12)
    stride = integer('origin_stride', 1, 60)
    try:
        threshold = Fraction(str(inputs.get('threshold', 0)))
        if not 0 <= threshold <= 10 ** 40:
            raise ValueError()
    except (TypeError, ValueError, ZeroDivisionError):
        raise ValueError('平稳阈值需要非负有限数值，单位与原数值相同')

    def period(value, label):
        try:
            value = str(value).strip()
            if frequency == 'Q-DEC' and re.fullmatch(r'\d{4}Q[1-4]', value):
                d = date(int(value[:4]), (int(value[-1]) - 1) * 3 + 1, 1)
            else:
                d = date.fromisoformat(value + '-01' if re.fullmatch(r'\d{4}-\d{2}', value) else value[:10])
            ordinal = d.toordinal() if frequency == 'D' else d.year * (12 if frequency == 'M' else 4) + (d.month - 1 if frequency == 'M' else (d.month - 1) // 3)
            return Period(ordinal, frequency)
        except (ValueError, TypeError):
            raise ValueError(label + '不是有效的周期；例如2025-01或2025Q1')

    def end(p):
        d = Period(p.ordinal + 1, frequency).start()
        return datetime(d.year, d.month, d.day, tzinfo=timezone.utc)

    groups, seen = {}, set()
    for i, row in enumerate(frame):
        record = int(i) + 2
        p = period(row[period_col], f'记录{record}的{period_col}')
        name = str(row[series_col]).strip() if series_col else '全部'
        if not name:
            raise ValueError(f'记录{record}缺少序列标识')
        try:
            available = datetime.fromisoformat(row[available_col].strip()) if available_col else end(p)
            available = available.replace(tzinfo=timezone.utc) if available.tzinfo is None else available.astimezone(timezone.utc)
        except (ValueError, TypeError):
            raise ValueError(f'记录{record}的实际可用时刻无效')
        if available < end(p):
            raise ValueError(f'记录{record}在周期结束前就声明完整周期数值可用，请核对周期或发布时间')
        key = (name, p.ordinal, available)
        if key in seen:
            raise ValueError(f'记录{record}在同序列、周期及可用时刻重复，请先明确汇总或修订关系')
        seen.add(key)
        try:
            value = Fraction(str(row[value_col]))
            if abs(value) > 10 ** 40:
                value = None
        except (ValueError, TypeError, ZeroDivisionError):
            value = None
        groups.setdefault(name, {}).setdefault(p.ordinal, []).append(
            {'value': value, 'available': available, 'record': record, 'period': p})
    for values in groups.values():
        for versions in values.values():
            versions.sort(key=lambda v: v['available'])

    first = period(inputs['first_origin'], '首个起点').ordinal if inputs.get('first_origin') else None
    last = period(inputs['last_origin'], '最后起点').ordinal if inputs.get('last_origin') else None
    if first is not None and last is not None and first > last:
        raise ValueError('首个起点不能晚于最后起点')
    samples, labels, excluded, prediction = [], [], [], []
    origins_count = 0
    for name, values in sorted(groups.items()):
        start = max(min(values) + history - 1, first if first is not None else min(values))
        stop = min(max(values), last if last is not None else max(values))
        if stop - start > 20000:
            raise ValueError('时间范围过长，请核对周期并缩小起点范围')
        last_selected = start + (stop - start) // stride * stride
        for origin in range(start, stop + 1, stride):
            origins_count += 1
            if origins_count * windows > 20000:
                raise ValueError('输出样本超过20000条，请增加起点间隔或缩小时间范围')
            p = Period(ordinal=origin, freq=frequency)
            origin_time = end(p)

            def collect(ordinals, cutoff=None):
                selected = []
                for ordinal in ordinals:
                    versions = [v for v in values.get(ordinal, []) if cutoff is None or v['available'] <= cutoff]
                    if not versions or versions[-1]['value'] is None:
                        return None, str(Period(ordinal=ordinal, freq=frequency))
                    selected.append(versions[-1])
                return selected, ''

            past, missing = collect(range(origin - history + 1, origin + 1), origin_time)
            if past is None:
                excluded.append(dict(series=name, origin=str(p), interval='', reason='历史窗口缺失或当时尚不可用：' + missing))
                continue
            base = dict(series=name, origin_time=origin_time.isoformat(),
                        **{f'lag_{i}': float(past[-1 - i]['value']) for i in range(history)})
            reference = past
            for number in range(1, windows + 1):
                ident = hashlib.sha256(f'{name}\0{p}\0{number}'.encode()).hexdigest()[:24]
                inputs_row = dict(sample_id=ident, **base, interval=number)
                if origin == last_selected:
                    prediction.append(inputs_row)
                begin = origin + (number - 1) * width + 1
                future, missing = collect(range(begin, begin + width))
                if reference is None or future is None:
                    excluded.append(dict(series=name, origin=str(p), interval=number,
                                         reason='未来标签区间或前一区间不完整：' + (missing or '前一区间')))
                    reference = future
                    continue
                reference_mean = sum(v['value'] for v in reference) / len(reference)
                future_mean = sum(v['value'] for v in future) / len(future)
                delta = future_mean - reference_mean
                target = '上涨' if delta > threshold else '下跌' if delta < -threshold else '平稳'
                label_time = max(v['available'] for v in past + reference + future)
                samples.append(dict(**inputs_row, label_available_time=label_time.isoformat(), target=target))
                labels.append(dict(sample_id=ident, series=name, origin=str(p), interval=number,
                                   window_start=str(Period(ordinal=begin, freq=frequency)),
                                   window_end=str(Period(ordinal=begin + width - 1, freq=frequency)),
                                   reference_mean=float(reference_mean), future_mean=float(future_mean), change=float(delta),
                                   threshold=float(threshold), target=target, label_available_time=label_time.isoformat(),
                                   feature_records=json.dumps([v['record'] for v in past]),
                                   reference_records=json.dumps([v['record'] for v in reference]),
                                   future_records=json.dumps([v['record'] for v in future])))
                reference = future
    if not origins_count:
        raise ValueError('所选起点范围没有足够历史周期，请调整起点或历史窗口')
    if not samples and not prediction:
        raise ValueError('没有可用历史窗口，请核对缺失周期、发布日期或历史窗口长度')
    output = Path('results') / ('interval-trends-' + str(uuid4()))
    output.mkdir(parents=True)
    snapshot = output / ('source' + path.suffix.lower())
    snapshot.write_bytes(raw)
    feature_fields = ['sample_id', 'series', 'origin_time', *[f'lag_{i}' for i in range(history)], 'interval']
    artifacts = [dict(label='原输入快照', file_path=str(snapshot))]

    def save(name, rows, fields=None):
        target = output / name
        with target.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        label = {'samples.csv': '有标签的历史样本', 'prediction-inputs.csv': '最后起点的模型输入',
                 'label-origins.csv': '每个标签的计算与原记录', 'excluded.csv': '未形成样本的具体原因'}[name]
        artifacts.append(dict(label=label, file_path=str(target)))
        return str(target)

    source_path = save('samples.csv', samples, feature_fields + ['label_available_time', 'target'])
    prediction_path = save('prediction-inputs.csv', prediction, feature_fields)
    save('label-origins.csv', labels, ['sample_id', 'series', 'origin', 'interval', 'window_start', 'window_end',
         'reference_mean', 'future_mean', 'change', 'threshold', 'target', 'label_available_time',
         'feature_records', 'reference_records', 'future_records'])
    save('excluded.csv', excluded, ['series', 'origin', 'interval', 'reason'])
    config = {k: v for k, v in inputs.items() if k != 'source_path'}
    class_counts = {label: sum(row['target'] == label for row in samples) for label in ['上涨', '平稳', '下跌']}
    manifest = dict(source_path=str(path), source_snapshot=str(snapshot), sha256=hashlib.sha256(raw).hexdigest(), config=config,
                    rows=len(samples), prediction_rows=len(prediction), excluded=len(excluded),
                    class_counts=class_counts,
                    features=[f'lag_{i}' for i in range(history)] + ['interval'],
                    mapping={'target': 'target', 'id_column': 'sample_id', 'prediction_time_column': 'origin_time',
                             'label_available_time_column': 'label_available_time'},
                    rule='首段与起点当时可知的历史均值比；后续段与前段比，绝对变化严格超过阈值才算涨跌。')
    (output / 'method.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    artifacts.append(dict(label='方法与输入版本', file_path=str(output / 'method.json')))
    markdown = (f'# 区间趋势样本\n\n得到 **{len(samples)}** 条有标签样本、**{len(prediction)}** 条最后起点输入；'
                f'有 **{len(excluded)}** 个窗口未形成标签或输入，原因见排除表。\n\n'
                + '标签数量：' + '，'.join(f'{label} {count} 条' for label, count in class_counts.items()) + '。\n\n'
                +
                '首段对比历史均值，之后每段对比前段均值。等于阈值算平稳。标签是历史实测构造，不是模型预测。\n\n'
                '样本特征只有历史滞后值与区间编号；未来均值、原记录号放在单独来源表，不得加入训练特征。'
                '后续训练必须按起点时间划分，并用label_available_time排除验证起点尚未知的标签。'
                '不同区间、重叠起点不是独立样本。这里只制备资料，没有训练、选择阈值或证明长期预测有效。\n\n'
                '日/月/季按UTC日历解释；每周期一条完整周期值，可有不同公布时刻的修订。'
                '未填可用时刻时假定周期结束即可使用，不能证明实际发布时间正确。\n\n' +
                '\n'.join(f'- [{a["label"]}]({a["file_path"]})' for a in artifacts))
    (output / 'report.md').write_text(markdown, encoding='utf-8')
    artifacts.append(dict(label='分析说明', file_path=str(output / 'report.md')))
    return dict(source_path=source_path, prediction_path=prediction_path, rows=len(samples),
                prediction_rows=len(prediction), excluded=len(excluded), class_counts=class_counts, features=manifest['features'],
                mapping=manifest['mapping'], artifacts=artifacts, markdown=markdown,
                source={'path': str(path), 'sha256': manifest['sha256']})

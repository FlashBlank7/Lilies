"""Sequential allocation of existing ranked alternatives under shared quantity limits."""
import csv
import hashlib
import json
from decimal import Decimal, InvalidOperation
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value):
    p = Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ('requirement-package', 'results'):
        raise ValueError('请选择当前项目的资料或结果文件')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p, *p.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size > 10_000_000:
        raise ValueError('文件不存在或超过10 MB')
    return p


def table(value, sheet, required, limit):
    p = source(value)
    if p.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(p, read_only=True, data_only=True)
        try:
            if not sheet and len(book.sheetnames) > 1:
                raise ValueError('Excel有多张表，请指定工作表：' + '、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:
                raise ValueError('工作表不存在：' + sheet)
            rows = list(islice((book[sheet] if sheet else book.active).values, limit+2))
        finally:
            book.close()
    elif p.suffix.lower() in ('.csv', '.tsv'):
        with p.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(islice(csv.reader(stream, delimiter='\t' if p.suffix.lower() == '.tsv' else ','), limit+2))
    else:
        raise ValueError('请选择CSV、TSV或XLSX')
    if not rows or len(rows) > limit+1:
        raise ValueError(f'{p.name} 需要表头，最多{limit}条记录')
    fields = ['' if v is None else str(v).strip() for v in rows[0]]
    if any(not v for v in fields) or len(fields) != len(set(fields)):
        raise ValueError('表头不能为空或重名')
    if set(required)-set(fields):
        raise ValueError(p.name + ' 缺少字段：' + '、'.join(sorted(set(required)-set(fields))))
    data = []
    for record, row in enumerate(rows[1:], 2):
        if all(v is None or not str(v).strip() for v in row):
            continue
        if len(row) != len(fields):
            raise ValueError(f'{p.name} 第{record}行列数与表头不同')
        data.append((record, dict(zip(fields, ['' if v is None else str(v).strip() for v in row]))))
    if not data:
        raise ValueError(p.name + ' 没有数据记录')
    return p, data


def number(value, label, signed=False):
    try:
        n = Decimal(str(value).strip())
        if n.is_finite() and (-Decimal('1e12') if signed else 0) <= n <= Decimal('1e12') and n.as_tuple().exponent >= -6:
            return n
    except (InvalidOperation, ValueError):
        pass
    raise ValueError(label + ' 需要有限数值，最多6位小数且绝对值不超过1万亿' + ('' if signed else '，不能为负'))


def prepare(inputs):
    defaults = {'id_column':'candidate_id','group_column':'group_id','rank_column':'comparison_rank',
                'resource_column':'resource_id','quantity_column':'quantity','capacity_column':'capacity'}
    cfg = {k: str(inputs.get(k, default) or '').strip() for k, default in defaults.items()}
    if any(not v for v in cfg.values()):
        raise ValueError('请填写方案、分组、顺位、资源、用量和上限字段')
    if len({cfg[k] for k in ('id_column','group_column','rank_column')}) != 3 or len({cfg[k] for k in ('id_column','resource_column','quantity_column')}) != 3 or cfg['resource_column'] == cfg['capacity_column']:
        raise ValueError('同一张表的配置字段不能重名')
    cfg['group_order_column'] = str(inputs.get('group_order_column') or '').strip()
    cfg['group_order_direction'] = inputs.get('group_order_direction') or '从小到大'
    if cfg['group_order_direction'] not in ('从小到大','从大到小'):
        raise ValueError('请选择分组处理方向')
    sources, tables = {}, {}
    for kind, key, fields, limit in [
        ('candidates','source_path',[cfg['id_column'],cfg['group_column'],cfg['rank_column']] + ([cfg['group_order_column']] if cfg['group_order_column'] else []),10000),
        ('usage','usage_path',[cfg['id_column'],cfg['resource_column'],cfg['quantity_column']],60000),
        ('limits','limits_path',[cfg['resource_column'],cfg['capacity_column']],10000),
    ]:
        path, tables[kind] = table(inputs.get(key), str(inputs.get(kind+'_sheet') or ''), fields, limit)
        sources[kind] = {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheet': inputs.get(kind+'_sheet') or ''}
    capacities, candidates, usages, group_priorities = {}, [], {}, {}
    for rownum, row in tables['limits']:
        identity = row[cfg['resource_column']]
        if not identity or identity in capacities:
            raise ValueError(f'上限表第{rownum}行资源标识为空或重复')
        capacities[identity] = str(number(row[cfg['capacity_column']], f'上限表第{rownum}行数量'))
    seen = set()
    for rownum, row in tables['candidates']:
        identity, group, rank = row[cfg['id_column']], row[cfg['group_column']], row[cfg['rank_column']]
        if not identity or identity in seen or not group:
            raise ValueError(f'候选表第{rownum}行方案标识为空/重复，或分组为空')
        seen.add(identity)
        if rank:
            n = number(rank, f'候选表第{rownum}行顺位')
            if n < 1 or n != int(n):
                raise ValueError(f'候选表第{rownum}行顺位需要正整数；空顺位表示不参与分配')
            rank = int(n)
        status = row.get('comparison_status', '')
        if status and status not in ('pass','unchecked','fail','unknown'):
            raise ValueError(f'候选表第{rownum}行comparison_status无效，请核对比较结果')
        priority = str(number(row[cfg['group_order_column']], f'候选表第{rownum}行分组顺序', signed=True)) if cfg['group_order_column'] else str(len(group_priorities))
        if group not in group_priorities:
            group_priorities[group] = priority
        elif cfg['group_order_column'] and Decimal(group_priorities[group]) != Decimal(priority):
            raise ValueError(f'候选表第{rownum}行，同一分组的处理顺序数值不一致')
        candidates.append({'candidate_id':identity,'group':group,'rank':rank or None,'source_row':rownum,
                           'comparison_status':status or 'unchecked','comparison_reason':row.get('comparison_reason','')})
        usages[identity] = {}
    for rownum, row in tables['usage']:
        identity, resource = row[cfg['id_column']], row[cfg['resource_column']]
        if identity not in usages:
            raise ValueError(f'用量表第{rownum}行方案不在本次候选中：{identity}')
        if resource not in capacities:
            raise ValueError(f'用量表第{rownum}行资源缺少明确上限：{resource}')
        if resource in usages[identity]:
            raise ValueError(f'用量表第{rownum}行方案和资源重复，请先核对并明确汇总')
        usages[identity][resource] = str(number(row[cfg['quantity_column']], f'用量表第{rownum}行用量'))
    for candidate in candidates:
        if not usages[candidate['candidate_id']]:
            raise ValueError('方案缺少用量记录：' + candidate['candidate_id'] + '；明确不消耗时也需填写0')
    data = {'config':cfg,'sources':sources,'candidates':candidates,'capacities':capacities,'usage':usages,'group_priorities':group_priorities}
    folder = Path('results') / ('allocation-input-'+uuid4().hex); folder.mkdir(parents=True)
    path = folder/'input.json'; path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
    return {'snapshot_path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}


def allocate(inputs):
    prepared = inputs['prepared']; path = source(prepared['snapshot_path'])
    if hashlib.sha256(path.read_bytes()).hexdigest() != prepared['sha256']:
        raise ValueError('输入快照已改变，请重新读取资料')
    data = json.loads(path.read_text(encoding='utf-8')); cfg = data['config']
    remaining = {k:Decimal(v) for k,v in data['capacities'].items()}
    groups = sorted(data['group_priorities'],key=lambda g:Decimal(data['group_priorities'][g]),reverse=cfg['group_order_direction']=='从大到小' and bool(cfg['group_order_column']))
    by_group = {group:[] for group in groups}
    for candidate in data['candidates']:
        by_group[candidate['group']].append(candidate)
    decisions, consumption, summary = [], [], []
    for group_index, group in enumerate(groups,1):
        selected = None
        group_decisions = []
        for candidate in sorted(by_group[group],key=lambda c:(c['rank'] is None,c['rank'] or 0,c['source_row'])):
            identity = candidate['candidate_id']; quantities = {r:Decimal(v) for r,v in data['usage'][identity].items()}
            if candidate['comparison_status'] in ('fail','unknown'):
                state, reason = '排除','前序条件未通过或未知；' + candidate['comparison_reason']
            elif candidate['rank'] is None:
                state, reason = '排除','没有明确的比较顺位，请先比较或补齐偏好'
            elif selected:
                state, reason = '未选','同组已选择方案：' + selected
            else:
                shortages = [f'{r}需{q:f}，当前剩余{remaining[r]:f}' for r,q in quantities.items() if q > remaining[r]]
                if shortages:
                    state, reason = '数量不足','；'.join(shortages)
                else:
                    state, reason, selected = '已选','按本组顺位及当前剩余数量选择',identity
                    for resource, quantity in quantities.items():
                        before = remaining[resource]; remaining[resource] -= quantity
                        consumption.append({'group_order':group_index,'group':group,'candidate_id':identity,'resource_id':resource,'quantity':format(quantity,'f'),
                                            'before':format(before,'f'),'after':format(remaining[resource],'f')})
            decision = {**candidate,'group_order':group_index,'allocation_status':state,'allocation_reason':reason}
            decisions.append(decision); group_decisions.append(decision)
        summary.append({'group_order':group_index,'group':group,'selected_candidate':selected or '',
                        'status':'已选择' if selected else '没有可选方案',
                        'reason':('当前数量足够；按组内顺位选择' if selected else '；'.join(d['candidate_id']+'：'+d['allocation_reason'] for d in group_decisions[:3]))})
    balances = [{'resource_id':r,'capacity':v,'used':format(Decimal(v)-remaining[r],'f'),'remaining':format(remaining[r],'f')} for r,v in data['capacities'].items()]
    folder = Path('results')/('candidate-allocation-'+uuid4().hex);folder.mkdir(parents=True);artifacts=[]
    def save_csv(name,fields,rows,label):
        with (folder/name).open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader();writer.writerows(rows)
        artifacts.append({'file_path':str(folder/name),'label':label})
    save_csv('decisions.csv',list(decisions[0]),decisions,'全部候选的选择与原因 CSV')
    save_csv('groups.csv',['group_order','group','selected_candidate','status','reason'],summary,'各组选择 CSV')
    save_csv('consumption.csv',['group_order','group','candidate_id','resource_id','quantity','before','after'],consumption,'本次逐项数量分配 CSV')
    save_csv('balances.csv',['resource_id','capacity','used','remaining'],balances,'共同数量与剩余 CSV')
    result={'groups':len(summary),'selected_groups':sum(bool(r['selected_candidate']) for r in summary),'candidates':len(decisions),
            'source_path':str(folder/'decisions.csv'),'groups_path':str(folder/'groups.csv'),'balances_path':str(folder/'balances.csv'),
            'config':cfg,'sources':data['sources'],'selections':summary}
    def md(value):return str(value).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('|','\\|').replace('\n',' ').replace('\r',' ')
    order = ('按'+cfg['group_order_column']+' '+cfg['group_order_direction']) if cfg['group_order_column'] else '按各组在候选表首次出现的顺序'
    lines=['# 候选按共享数量分配','',f'本次{len(decisions)}个候选、{len(summary)}组，{result["selected_groups"]}组选择了方案。原文件及库存未修改。',
           '',f'处理顺序：{md(order)}。每组按顺位从小到大选择第一个满足剩余数量的方案，每组最多一个；并列顺位按原表行顺序。分组顺序相同时同样按首次出现顺序，不将这个打破并列的方法当业务最优。',
           '', '## 每组结果（前30组，完整记录见下载）','', '| 顺序 | 分组 | 已选方案 | 结果与原因 |','|---:|---|---|---|']
    lines.extend('| '+' | '.join(md(v) for v in (r['group_order'],r['group'],r['selected_candidate'],r['status']+'：'+r['reason']))+' |' for r in summary[:30])
    lines += ['', '## 数量核对（前30项）','', '| 项目 | 原上限 | 本次分配 | 剩余 |','|---|---:|---:|---:|']
    lines.extend('| '+' | '.join(md(r[k]) for k in ('resource_id','capacity','used','remaining'))+' |' for r in balances[:30])
    lines += ['', '## 怎样理解结果','',
              '数量按同一资源标识共同扣减，超出剩余数量的方案保留原因；前序fail/unknown或无顺位的候选不选。输入带有unchecked或没有comparison_status时，仅说明给定顺位，没有检查其他业务限制。',
              '这是一种顺序贪心分配，处理顺序可能改变结果，可能没有充分利用资源或满足最多组；没有可选项不证明全局无解。不回退、更换早先选择，也不重新生成候选。',
              '数量既可以是库存上限，也可以是尚可满足的需求；同一资源必须使用相同单位，不能混用。这里只检查总量，不检查时间冲突、设备、配对、长度、现场锁定或预留事务，不构成完整生产排程。',
              '更换上限、顺序或前序比较结果可重新分配，复用原候选和用量；不会训练或重新预测。每次新建结果，旧报告保持。实际采用方案与写入业务系统仍需明确的后续任务。']
    result['markdown']='\n'.join(lines)
    (folder/'report.md').write_text(result['markdown'],encoding='utf-8')
    (folder/'summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    artifacts += [{'file_path':str(folder/'report.md'),'label':'分配说明 Markdown'},
                  {'file_path':str(folder/'summary.json'),'label':'本次配置与结果 JSON'},
                  {'file_path':str(path),'label':'本次输入快照 JSON'}]
    result['artifacts']=artifacts
    return result


def main(inputs):
    if inputs.get('operation')=='prepare':return prepare(inputs)
    if inputs.get('operation')=='allocate':return allocate(inputs)
    raise ValueError('未知处理步骤')

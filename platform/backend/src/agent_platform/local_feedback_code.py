"""Editable neighborhood score updates and two-action segmentation, without fitting."""
import csv
import hashlib
import json
import math
from decimal import Decimal
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value):
    path = Path(str(value))
    if path.is_absolute() or '..' in path.parts or not path.parts or path.parts[0] not in ('requirement-package', 'results'):
        raise ValueError('请选择当前项目的资料或结果文件')
    if not path.resolve().is_relative_to(Path.cwd().resolve()) or any(p.is_symlink() for p in [path, *path.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not path.is_file() or path.stat().st_size > 10_000_000:
        raise ValueError('文件不存在或超过10 MB')
    return path


def table(value, sheet, columns, limit, allow_empty=False):
    path = source(value)
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if not sheet and len(book.sheetnames) > 1:
                raise ValueError('Excel有多张表，请指定工作表：' + '、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:
                raise ValueError('工作表不存在：' + sheet)
            rows = list(islice((book[sheet] if sheet else book.active).values, limit + 2))
        finally:
            book.close()
    elif path.suffix.lower() in ('.csv', '.tsv'):
        with path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(islice(csv.reader(stream, delimiter='\t' if path.suffix.lower() == '.tsv' else ','), limit + 2))
    else:
        raise ValueError('请选择CSV、TSV或XLSX')
    if not rows or len(rows) > limit + 1:
        raise ValueError(f'{path.name} 需要表头，最多{limit}条记录')
    fields = ['' if v is None else str(v).strip() for v in rows[0]]
    if any(not v for v in fields) or len(fields) != len(set(fields)):
        raise ValueError('表头不能为空或重名')
    if set(columns) - set(fields):
        raise ValueError(path.name + ' 缺少字段：' + '、'.join(sorted(set(columns) - set(fields))))
    data = []
    for index, row in enumerate(rows[1:], 2):
        if all(v is None or not str(v).strip() for v in row):
            continue
        if len(row) != len(fields):
            raise ValueError(f'{path.name} 第{index}行列数与表头不同')
        data.append((index, dict(zip(fields, ['' if v is None else str(v).strip() for v in row]))))
    if not data and not allow_empty:
        raise ValueError(path.name + ' 没有数据记录')
    return path, data


def number(value, label, lo=0, hi=1e9, positive=False):
    try:
        result = float(value)
        if math.isfinite(result) and lo <= result <= hi and (not positive or result > 0):
            return result
    except (TypeError, ValueError, OverflowError):
        pass
    raise ValueError(f'{label} 需要{lo:g}至{hi:g}的有限数值' + ('，且必须大于0' if positive else ''))


def prepare(inputs):
    config = {'unit': str(inputs.get('unit') or '').strip()}
    if not config['unit']:
        raise ValueError('请填写坐标单位，两张表必须使用同一单位')
    for key, label, hi, positive in [
        ('grid_step', '等间距网格宽度', 1e9, True),
        ('influence_width', '观测影响宽度', 1e9, True),
        ('strength', '观测修正力度', 1, False),
        ('keep_cost', '保留风险单位代价', 1e9, False),
        ('discard_cost', '排除单位代价', 1e9, False),
        ('switch_cost', '相邻状态切换代价', 1e9, False),
        ('score_floor', '修正分数下限', 1, False),
        ('score_ceiling', '修正分数上限', 1, False),
    ]:
        config[key] = number(inputs.get(key), label, hi=hi, positive=positive)
    if config['score_floor'] >= config['score_ceiling']:
        raise ValueError('修正分数下限必须小于上限')
    columns = {key: str(inputs.get(key) or '').strip() for key in ('position_column', 'score_column', 'observed_column')}
    if any(not v for v in columns.values()):
        raise ValueError('请填写坐标列、原分数列和观测结论列')
    if columns['position_column'] in (columns['score_column'], columns['observed_column']):
        raise ValueError('坐标列不能同时作为分数或观测结论列')
    config.update(columns)
    curves, observations, sources = [], [], {}
    for kind, key, limit in [('curve', 'source_path', 10000), ('observations', 'observations_path', 200)]:
        if kind == 'observations' and not inputs.get(key):
            continue
        required = [columns['position_column'], columns['score_column']] if kind == 'curve' else ['observation_id', columns['position_column'], columns['observed_column']]
        path, rows = table(inputs.get(key), str(inputs.get(kind + '_sheet') or ''), required, limit, kind == 'observations')
        sources[kind] = {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheet': inputs.get(kind + '_sheet') or ''}
        seen = set()
        previous_position = None
        for index, row in rows:
            label = f'{path.name} 第{index}行'
            if 'unit' in row and row['unit'] != config['unit']:
                raise ValueError(label + ' 单位缺失或不一致，请先统一坐标单位')
            position = number(row[columns['position_column']], label + ' 坐标', -1e9, 1e9)
            if kind == 'curve':
                exact_position = Decimal(row[columns['position_column']])
                if previous_position is not None and exact_position-previous_position != Decimal(str(inputs['grid_step'])):
                    raise ValueError(label + ' 坐标必须严格递增、无重复且符合网格宽度；不自动补齐缺口')
                if not position-config['grid_step']/2 < position < position+config['grid_step']/2:
                    raise ValueError(label + ' 网格宽度太小，不能在当前坐标数值精度下表示，请改用局部坐标')
                previous_position = exact_position
                curves.append({'position': position, 'score': number(row[columns['score_column']], label + ' 原分数', hi=1), 'source_row': index})
            else:
                identity = row['observation_id']
                if not identity or identity in seen:
                    raise ValueError(label + ' observation_id为空或重复，请先核对是否重复导入')
                seen.add(identity)
                value = row[columns['observed_column']]
                if value not in ('1', '0'):
                    raise ValueError(label + ' 观测结论必须明确为1（高风险）或0（低风险），未知项请先复核')
                if not curves[0]['position']-config['grid_step']/2 <= position <= curves[-1]['position']+config['grid_step']/2:
                    raise ValueError(label + ' 观测坐标不在本次网格覆盖区间')
                observations.append({'observation_id': identity, 'position': position, 'high_risk': int(value), 'source_row': index})
    if len(curves) * len(observations) > 2_000_000:
        raise ValueError('本次网格和观测过多，请缩小本次分析范围')
    folder = Path('results') / ('local-feedback-input-' + uuid4().hex)
    folder.mkdir(parents=True)
    path = folder / 'input.json'
    path.write_text(json.dumps({'config': config, 'sources': sources, 'curve': curves, 'observations': observations}, ensure_ascii=False), encoding='utf-8')
    return {'snapshot_path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


def update_scores(positions, scores, observations, config):
    updated = list(scores)
    for observation in observations:
        for i, position in enumerate(positions):
            ratio = (position - observation['position']) / config['influence_width']
            weight = math.exp(-0.5 * ratio * ratio) * config['strength']
            updated[i] += weight * (observation['high_risk'] - updated[i])
    # No observations means no adjustment, including no clipping of original scores.
    if observations:
        updated = [min(config['score_ceiling'], max(config['score_floor'], s)) for s in updated]
    return updated


def plan_cost(labels, scores, config):
    local = sum(config['grid_step'] * (config['keep_cost'] * score if label == 0 else config['discard_cost']) for label, score in zip(labels, scores))
    switches = sum(a != b for a, b in zip(labels, labels[1:]))
    return local + switches * config['switch_cost']


def segment(positions, scores, config):
    previous = None
    back = []
    for score in scores:
        costs = [config['grid_step'] * config['keep_cost'] * score, config['grid_step'] * config['discard_cost']]
        if previous is None:
            previous = costs
            continue
        current, predecessors = [], []
        for state in (0, 1):
            alternatives = [previous[p] + (config['switch_cost'] if p != state else 0) + costs[state] for p in (0, 1)]
            predecessor = 0 if alternatives[0] <= alternatives[1] else 1
            predecessors.append(predecessor)
            current.append(alternatives[predecessor])
        previous = current
        back.append(predecessors)
    state = 0 if previous[0] <= previous[1] else 1
    labels = [state]
    for predecessors in reversed(back):
        state = predecessors[state]
        labels.append(state)
    labels.reverse()
    segments, cuts = [], []
    start = 0
    half = config['grid_step'] / 2
    for i in range(1, len(labels) + 1):
        if i == len(labels) or labels[i] != labels[start]:
            left, right = positions[start]-half, positions[i-1]+half
            segments.append({'action': '保留' if labels[start] == 0 else '排除', 'start': left, 'end': right, 'length': right-left, 'points': i-start})
            if i < len(labels):
                cuts.append(right)
            start = i
    return {'labels': labels, 'segments': segments, 'cuts': cuts, 'cost': plan_cost(labels, scores, config),
            'kept_length': sum(s['length'] for s in segments if s['action'] == '保留'),
            'discarded_length': sum(s['length'] for s in segments if s['action'] == '排除')}


def evaluate(inputs):
    prepared = inputs['prepared']; path = source(prepared['snapshot_path'])
    if hashlib.sha256(path.read_bytes()).hexdigest() != prepared['sha256']:
        raise ValueError('输入快照已改变，请重新读取资料')
    data = json.loads(path.read_text(encoding='utf-8')); cfg = data['config']; curve = data['curve']
    positions = [r['position'] for r in curve]; scores = [r['score'] for r in curve]
    updated = update_scores(positions, scores, data['observations'], cfg)
    before = segment(positions, scores, cfg); after = segment(positions, updated, cfg)
    changed = sum(a != b for a, b in zip(before['labels'], after['labels']))
    folder = Path('results') / ('local-feedback-' + uuid4().hex); folder.mkdir(parents=True); artifacts = []
    def save_csv(name, fields, rows, label):
        with (folder / name).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
        artifacts.append({'file_path': str(folder / name), 'label': label})
    rows = [{'source_row': row['source_row'], 'position': row['position'], 'original_score': row['score'], 'adjusted_score': updated[i],
             'change': updated[i]-row['score'], 'before_action': '保留' if before['labels'][i] == 0 else '排除',
             'after_action': '保留' if after['labels'][i] == 0 else '排除', 'unit': cfg['unit']} for i, row in enumerate(curve)]
    save_csv('scores.csv', list(rows[0]), rows, '逐点分数与决策 CSV')
    segments = [{'version': version, **s, 'unit': cfg['unit']} for version, plan in [('修正前', before), ('修正后', after)] for s in plan['segments']]
    save_csv('segments.csv', ['version', 'action', 'start', 'end', 'length', 'points', 'unit'], segments, '前后分段 CSV')
    save_csv('observations.csv', ['order', 'observation_id', 'position', 'high_risk', 'source_row'],
             [{'order': i+1, **r} for i, r in enumerate(data['observations'])], '实际使用的观测顺序 CSV')
    compact = lambda plan: {key: value for key, value in plan.items() if key not in ('labels', 'segments')}
    result = {'rows': len(curve), 'observations': len(data['observations']), 'changed_points': changed, 'before': compact(before), 'after': compact(after),
              'previous_plan_cost_on_adjusted_scores': plan_cost(before['labels'], updated, cfg), 'config': cfg, 'sources': data['sources'],
              'source_path': str(folder / 'scores.csv'), 'segments_path': str(folder / 'segments.csv')}
    def md(value):
        return str(value).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('|','\\|').replace('\n',' ').replace('\r',' ')
    lines = ['# 实测反馈修正与分段建议', '', f'读取{len(curve)}个网格点、{len(data["observations"])}条观测；{changed}个点的建议发生改变。坐标单位：{md(cfg["unit"])}。', '',
             '## 前后分段', '', '| 版本 | 保留长度 | 排除长度 | 内部边界数 | 声明代价 |', '|---|---:|---:|---:|---:|']
    for name, plan in [('修正前', before), ('修正后', after)]:
        lines.append(f'| {name} | {plan["kept_length"]:g} | {plan["discarded_length"]:g} | {len(plan["cuts"])} | {plan["cost"]:.6g} |')
    lines += ['', f'在相同修正分数上，原分段代价为{result["previous_plan_cost_on_adjusted_scores"]:.6g}，重算分段代价为{after["cost"]:.6g}。这是所填代价函数的比较，不是实际收益或准确率的提升。', '',
              '## 计算依据', '',
              f'依观测表行顺序，按距离和影响宽度{cfg["influence_width"]:g}计算高斯权重，乘修正力度{cfg["strength"]:g}，将附近分数拉向观测的0或1；全部更新后限制到[{cfg["score_floor"]:g}, {cfg["score_ceiling"]:g}]。没有观测时原分数完全保留。',
              '这是一种顺序相关的经验修正，不是贝叶斯推断或校准概率。相反观测的顺序可能改变结果；相同坐标的不同观测不会自动去重。影响宽度、力度和裁剪需用业务资料验证，不能仅凭本报告确定。',
              f'每格保留代价=网格宽度×{cfg["keep_cost"]:g}×风险分数；排除代价=网格宽度×{cfg["discard_cost"]:g}；相邻决策变化再加{cfg["switch_cost"]:g}。动态规划在这两个状态和本次网格上最小化总和；同值按保留优先的稳定顺序处理。',
              '坐标表示等宽单元的中心，边界向两端延伸半格；输出的是内部状态边界，不含端部切割。没有设置设备最短长度、切缝、总需求、已切除区域或不可逆现场约束，不能直接当作生产切割指令。', '',
              '## 分段明细（最多展示前30段，完整记录见下载）', '', '| 版本 | 建议 | 起点 | 终点 | 长度 |', '|---|---|---:|---:|---:|']
    lines.extend('| '+' | '.join(md(r[k]) for k in ('version','action','start','end','length'))+' |' for r in segments[:30])
    lines += ['', '## 下一步', '', '核对观测是否来自同一对象、相同坐标系，检查原分数含义与实际代价，再决定下一步。换观测或参数会新建结果；始终从原始曲线应用本次完整观测表，避免重复反馈。',
              '原模型与输入保持不变，没有重训、预测模型调用、库存修改或现场回写。泛化效果仍需独立实测评价；仅想比较预测和实测误差时使用已有“预测与实测反馈对照”。']
    result['markdown'] = '\n'.join(lines)
    (folder / 'report.md').write_text(result['markdown'], encoding='utf-8')
    (folder / 'summary.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    artifacts.extend([{'file_path': str(folder / 'report.md'), 'label': '反馈说明 Markdown'},
                      {'file_path': str(folder / 'summary.json'), 'label': '配置与比较 JSON'},
                      {'file_path': str(path), 'label': '本次输入快照 JSON'}])
    result['artifacts'] = artifacts
    return result


def main(inputs):
    if inputs.get('operation') == 'prepare':
        return prepare(inputs)
    if inputs.get('operation') == 'evaluate':
        return evaluate(inputs)
    raise ValueError('未知处理步骤')

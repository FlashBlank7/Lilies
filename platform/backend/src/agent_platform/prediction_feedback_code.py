"""Portable code nodes for comparing saved predictions with measured feedback."""
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value, limit=10_000_000):
    p = Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ('requirement-package', 'results'):
        raise ValueError('请选择当前项目资料或结果文件')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p, *p.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size > limit:
        raise ValueError(f'资料不存在或超过{limit // 1_000_000} MB，请检查文件或拆分')
    return p


def table(value, sheet=''):
    p = source(value)
    if p.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(p, read_only=True, data_only=True)
        try:
            if not sheet and len(book.sheetnames) > 1:
                raise ValueError('Excel有多张表，请指定工作表：' + '、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:
                raise ValueError('工作表不存在：' + sheet)
            rows = list(islice((book[sheet] if sheet else book.active).values, 100002))
        finally:
            book.close()
    elif p.suffix.lower() in ('.csv', '.tsv'):
        with p.open(encoding='utf-8-sig', newline='') as f:
            rows = list(islice(csv.reader(f, delimiter='\t' if p.suffix.lower() == '.tsv' else ','), 100002))
    else:
        raise ValueError('请选择CSV、TSV或XLSX')
    if not rows or len(rows) > 100001:
        raise ValueError('需要表头和数据，单表最多10万条记录')
    fields = [str(v).strip() if v is not None else '' for v in rows[0]]
    if not fields or any(not f for f in fields) or len(set(fields)) != len(fields):
        raise ValueError('表头不能为空或重名')
    result = []
    for n, raw in enumerate(rows[1:], 2):
        if all(v is None or not str(v).strip() for v in raw):
            continue
        if len(raw) != len(fields):
            raise ValueError(f'{p.name} 第{n}条记录列数与表头不同')
        result.append((n, dict(zip(fields, ['' if v is None else str(v) for v in raw]))))
    if not result:
        raise ValueError('表格没有数据记录')
    return fields, result


def columns(value):
    names = [v.strip() for v in re.split('[,，\n]', str(value or '')) if v.strip()]
    if len(names) != len(set(names)):
        raise ValueError('标识列或分组列重复，请每列只填写一次')
    return names


def index_rows(rows, keys, label):
    found = {}
    for record, row in rows:
        key = tuple(row[k].strip() for k in keys)
        if any(not k for k in key):
            raise ValueError(f'{label}第{record}条记录的样本标识为空')
        if key in found:
            raise ValueError(f'{label}第{record}与第{found[key][0]}条记录标识重复；请选定版本或补充组合标识，不自动去重')
        found[key] = (record, row)
    return found


def prepare(inputs):
    prediction_path = source(inputs['source_path'])
    separate = bool(str(inputs.get('actual_path') or '').strip())
    actual_path = source(inputs['actual_path']) if separate else prediction_path
    pf, predictions = table(prediction_path, inputs.get('prediction_sheet') or '')
    af, actuals = table(actual_path, inputs.get('actual_sheet') or '') if separate else (pf, predictions)
    names = {k:str(inputs.get(k) or '').strip() for k in ('prediction_column','actual_column','baseline_column')}
    keys, groups = columns(inputs.get('key_columns')), columns(inputs.get('group_columns'))
    if not keys:
        raise ValueError('请指定样本标识列；不按行号猜测对应关系')
    if not names['prediction_column'] or not names['actual_column']:
        raise ValueError('请指定预测结果列和实测结果列')
    if not separate and names['prediction_column'] == names['actual_column']:
        raise ValueError('同表的预测列与实测列不能相同')
    for fields, required, label in [(pf,keys+groups+[names['prediction_column']]+([names['baseline_column']] if names['baseline_column'] else []),'预测表'),
                                     (af,keys+[names['actual_column']],'实测表')]:
        missing = [k for k in required if k not in fields]
        if missing:
            raise ValueError(label+'缺少字段：'+'、'.join(missing))
        if any(k.startswith('eval_') for k in fields):
            raise ValueError(label+'已有eval_结果字段；请选择原始预测/实测文件，或先重命名这些列')
    index_rows(predictions, keys, '预测表')
    index_rows(actuals, keys, '实测表')
    mode = str(inputs.get('mode') or '')
    if mode not in ('数值预测','类别判断'):
        raise ValueError('评价方式请选择数值预测或类别判断')
    raw_tolerance = str(inputs.get('absolute_tolerance') if inputs.get('absolute_tolerance') is not None else '').strip()
    tolerance = None
    if raw_tolerance:
        if mode != '数值预测':
            raise ValueError('类别判断不使用数值容差，请清空容差')
        try:
            tolerance = float(raw_tolerance)
        except ValueError:
            raise ValueError('容差应为非负有限数值，或留空') from None
        if not math.isfinite(tolerance) or tolerance < 0:
            raise ValueError('容差应为非负有限数值，或留空')
    data = dict(mode=mode,keys=keys,groups=groups,**names,tolerance=tolerance,
                prediction_fields=pf,actual_fields=af,predictions=predictions,actuals=actuals,
                sources=[{'role':role,'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheet':sheet or ''}
                         for role,p,sheet in [('prediction',prediction_path,inputs.get('prediction_sheet')),
                                             ('actual',actual_path,inputs.get('actual_sheet') if separate else inputs.get('prediction_sheet'))]])
    folder = Path('results') / ('feedback-input-' + uuid4().hex)
    folder.mkdir(parents=True)
    snapshot = folder / 'input.json'
    snapshot.write_text(json.dumps(data,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    return {'snapshot_path':str(snapshot),'sha256':hashlib.sha256(snapshot.read_bytes()).hexdigest(),
            'prediction_rows':len(predictions),'actual_rows':len(actuals)}


def parsed(value, mode, label):
    value = str(value).strip()
    if not value:
        return None, label+'缺失'
    if mode == '类别判断':
        return value, ''
    try:
        number = float(value)
        if math.isfinite(number) and abs(number) <= 1e150:
            return number, ''
    except ValueError:
        pass
    return None, label+'不是有效有限数值（绝对值需不超过1e150）'


def metrics(pairs, mode, tolerance, labels=None):
    n = len(pairs)
    if mode == '类别判断':
        counts = Counter((actual,prediction) for prediction,actual in pairs)
        labels = sorted(set(labels or []) | {v for pair in pairs for v in pair})
        if len(labels) > 100:
            raise ValueError('类别超过100种，请检查是否把连续数值选成了类别判断，或选错了标签列')
        per_class = []
        for label in labels:
            support = sum(v for (a,p),v in counts.items() if a==label)
            predicted = sum(v for (a,p),v in counts.items() if p==label)
            tp = counts[label,label]
            precision = tp / predicted if predicted else 0.0
            recall = tp / support if support else 0.0
            per_class.append(dict(label=label,support=support,predicted=predicted,precision=precision,recall=recall,
                                  f1=2*precision*recall/(precision+recall) if precision+recall else 0.0))
        return dict(n=n,accuracy=sum(p==a for p,a in pairs)/n if n else None,
                    macro_f1=sum(v['f1'] for v in per_class)/len(labels) if labels and n else None,
                    classes=per_class,confusion=[dict(actual=a,prediction=p,count=c) for (a,p),c in sorted(counts.items())])
    if not n:
        return dict(n=0,mae=None,rmse=None,bias=None,r2=None,mape=None,mape_n=0,within_tolerance=None)
    errors = [p-a for p,a in pairs]
    squared = math.fsum(e*e for e in errors)
    mean_actual = math.fsum(a for p,a in pairs)/n
    total = math.fsum((a-mean_actual)**2 for p,a in pairs)
    # No epsilon is added to zero actuals; MAPE has its own explicit denominator.
    percentages = [abs((p-a)/a) for p,a in pairs if a != 0]
    ratio = squared/total if total else None
    relative = math.fsum(v/len(percentages) for v in percentages) if percentages and all(math.isfinite(v) and v<1e300 for v in percentages) else None
    return dict(n=n,mae=math.fsum(abs(e) for e in errors)/n,rmse=math.sqrt(squared/n),
                bias=math.fsum(errors)/n,r2=1-ratio if ratio is not None and math.isfinite(ratio) else None,
                mape=relative*100 if relative is not None and relative<1e298 else None,mape_n=len(percentages),
                within_tolerance=sum(abs(e)<=tolerance for e in errors)/n if tolerance is not None else None)


def evaluate(inputs):
    prepared = inputs['prepared']
    path = source(prepared['snapshot_path'],100_000_000)
    if hashlib.sha256(path.read_bytes()).hexdigest() != prepared['sha256']:
        raise ValueError('本次输入快照已改变，请重新读取资料')
    data = json.loads(path.read_text(encoding='utf-8'))
    mode, tolerance = data['mode'], data['tolerance']
    actuals = index_rows(data['actuals'],data['keys'],'实测表')
    seen, rows, pairs, comparable, baselines = set(), [], [], [], []
    grouped = defaultdict(lambda: {'rows':0,'pairs':[]})
    statuses = Counter()
    for record, row in data['predictions']:
        key = tuple(row[k].strip() for k in data['keys'])
        seen.add(key)
        match = actuals.get(key)
        actual_record, actual_row = match if match else ('',{})
        p, pe = parsed(row[data['prediction_column']],mode,'预测结果')
        a, ae = parsed(actual_row.get(data['actual_column'],''),mode,'实测结果') if match else (None,'没有匹配的实测记录')
        valid = not pe and not ae
        status = 'evaluated' if valid else 'no_actual_record' if not match else 'invalid_values'
        result = {**row,'eval_prediction_record':record,'eval_actual_record':actual_record,'eval_prediction':row[data['prediction_column']],
                  'eval_actual':actual_row.get(data['actual_column'],''),'eval_status':status,
                  'eval_reason':'；'.join(x for x in (pe,ae) if x),'eval_error':'','eval_absolute_error':'','eval_correct':'','eval_within_tolerance':'',
                  'eval_baseline_reason':''}
        group = tuple(row[k].strip() for k in data['groups'])
        grouped[group]['rows'] += 1
        if valid:
            pairs.append((p,a)); grouped[group]['pairs'].append((p,a))
            if mode == '数值预测':
                result.update(eval_error=p-a,eval_absolute_error=abs(p-a),
                              eval_within_tolerance=abs(p-a)<=tolerance if tolerance is not None else '')
            else:
                result['eval_correct'] = p==a
        if data['baseline_column']:
            b, be = parsed(row[data['baseline_column']],mode,'基线结果')
            result['eval_baseline_reason'] = be or ('' if valid else '预测或实测无效，不进入共同评价')
            if valid and not be:
                comparable.append((p,a)); baselines.append((b,a))
        rows.append(result); statuses[status] += 1
    overall = metrics(pairs,mode,tolerance)
    group_labels = [c['label'] for c in overall.get('classes',[])]
    group_results = [{**dict(zip(data['groups'],group)),'eval_rows':value['rows'],
                      **{'eval_'+k:v for k,v in metrics(value['pairs'],mode,tolerance,group_labels).items() if k not in ('classes','confusion')}}
                     for group,value in grouped.items()] if data['groups'] else []
    common_labels = sorted({v for pair in comparable+baselines for v in pair}) if mode=='类别判断' else []
    unmatched = [{**row,'eval_actual_record':n} for key,(n,row) in actuals.items() if key not in seen]
    summary = dict(mode=mode,prediction_rows=len(rows),actual_rows=len(data['actuals']),counts=dict(statuses),
                   evaluated_rows=len(pairs),coverage=len(pairs)/len(rows),unmatched_actual_rows=len(unmatched),
                   metrics=overall,groups=group_results,group_columns=data['groups'],tolerance=tolerance,sources=data['sources'],
                   baseline_comparison={'n':len(comparable),'prediction':metrics(comparable,mode,tolerance,common_labels),'baseline':metrics(baselines,mode,tolerance,common_labels)} if data['baseline_column'] else None,
                   scope='已有预测与实测的对照；未核实原训练/测试来源，不证明泛化表现。')
    folder = Path('results') / ('prediction-feedback-' + uuid4().hex)
    folder.mkdir(parents=True)
    artifacts = []
    def save_csv(name,fields,values,label):
        with (folder/name).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(values)
        artifacts.append({'file_path':str(folder/name),'label':label})
    save_csv('comparison.csv',list(rows[0]),rows,'逐条预测与实测对照 CSV')
    save_csv('unmatched-actuals.csv',data['actual_fields']+['eval_actual_record'],unmatched,'没有对应预测的实测记录 CSV')
    if group_results:
        save_csv('groups.csv',list(group_results[0]),group_results,'按批次或设备分组的结果 CSV')
    if mode == '类别判断' and pairs:
        save_csv('classes.csv',list(overall['classes'][0]),overall['classes'],'每个类别的数量与评价 CSV')
        save_csv('confusion.csv',['actual','prediction','count'],overall['confusion'],'实际与预测类别对照 CSV')
    markdown = report(summary,rows,data['keys'])
    for name,content,label in [('summary.json',json.dumps(summary,ensure_ascii=False,indent=2,allow_nan=False),'评价口径与数据版本 JSON'),
                               ('report.md',markdown,'反馈分析说明 Markdown')]:
        (folder/name).write_text(content,encoding='utf-8');artifacts.append({'file_path':str(folder/name),'label':label})
    return {**summary,'source_path':str(folder/'comparison.csv'),'markdown':markdown,'artifacts':artifacts,'preview':rows[:12]}


def report(summary, rows, keys):
    def number(value):
        return '无法计算' if value is None else f'{value:.6g}'
    def cell(value):
        return str(value).replace('\\','\\\\').replace('|','\\|').replace('\r',' ').replace('\n',' ')[:120]
    def metric_text(m):
        if summary['mode']=='类别判断':
            return f"准确率 {number(m['accuracy'])}，宏平均F1 {number(m['macro_f1'])}"
        return f"平均绝对误差 {number(m['mae'])}，均方根误差 {number(m['rmse'])}，平均偏差 {number(m['bias'])}"
    m = summary['metrics']
    text = f"# 预测与实测反馈对照\n\n{summary['prediction_rows']} 条预测中，{m['n']} 条具备有效预测与实测，可评价比例 {summary['coverage']:.1%}。"
    text += f"没有匹配实测的预测记录 {summary['counts'].get('no_actual_record',0)} 条，匹配后缺值或无效 {summary['counts'].get('invalid_values',0)} 条；另有 {summary['unmatched_actual_rows']} 条实测找不到对应预测。\n\n"
    text += f"## 这些结果说明什么\n\n本次 {summary['mode']}：{metric_text(m)}。\n\n"
    if summary['mode']=='数值预测':
        text += '平均绝对误差表示通常差多少，均方根误差对大偏差更敏感，单位与实测值相同。误差=预测−实测，平均偏差为正表示整体偏高，为负表示偏低。\n\n'
        text += f"R² {number(m['r2'])}；实测为常量或没有可评价行时不定义。MAPE（平均绝对百分比误差）{number(m['mape'])}%，计算样本 {m['mape_n']} 条，排除实测0；非常接近0时可能失真或无法表示，应优先看绝对误差。\n\n"
        text += '没有设置业务容差，因此本流程不判断是否合格。\n\n' if summary['tolerance'] is None else f"用户设置绝对误差容差 {summary['tolerance']}，有效样本中容差内比例 {number(m['within_tolerance'])}（0到1）。这不是产品放行决定。\n\n"
    else:
        text += '准确率表示判断正确的比例；宏平均F1给每个实际或预测类别相同权重。未知预测类别也计入错误，少数类请结合样本量看。没有真实样本或没有预测样本的类别，其不可定义的精确率/召回率按0计入宏平均；不隐藏缺类问题。\n\n'
        text += '| 类别 | 实测数量 | 预测数量 | 精确率 | 召回率 | F1 |\n| --- | --- | --- | --- | --- | --- |\n'
        text += '\n'.join('| '+' | '.join([cell(c['label']),str(c['support']),str(c['predicted']),number(c['precision']),number(c['recall']),number(c['f1'])])+' |' for c in m['classes'][:20])+'\n\n最多展示20类，完整类别表与混淆计数见下载文件。\n\n'
    if summary['baseline_comparison']:
        b = summary['baseline_comparison']
        text += f"## 和已有基线比较\n\n仅在双方结果与实测都有效的同一组 {b['n']} 条记录上比较。模型：{metric_text(b['prediction'])}；基线：{metric_text(b['baseline'])}。其覆盖范围可能小于上面的整体结果，不跨样本比较。"
        text += ('类别比较使用相同类别集合（双方预测与实测的并集）。' if summary['mode']=='类别判断' else '')+'\n\n'
    if summary['groups']:
        text += '## 分组观察\n\n| 分组 | 预测条数 | 有效条数 | 评价 |\n| --- | --- | --- | --- |\n'
        for g in summary['groups'][:20]:
            label=' / '.join(k+'='+str(g[k] or '（空）') for k in summary['group_columns'])
            gm={k.removeprefix('eval_'):v for k,v in g.items() if k.startswith('eval_')}
            text += f"| {cell(label)} | {g['eval_rows']} | {g['eval_n']} | {metric_text(gm)} |\n"
        text += '\n最多展示20组，完整结果见CSV。'
        text += '各组沿用整体类别集合；没有该类的组也保留相同计算口径。' if summary['mode']=='类别判断' else ''
        text += '先看各组样本数，再看差异；小样本或当前分组差异不能证明原因及生产稳定性。\n\n'
    text += '## 建议下一步\n\n先补齐缺失实测并核对标识及单位，再查看误差较大或类别混淆的记录。若要调整采纳规则，使用已有规则重算流程；若要改模型，再提出训练任务。这里不会自动启动下一步。\n\n'
    text += summary['scope']+' 本次不重新预测、不训练或改写模型，原资料保持不变。\n\n'
    fields=keys[:3]+['eval_prediction','eval_actual','eval_error' if summary['mode']=='数值预测' else 'eval_correct','eval_status','eval_reason']
    labels={'eval_prediction':'预测','eval_actual':'匹配实测','eval_error':'预测−实测','eval_correct':'是否正确','eval_status':'结果','eval_reason':'说明'}
    status_labels={'evaluated':'已评价','no_actual_record':'无匹配实测','invalid_values':'缺值或无效'}
    text+='## 逐条预览\n\n| '+' | '.join(cell(labels.get(k,k)) for k in fields)+' |\n| '+' | '.join('---' for _ in fields)+' |\n'
    def preview_value(row,k):
        if k=='eval_status':return status_labels.get(row[k],row[k])
        if k=='eval_correct' and isinstance(row[k],bool):return '是' if row[k] else '否'
        return row[k]
    text+='\n'.join('| '+' | '.join(cell(preview_value(r,k)) for k in fields)+' |' for r in rows[:12])
    return text+'\n\n预览最多12条、3个标识字段，完整预测原字段、匹配记录号与误差见逐条CSV。'


def main(inputs):
    if inputs['operation']=='prepare':
        return prepare(inputs)
    if inputs['operation']=='evaluate':
        return evaluate(inputs)
    raise ValueError('未知的反馈分析步骤')

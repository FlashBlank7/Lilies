"""Portable candidate checks and comparison; no solver, model call or production write."""
import csv
import hashlib
import json
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, localcontext
from fractions import Fraction
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value, limit=10_000_000):
    p = Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ('requirement-package', 'results'):
        raise ValueError('请选择当前项目资料或结果文件')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p, *p.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size > limit:
        raise ValueError(f'资料不存在或超过{limit // 1_000_000} MB，请检查文件或拆分')
    return p


def table(value, sheet=''):
    p = source(value)
    if p.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(p, read_only=True, data_only=True)
        try:
            if not sheet and len(book.sheetnames) > 1:
                raise ValueError('Excel有多张表，请指定工作表：' + '、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:
                raise ValueError('工作表不存在：' + sheet)
            rows = list(islice((book[sheet] if sheet else book.active).values, 100002))
        finally:
            book.close()
    elif p.suffix.lower() in ('.csv', '.tsv'):
        with p.open(encoding='utf-8-sig', newline='') as f:
            rows = list(islice(csv.reader(f, delimiter='\t' if p.suffix.lower() == '.tsv' else ','), 100002))
    else:
        raise ValueError('请选择CSV、TSV或XLSX')
    if not rows or len(rows) > 100001:
        raise ValueError('需要表头和数据，单表最多10万条记录')
    fields = [str(v).strip() if v is not None else '' for v in rows[0]]
    if not fields or any(not f for f in fields) or len(set(fields)) != len(fields):
        raise ValueError('表头不能为空或重名')
    result = []
    for n, raw in enumerate(rows[1:], 2):
        if all(v is None or not str(v).strip() for v in raw):
            continue
        if len(raw) != len(fields):
            raise ValueError(f'{p.name} 第{n}条记录列数与表头不同')
        result.append((n, dict(zip(fields, ['' if v is None else str(v) for v in raw]))))
    if not result:
        raise ValueError('表格没有数据记录')
    return fields, result


OPERATORS = ['不大于','小于','不小于','大于','数值等于','数值不等于','文字等于','文字不等于']
DIRECTIONS = ['越大越好','越小越好']


def number(value):
    try:
        n=Decimal(str(value).strip())
        if n.is_finite() and len(n.as_tuple().digits)<=40 and (n==0 or -40<=n.adjusted()<=40):
            return n
    except (InvalidOperation,ValueError):
        pass
    raise ValueError('需要有限数值（最多40位有效数字，非零值数量级在1e-40至1e40之间）')


def prepare(inputs):
    path=source(inputs['source_path'])
    fields, rows=table(path,inputs.get('sheet') or '')
    ident=str(inputs.get('id_column') or '').strip(); group=str(inputs.get('group_column') or '').strip()
    if ident not in fields or (group and group not in fields):
        raise ValueError('请检查方案标识列和分组列是否存在')
    if any(f.startswith('comparison_') for f in fields):
        raise ValueError('已有comparison_结果字段，请选择原候选表或重命名这些列')
    constraints=inputs.get('constraints') or []; objectives=inputs.get('objectives') or []
    if not isinstance(constraints,list) or not isinstance(objectives,list):
        raise ValueError('条件和比较目标需要使用行表配置')
    if len(rows)>20000 or len(constraints)>50 or len(objectives)>10 or len(rows)*max(1,len(constraints))>500000:
        raise ValueError('首版最多2万候选、50个条件、10个目标及50万次条件检查，请按任务分组拆分')
    mode=inputs.get('mode','按优先级逐项比较')
    if mode not in ['按优先级逐项比较','按明确权重评分']:
        raise ValueError('请选择有效的比较方式')
    normalized=[]
    for i,rule in enumerate(constraints,1):
        if not isinstance(rule,dict):raise ValueError(f'条件第{i}行不是字段配置')
        field=str(rule.get('field') or '').strip(); op=rule.get('operator')
        other=str(rule.get('other_field') or '').strip(); value=str(rule.get('value') if rule.get('value') is not None else '').strip()
        if field not in fields or (other and other not in fields):raise ValueError(f'条件第{i}行引用的字段不存在')
        if op not in OPERATORS:raise ValueError(f'条件第{i}行请选择比较方式')
        if bool(other)==bool(value):raise ValueError(f'条件第{i}行的固定值和对比列应恰好填写一个')
        if not other and not op.startswith('文字'):
            try:number(value)
            except ValueError as exc:raise ValueError(f'条件第{i}行的固定值：{exc}') from None
        normalized.append(dict(field=field,operator=op,value=value,other_field=other))
    targets=[]
    for i,obj in enumerate(objectives,1):
        if not isinstance(obj,dict):raise ValueError(f'目标第{i}行不是字段配置')
        field=str(obj.get('field') or '').strip(); direction=obj.get('direction')
        if field not in fields:raise ValueError(f'目标第{i}行字段不存在')
        if field in [t['field'] for t in targets]:raise ValueError('同一比较字段不能重复填写')
        if direction not in DIRECTIONS:raise ValueError(f'目标第{i}行请选择越大或越小越好')
        weight=scale=Decimal(1)
        if mode=='按明确权重评分':
            try:weight=number(obj.get('weight'));scale=number(obj.get('scale'))
            except ValueError as exc:raise ValueError(f'目标第{i}行的权重及单位尺度：{exc}') from None
            if weight<=0 or scale<=0:raise ValueError(f'目标第{i}行的权重及单位尺度必须大于0')
        targets.append(dict(field=field,direction=direction,weight=str(weight),scale=str(scale)))
    seen={}
    for record,row in rows:
        key=(row[group].strip() if group else '',row[ident].strip())
        if not key[1] or (group and not key[0]):raise ValueError(f'第{record}条记录的方案标识或分组为空')
        if key in seen:raise ValueError(f'第{record}与第{seen[key]}条记录的方案标识在同组内重复')
        seen[key]=record
    data=dict(fields=fields,rows=rows,id_column=ident,group_column=group,constraints=normalized,objectives=targets,mode=mode,
              source={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':inputs.get('sheet') or ''})
    folder=Path('results')/('candidate-input-'+uuid4().hex);folder.mkdir(parents=True)
    snapshot=folder/'input.json';snapshot.write_text(json.dumps(data,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    return {'snapshot_path':str(snapshot),'sha256':hashlib.sha256(snapshot.read_bytes()).hexdigest(),'rows':len(rows)}


def check(row, rule):
    left=row[rule['field']].strip();right=row[rule['other_field']].strip() if rule['other_field'] else rule['value']
    if not left or not right:return 'unknown','缺少用于比较的值',left,right
    op=rule['operator']
    if op.startswith('文字'):
        passed=left==right if op=='文字等于' else left!=right
    else:
        try:a,b=number(left),number(right)
        except ValueError:return 'unknown','用于比较的字段不是有效数值',left,right
        passed={'不大于':a<=b,'小于':a<b,'不小于':a>=b,'大于':a>b,'数值等于':a==b,'数值不等于':a!=b}[op]
    return 'pass' if passed else 'fail','满足' if passed else '不满足',left,right


def fraction_text(value):
    with localcontext() as ctx:
        ctx.prec=24
        return str(Decimal(value.numerator)/Decimal(value.denominator))


def compare(inputs):
    prepared=inputs['prepared'];snapshot=source(prepared['snapshot_path'],100_000_000)
    if hashlib.sha256(snapshot.read_bytes()).hexdigest()!=prepared['sha256']:raise ValueError('候选或配置快照已改变，请重新读取资料')
    data=json.loads(snapshot.read_text(encoding='utf-8'));ident=data['id_column'];group_col=data['group_column']
    details=[];results=[];groups=defaultdict(list);totals=Counter()
    for record,row in data['rows']:
        group=row[group_col].strip() if group_col else '全部方案'
        state=[];reasons=[]
        for i,rule in enumerate(data['constraints'],1):
            status,reason,left,right=check(row,rule);state.append(status)
            expression=rule['field']+' '+rule['operator']+' '+(rule['other_field'] or rule['value'])
            if status!='pass':reasons.append(f'条件{i}（{expression}）：{reason}')
            details.append(dict(group=group,candidate_id=row[ident],source_record=record,condition=i,field=rule['field'],operator=rule['operator'],
                                compared_with=rule['other_field'] or '固定值',actual_value=left,expected_value=right,status=status,reason=reason))
        status='fail' if 'fail' in state else 'unknown' if 'unknown' in state else 'pass' if state else 'unchecked'
        values=[];contributions=[];objective_errors=[]
        for i,obj in enumerate(data['objectives'],1):
            try:
                n=number(row[obj['field']]); signed=Fraction(n)*(1 if obj['direction']=='越大越好' else -1)
                contribution=signed*Fraction(obj['weight'])/Fraction(obj['scale'])
                values.append(signed);contributions.append(contribution)
            except ValueError:
                objective_errors.append(f"目标{i}（{obj['field']}）缺失或不是有效数值")
        key=tuple(values) if data['mode']=='按优先级逐项比较' else (sum(contributions,Fraction(0)),)
        result={**row,'comparison_source_record':record,'comparison_status':status,'comparison_reason':'；'.join(reasons),
                'comparison_rank':'','comparison_score':'','comparison_objective_issue':'；'.join(objective_errors)}
        if data['mode']=='按明确权重评分' and data['objectives'] and not objective_errors:
            result['comparison_score']=fraction_text(key[0])
        results.append(result);totals[status]+=1
        groups[group].append(dict(row=result,key=key,ready=not objective_errors and bool(data['objectives']) and status in ('pass','unchecked')))
    summaries=[]
    for group,entries in groups.items():
        eligible=sorted((e for e in entries if e['ready']),key=lambda e:e['key'],reverse=True)
        previous=None;rank=0;top=[]
        for position,entry in enumerate(eligible,1):
            if entry['key']!=previous:rank=position;previous=entry['key']
            entry['row']['comparison_rank']=rank
            if rank==1:top.append(entry['row'][ident])
        count=Counter(e['row']['comparison_status'] for e in entries)
        summaries.append(dict(group=group,rows=len(entries),passed=count['pass'],failed=count['fail'],unknown=count['unknown'],unchecked=count['unchecked'],
                              ranked=len(eligible),top_candidates=top,
                              status='only_sorted_without_constraints' if not data['constraints'] else 'no_candidate_passes' if not count['pass'] else
                              'no_objectives' if not data['objectives'] else 'missing_objective_values' if not eligible else 'compared'))
    summary=dict(rows=len(results),counts=dict(totals),groups=summaries,mode=data['mode'],constraints=data['constraints'],objectives=data['objectives'],source=data['source'],
                 scope='只比较所给候选及已声明条件，不证明全局最优、现场可行或实际安全，不执行生产回写。')
    folder=Path('results')/('candidate-comparison-'+uuid4().hex);folder.mkdir(parents=True)
    artifacts=[]
    def save_csv(name,fields,rows,label):
        with (folder/name).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(rows)
        artifacts.append({'file_path':str(folder/name),'label':label})
    save_csv('candidates.csv',list(results[0]),results,'全部候选与比较结果 CSV')
    save_csv('conditions.csv',['group','candidate_id','source_record','condition','field','operator','compared_with','actual_value','expected_value','status','reason'],details,'每条条件的实际值与原因 CSV')
    # Fixed scale means scores do not change merely because another candidate is added.
    score_details=[]
    for row in results:
        if row['comparison_objective_issue']:continue
        for i,obj in enumerate(data['objectives'],1):
            value=number(row[obj['field']]);part=Fraction(value)*Fraction(obj['weight'])/Fraction(obj['scale'])*(1 if obj['direction']=='越大越好' else -1)
            score_details.append(dict(candidate_id=row[ident],group=row[group_col] if group_col else '全部方案',priority=i,field=obj['field'],value=str(value),
                                      direction=obj['direction'],weight=obj['weight'],scale=obj['scale'],contribution=fraction_text(part) if data['mode']=='按明确权重评分' else ''))
    save_csv('objectives.csv',['candidate_id','group','priority','field','value','direction','weight','scale','contribution'],score_details,'比较目标与评分分项 CSV')
    markdown=report(summary,results,ident,group_col)
    for name,content,label in [('summary.json',json.dumps(summary,ensure_ascii=False,indent=2),'条件、配置与结果汇总 JSON'),('report.md',markdown,'方案比较说明 Markdown')]:
        (folder/name).write_text(content,encoding='utf-8');artifacts.append({'file_path':str(folder/name),'label':label})
    return {**summary,'source_path':str(folder/'candidates.csv'),'markdown':markdown,'preview':results[:12],'artifacts':artifacts}


def report(summary,rows,ident,group_col):
    def cell(value):return str(value).replace('\\','\\\\').replace('|','\\|').replace('\n',' ').replace('\r',' ')[:200]
    labels={'pass':'满足已声明条件','fail':'不满足','unknown':'待补充条件数据','unchecked':'未检查限制'}
    counts=summary['counts']
    text=f"# 候选方案约束检查与比较\n\n共 {summary['rows']} 个候选；满足已声明条件 {counts.get('pass',0)} 个，不满足 {counts.get('fail',0)} 个，待补充条件数据 {counts.get('unknown',0)} 个，未检查限制 {counts.get('unchecked',0)} 个。\n\n"
    text+='## 如何读这个结果\n\n'+('没有声明限制，本次仅排序，不能据此判断方案可行。\n\n' if not summary['constraints'] else '所有声明条件都满足才参加后续比较；缺值不会当作0或自动放行。全部不满足时不回退选择高分方案。\n\n')
    if summary['mode']=='按优先级逐项比较':text+='比较目标按表单从上到下逐项比较：只有前面的目标相同，才比较下一项，不把各指标相加。\n\n'
    else:text+='分数=各指标的（数值÷单位尺度×权重×方向）之和，越大越优先；越大越好取正，越小越好取负。尺度和权重由本次配置明确给定，不根据候选集合自动变化。不同配置的分数不直接比较。硬条件不会被高分抵消。\n\n'
    if summary['objectives']:
        text+='本次目标：'+'；'.join(f"{i}. {cell(o['field'])} {o['direction']}"+(f"，权重{o['weight']}，尺度{o['scale']}" if summary['mode']=='按明确权重评分' else '') for i,o in enumerate(summary['objectives'],1))+'。\n\n'
    else:text+='没有设置比较目标，保留条件检查结果，不代替员工选择方案。\n\n'
    text+='## 各组结论\n\n'
    for g in summary['groups'][:20]:
        text+=f"- {cell(g['group'])}：{g['rows']} 个候选，满足条件 {g['passed']} 个，进入比较 {g['ranked']} 个。"
        if g['status']=='no_candidate_passes':text+='当前没有可确认满足全部已声明条件的候选，请查看失败/缺失原因，补充方案或修正真实输入。'
        elif g['status']=='missing_objective_values':text+='已通过条件的候选缺少比较目标，暂不能排序。'
        elif g['top_candidates']:text+='当前比较口径下排在首位：'+'、'.join(cell(k) for k in g['top_candidates'][:20])+('（同分并列）' if len(g['top_candidates'])>1 else '')+'。'
        text+='\n'
    text+='\n分组摘要最多20组，每组首位最多显示20个；完整候选排名见CSV。同分保留并列，不任意声称某个方案唯一最佳。\n\n'
    preview_targets=summary['objectives'][:5]
    headings=['方案','分组','条件检查','名次','分数']+[o['field']+'（'+o['direction']+'）' for o in preview_targets]+['原因或待补充内容']
    text+='## 候选预览\n\n| '+' | '.join(map(cell,headings))+' |\n| '+' | '.join('---' for _ in headings)+' |\n'
    for r in rows[:12]:
        values=[r[ident],r[group_col] if group_col else '全部方案',labels[r['comparison_status']],r['comparison_rank'],r['comparison_score']]+[r[o['field']] for o in preview_targets]+['；'.join(v for v in [r['comparison_reason'],r['comparison_objective_issue']] if v)]
        text+='| '+' | '.join(map(cell,values))+' |\n'
    text+='\n预览保留输入顺序、最多12行；下载表保留全部原字段及每条限制的实际值和依据。\n\n## 下一步\n\n核对指标单位、候选覆盖范围和现场约束；有缺项先补齐。修改条件或目标可重新比较，不需要重新训练。若不同任务争用库存、设备或时间段，各组首位不能直接组成总体排程，需额外核对共同资源与时序约束。若需要生成更多方案，可另行使用代码或优化流程；本流程不会自动生成、执行或回写方案。\n\n'+summary['scope']
    return text


def main(inputs):
    if inputs['operation']=='prepare':return prepare(inputs)
    if inputs['operation']=='compare':return compare(inputs)
    raise ValueError('未知的方案比较步骤')

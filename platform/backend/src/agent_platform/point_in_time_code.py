"""Portable code nodes: build samples using only feature versions known then."""
import csv
import hashlib
import json
import math
import re
from collections import Counter
from datetime import datetime, timedelta, timezone
from itertools import islice
from pathlib import Path
from uuid import uuid4


def source(value, limit=10_000_000):
    p = Path(str(value))
    if p.is_absolute() or '..' in p.parts or not p.parts or p.parts[0] not in ('requirement-package', 'results'):
        raise ValueError('请选择当前项目资料或结果文件')
    if not p.resolve().is_relative_to(Path.cwd().resolve()) or any(x.is_symlink() for x in [p, *p.parents]):
        raise ValueError('不支持项目外文件或符号链接')
    if not p.is_file() or p.stat().st_size > limit:
        raise ValueError(f'资料不存在或超过{limit // 1_000_000} MB，请检查文件或拆分')
    return p


def table(value, sheet=''):
    p = source(value)
    if p.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(p, read_only=True, data_only=True)
        try:
            if not sheet and len(book.sheetnames) > 1:
                raise ValueError('Excel有多张表，请指定工作表：' + '、'.join(book.sheetnames))
            if sheet and sheet not in book.sheetnames:
                raise ValueError('工作表不存在：' + sheet)
            rows = list(islice((book[sheet] if sheet else book.active).values, 100002))
        finally:
            book.close()
    elif p.suffix.lower() in ('.csv', '.tsv'):
        with p.open(encoding='utf-8-sig', newline='') as f:
            rows = list(islice(csv.reader(f, delimiter='\t' if p.suffix.lower() == '.tsv' else ','), 100002))
    else:
        raise ValueError('请选择CSV、TSV或XLSX')
    if not rows or len(rows) > 100001:
        raise ValueError('需要表头和数据，单表最多10万条记录')
    fields = [str(v).strip() if v is not None else '' for v in rows[0]]
    if not fields or any(not f for f in fields) or len(set(fields)) != len(fields):
        raise ValueError('表头不能为空或重名')
    result = []
    for n, raw in enumerate(rows[1:], 2):
        if all(v is None or not str(v).strip() for v in raw):
            continue
        if len(raw) != len(fields):
            raise ValueError(f'{p.name} 第{n}条记录列数与表头不同')
        result.append((n, dict(zip(fields, ['' if v is None else str(v) for v in raw]))))
    if not result:
        raise ValueError('表格没有数据记录')
    return fields, result


def offset(value):
    if not re.fullmatch(r'[+-]\d{2}:\d{2}', value):
        raise ValueError('默认时区请填写UTC偏移，例如 +08:00 或 +00:00')
    hour, minute = int(value[1:3]), int(value[4:6])
    if hour > 14 or minute > 59 or (hour == 14 and minute):
        raise ValueError('UTC偏移超出支持范围')
    return timezone((1 if value[0] == '+' else -1) * timedelta(hours=hour, minutes=minute))


def stamp(value, zone, label, precise=True):
    text = str(value).strip()
    if precise and not re.search(r'[T ]\d{2}:\d{2}', text):
        raise ValueError(label + '需要实际时分；不能把未知发布时间补成当天零点')
    try:
        time = datetime.fromisoformat(text.replace('Z', '+00:00'))
        if time.tzinfo is None:
            time = time.replace(tzinfo=zone)
        return time.astimezone(timezone.utc).isoformat()
    except ValueError:
        raise ValueError(label + '不是有效ISO日期时间：' + text) from None


def prepare(inputs):
    sample_path, records_path = source(inputs['source_path']), source(inputs['records_path'])
    sf, samples = table(sample_path, inputs.get('sample_sheet', ''))
    rf, records = table(records_path, inputs.get('records_sheet', ''))
    names = {k: str(inputs.get(k, '')).strip() for k in ('sample_id', 'prediction_time', 'entity', 'feature', 'observed_time', 'available_time', 'value')}
    for fields, required, label in [(sf, ['sample_id', 'prediction_time', 'entity'], '样本表'),
                                   (rf, ['entity', 'feature', 'observed_time', 'available_time', 'value'], '记录表')]:
        for k in required:
            if not names[k] and k == 'entity':
                continue
            if names[k] not in fields:
                raise ValueError(label + '缺少字段：' + (names[k] or k))
        actual = [names[k] for k in required if names[k]]
        if len(set(actual)) != len(actual):
            raise ValueError(label + '中的标识、时间、指标和值应使用不同字段')
    zone = offset(str(inputs.get('utc_offset', '+00:00')))
    age = float(inputs.get('max_age_hours') or 0)
    if not math.isfinite(age) or age < 0:
        raise ValueError('最长历史时长应为非负小时数，0表示不限制')
    wanted = list(dict.fromkeys(x.strip() for x in re.split(r'[,\n]', str(inputs.get('features') or '')) if x.strip()))
    sample_ids, version_keys, prepared, events = set(), set(), [], []
    for pos, row in samples:
        ident = row[names['sample_id']].strip()
        if not ident or ident in sample_ids:
            raise ValueError(f'样本表第{pos}条记录的样本标识为空或重复：{ident}')
        sample_ids.add(ident)
        entity = row[names['entity']].strip() if names['entity'] else ''
        if names['entity'] and not entity:
            raise ValueError(f'样本表第{pos}条记录缺少实体标识')
        prepared.append({'record': pos, 'id': ident, 'entity': entity, 'values': row,
                         'cutoff': stamp(row[names['prediction_time']], zone, f'样本表第{pos}条记录的预测时点')})
    for pos, row in records:
        entity = row[names['entity']].strip() if names['entity'] else ''
        feature = row[names['feature']].strip()
        if not feature or (names['entity'] and not entity):
            raise ValueError(f'记录表第{pos}条记录缺少实体或指标名称')
        observed = stamp(row[names['observed_time']], zone, f'记录表第{pos}条记录的所属时间', precise=False)
        available = stamp(row[names['available_time']], zone, f'记录表第{pos}条记录的可用时间')
        key = (entity, feature, observed, available)
        if key in version_keys:
            raise ValueError(f'记录表第{pos}条记录的实体、指标、所属时间、可用时间重复；请先确定唯一来源版本')
        version_keys.add(key)
        events.append({'entity': entity, 'feature': feature, 'observed': observed, 'available': available,
                       'eligible': max(observed, available), 'record': pos, 'value': row[names['value']]})
    available_features = sorted({e['feature'] for e in events})
    wanted = wanted or available_features
    if set(wanted) - set(available_features):
        raise ValueError('记录表没有所选指标：' + '、'.join(sorted(set(wanted) - set(available_features))))
    if len(wanted) > 100 or len(wanted) * len(samples) > 1_000_000:
        raise ValueError('最多100个指标、100万个样本指标组合，请分组处理')
    if any('asof_' + f in sf for f in wanted):
        raise ValueError('样本表已有同名asof_字段，请改名或选择原样本表，避免覆盖旧特征')
    snapshot = {'samples': prepared, 'events': [e for e in events if e['feature'] in wanted], 'features': wanted,
                'sample_fields': sf, 'max_age_hours': age, 'utc_offset': str(inputs.get('utc_offset', '+00:00')),
                'sources': [{'path': str(p), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in (sample_path, records_path)]}
    folder = Path('results/point-in-time') / str(uuid4()); folder.mkdir(parents=True)
    path = folder / 'prepared.json'; path.write_text(json.dumps(snapshot, ensure_ascii=False), encoding='utf-8')
    return {'snapshot_path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'samples': len(samples), 'records': len(records), 'features': wanted}


def align(inputs):
    prepared = inputs['prepared']; path = source(prepared['snapshot_path'], limit=100_000_000)
    if hashlib.sha256(path.read_bytes()).hexdigest() != prepared['sha256']:
        raise ValueError('已检查的中间资料发生改变，请重新读取资料')
    data = json.loads(path.read_text()); samples = data['samples']; events = sorted(data['events'], key=lambda e: e['eligible'])
    features = data['features']; latest = {}; cursor = 0; counts = Counter(); aligned = {}; audit = []
    max_age = data['max_age_hours'] * 3600
    for sample in sorted(samples, key=lambda s: s['cutoff']):
        while cursor < len(events) and events[cursor]['eligible'] <= sample['cutoff']:
            event = events[cursor]; cursor += 1; key = (event['entity'], event['feature']); old = latest.get(key)
            if old is None or (event['observed'], event['available']) > (old['observed'], old['available']):
                latest[key] = event
        output = dict(sample['values'])
        for feature in features:
            event = latest.get((sample['entity'], feature)); age = None
            reason = 'no_available_record'
            if event:
                age = (datetime.fromisoformat(sample['cutoff']) - datetime.fromisoformat(event['observed'])).total_seconds()
                reason = 'stale' if max_age and age > max_age else 'missing_value' if not event['value'].strip() else 'matched'
            output['asof_' + feature] = event['value'] if reason == 'matched' else ''
            counts[reason] += 1
            audit.append({'sample_id': sample['id'], 'sample_record': sample['record'], 'prediction_time': sample['cutoff'],
                'entity': sample['entity'], 'feature': feature, 'status': reason, 'source_path': data['sources'][1]['path'],
                'source_record': event['record'] if event else '', 'observed_time': event['observed'] if event else '',
                'available_time': event['available'] if event else '', 'age_hours': age / 3600 if age is not None else '',
                'value': output['asof_' + feature]})
        aligned[sample['record']] = output
    rows = [aligned[s['record']] for s in samples]
    folder = Path('results/point-in-time') / str(uuid4()); folder.mkdir(parents=True)
    def save_csv(name, fields, values):
        with (folder / name).open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fields); writer.writeheader(); writer.writerows(values)
    save_csv('samples.csv', data['sample_fields'] + ['asof_' + f for f in features], rows)
    save_csv('matches.csv', list(audit[0]), audit)
    summary = {'samples': len(rows), 'features': features, 'counts': dict(counts), 'sources': data['sources'],
               'max_age_hours': data['max_age_hours'], 'utc_offset': data['utc_offset']}
    (folder / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    markdown = f"# 按当时已知信息制备数据\n\n保留 {len(rows)} 个样本，新增 {len(features)} 个特征。成功匹配 {counts['matched']} 项；无当时可用记录 {counts['no_available_record']} 项；超出时限 {counts['stale']} 项；来源值缺失 {counts['missing_value']} 项。\n\n"
    markdown += '仅使用所属时间和可用时间均不晚于预测时点的版本；优先最新所属时间，再取该期当时最新版本。缺失保留为空，不填0、不自动删除样本。\n\n'
    markdown += '样本表可供后续分析或训练。新增特征以 asof_ 开头；逐项匹配表单独保存来源及时间。原样本自带字段仍需检查预测时是否已知，时间匹配不能保证这些字段没有泄漏。\n\n'
    markdown += f"无时区日期按 {data['utc_offset']} 解释；最长历史时长为 {data['max_age_hours']} 小时（0表示未限制）。当前流程不训练或调用大模型。"
    preview_fields = data['sample_fields'][:3] + ['asof_' + f for f in features[:5]]
    def cell(value):
        return str(value).replace('\\', '\\\\').replace('|', '\\|').replace('\n', ' ').replace('\r', ' ')[:120]
    markdown += '\n\n## 样本预览\n\n' + '| ' + ' | '.join(map(cell, preview_fields)) + ' |\n'
    markdown += '| ' + ' | '.join('---' for _ in preview_fields) + ' |\n'
    markdown += '\n'.join('| ' + ' | '.join(cell(row[k]) for k in preview_fields) + ' |' for row in rows[:12])
    markdown += '\n\n预览最多12行、3个原字段及5个新特征，完整数据见CSV。匹配明细中的 matched 表示已匹配，no_available_record 表示当时无可用记录，stale 表示超时，missing_value 表示来源值缺失。'
    (folder / 'report.md').write_text(markdown, encoding='utf-8')
    return {**summary, 'source_path': str(folder / 'samples.csv'), 'preview': rows[:12], 'markdown': markdown,
            'artifacts': [{'file_path': str(folder / file), 'label': label} for file, label in [
                ('samples.csv', '可供分析训练的样本表 CSV'), ('matches.csv', '每个匹配值的来源与原因 CSV'),
                ('summary.json', '数据版本与匹配汇总 JSON'), ('report.md', '方法与结果说明 Markdown')]]}


def main(inputs):
    if inputs['operation'] == 'prepare':
        return prepare(inputs)
    if inputs['operation'] == 'align':
        return align(inputs)
    raise ValueError('未知的数据制备步骤')
